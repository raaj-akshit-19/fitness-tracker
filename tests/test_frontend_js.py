"""Frontend tests.

Builds temporary workbooks with known scenarios, takes the real backend's
responses for them, and runs the frontend scripts against those responses
under Node (tests/frontend/frontend.test.js). No browser and no extra packages
are needed, only Node itself. The real workbook is not used.

Run from the project root:  python -m unittest discover -s tests -v
"""

import json
import os
import re
import shutil
import subprocess
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from backend import workbook as wbk
from backend.app import create_app
from backend.store import TrackerStore
from tests.earlier_layout import add_month

JS_TESTS = Path(__file__).resolve().parent / "frontend" / "frontend.test.js"
TODAY = date(2026, 10, 20)
MONTHS = [(2026, 10), (2026, 11), (2026, 12), (2027, 1)]


def build_scenario(path):
    """Four months: October with entries, November with one day and one Sunday,
    December empty, January 2027 with a recorded zero for cardio and for weight."""
    wbk.create_workbook(path)
    wb = load_workbook(path)
    for year, month in MONTHS[1:]:
        wbk.add_month_sheet(wb, year, month)
    october, november = wb["October 2026"], wb["November 2026"]
    wb["January 2027"]["D2"], wb["January 2027"]["G2"] = "0:00", 0

    # Row is day + 1. B exercise, C junk food, D cardio, G weight lifted.
    days = {
        1: ("Completed", "None", "25:30", 2450),
        2: ("Completed", "Had", "0:45", 1820.5),
        5: ("Completed", "None", None, None),
        6: ("Completed", None, "60:00", 3000),
        7: ("Completed", "None", None, None),
        18: ("Completed", None, None, None),
        19: ("Completed", "None", None, None),
        20: ("Completed", None, None, None),   # today: junk food not entered yet
    }
    for day, values in days.items():
        for col, value in zip((2, 3, 4, 7), values):
            if value is not None:
                october.cell(day + 1, col, value)
    # Sundays 4, 11, 18, 25 Oct are rows 2 to 5. K weight, L waist.
    october["K2"], october["L2"] = 72.5, 84
    october["K3"] = 72.0
    october["K4"], october["L4"] = 71.1, 82.5

    november["B3"], november["D3"], november["G3"] = "Completed", "10:00", 2000
    november["K2"] = 80   # 1 Nov, the month's first Sunday
    wb.save(path)


def collect_fixtures(path):
    """Backend responses for the scenario."""
    client = create_app(path, today=lambda: TODAY).test_client()
    get = lambda url: client.get(url).get_json()
    return {
        "months": get("/api/months"),
        "status": get("/api/status"),
        "data": {f"{year}-{number}": {"month": get(f"/api/months/{year}/{number}"),
                                      "analytics": get(f"/api/months/{year}/{number}/analytics")}
                 for year, number in MONTHS},
    }


# The edit the editor tests make to 2 October, and the same edit sent to the real backend.
SAMPLE_EDIT = {"exercise": "Completed", "cardio": "30:00", "calories_kcal": None, "protein_g": 120.5}
# The same for the measurements of Sunday 4 October.
SAMPLE_MEASUREMENTS = {"weight_kg": 70.5, "waist_cm": None, "chest_cm": 98}


def build_editing_scenario(path):
    """A workbook for the editors' tests: October with entries of every kind, November (all in the future) empty."""
    wbk.create_workbook(path, 2026, 10)
    TrackerStore(path).create_month(2026, 11)
    wb = load_workbook(path)
    october = wb["October 2026"]
    # Row is day + 1. B exercise, C junk food, D cardio, E calories, F protein, G weight lifted.
    days = {
        1: ("Completed", "None", "25:30", 2200, 140.5, 2450),
        2: ("Partial", "Controlled", None, 1800, None, None),
        3: ("Rest Day", "Had", None, None, None, None),
        4: ("Missed", None, "0:00", 0, 0, 0),
        6: ("Done", None, "25:75", "lots", None, None),      # typed wrongly in Excel
    }
    for day, values in days.items():
        for col, value in zip(range(2, 8), values):
            if value is not None:
                october.cell(day + 1, col, value)
    # Sundays 4, 11, 18, 25 Oct are rows 2 to 5. K weight, L waist, M chest, P forearm.
    october["K2"], october["L2"] = 72.5, 84
    october["K3"], october["M3"], october["P3"] = 72, "wide", 28.5       # chest typed wrongly in Excel
    wb.save(path)


def build_other_layout(path):
    """A workbook whose October sheet is not laid out the way the tracker lays one out."""
    wb = Workbook()
    wb.remove(wb.active)
    add_month(wb, 2026, 10, earlier=True)
    wb.save(path)


def collect_editing_fixtures(path, other_path):
    """Backend responses for the editing scenario, including what it says to each kind of edit."""
    client = create_app(path, today=lambda: TODAY).test_client()
    get = lambda url: client.get(url).get_json()

    def put(url, body, app=client):
        response = app.put(url, json=body)
        return {"status": response.status_code, "body": response.get_json()}

    fixtures = {
        "months": get("/api/months"),
        "status": get("/api/status"),
        "data": {f"2026-{number}": {"month": get(f"/api/months/2026/{number}"),
                                    "analytics": get(f"/api/months/2026/{number}/analytics")}
                 for number in (10, 11)},
        "analytics": {"status": client.get("/api/months/2026/10/analytics").status_code},
    }
    other_client = create_app(other_path, today=lambda: TODAY).test_client()
    with mock.patch("builtins.open", side_effect=PermissionError("locked")):
        locked = put("/api/months/2026/10/days/2", SAMPLE_EDIT)
    fixtures["replies"] = {
        "invalid": put("/api/months/2026/10/days/2", {"exercise": "Done", "calories_kcal": 12.5, "cardio": "9:99"}),
        "future": put("/api/months/2026/10/days/25", {"exercise": "Completed"}),
        "missing": put("/api/months/2026/8/days/2", SAMPLE_EDIT),
        "locked": locked,
        "other_layout": put("/api/months/2026/10/days/2", {"exercise": "Completed"}, other_client),
    }
    with mock.patch("builtins.open", side_effect=PermissionError("locked")):
        locked = put("/api/months/2026/10/measurements/4", SAMPLE_MEASUREMENTS)
    fixtures["measure"] = {"replies": {
        "invalid": put("/api/months/2026/10/measurements/4", {"weight_kg": 0, "waist_cm": 501, "bicep_cm": 35.5}),
        "not_sunday": put("/api/months/2026/10/measurements/5", {"weight_kg": 70}),
        "future": put("/api/months/2026/10/measurements/25", {"weight_kg": 70}),
        "missing": put("/api/months/2026/8/measurements/2", SAMPLE_MEASUREMENTS),
        "locked": locked,
        "other_layout": put("/api/months/2026/10/measurements/4", {"weight_kg": 70}, other_client),
    }}
    # What the backend says about a month it cannot read, for the page's handling of it.
    fixtures["other_layout_month"] = {
        "status": other_client.get("/api/months/2026/10").status_code,
        "body": other_client.get("/api/months/2026/10").get_json(),
    }
    # Last, because these change the workbook.
    fixtures["edit"] = {"request": SAMPLE_EDIT, "reply": put("/api/months/2026/10/days/2", SAMPLE_EDIT)}
    fixtures["measure"]["edit"] = {"request": SAMPLE_MEASUREMENTS,
                                   "reply": put("/api/months/2026/10/measurements/4", SAMPLE_MEASUREMENTS)}
    return fixtures


def build_fixtures(tmp):
    """Every fixture the JavaScript suite uses, made by the real backend in tmp."""
    build_scenario(tmp / "Fitness_Tracker.xlsx")
    (tmp / "editing").mkdir()
    build_editing_scenario(tmp / "editing" / "Fitness_Tracker.xlsx")
    (tmp / "other").mkdir()
    build_other_layout(tmp / "other" / "Fitness_Tracker.xlsx")
    collected = collect_fixtures(tmp / "Fitness_Tracker.xlsx")
    collected["editing"] = collect_editing_fixtures(tmp / "editing" / "Fitness_Tracker.xlsx",
                                                    tmp / "other" / "Fitness_Tracker.xlsx")
    return collected


class ScenarioTests(unittest.TestCase):
    """Pins the backend figures the frontend tests rely on."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        cls.addClassCleanup(shutil.rmtree, cls.tmp, ignore_errors=True)
        cls.fixtures = build_fixtures(cls.tmp)
        cls.editing = cls.fixtures["editing"]

    def test_nothing_served_names_a_version(self):
        text = json.dumps(self.fixtures).lower()
        for word in ('"version"', "upgrade", "migrat", "legacy"):
            self.assertNotIn(word, text, word)
        self.assertEqual(self.fixtures["months"][0], {"year": 2026, "month": 10, "label": "October 2026"})

    def test_editing_scenario(self):
        october = self.editing["data"]["2026-10"]["month"]
        self.assertEqual(set(october), {"year", "month", "label", "days", "measurements", "issues"})
        self.assertEqual(self.editing["status"]["today"], "2026-10-20")
        self.assertEqual(october["days"][0], {
            "date": "2026-10-01", "exercise": "Completed", "junk_food": "None",
            "cardio_display": "25:30", "cardio_seconds": 1530,
            "calories_kcal": 2200, "protein_g": 140.5, "weight_lifted_kg": 2450,
        })
        self.assertEqual(sorted((i["date"], i["field"]) for i in october["issues"]), [
            ("2026-10-06", "calories_kcal"), ("2026-10-06", "cardio"), ("2026-10-06", "exercise"),
            ("2026-10-11", "chest_cm")])
        self.assertEqual([m["sunday"] for m in october["measurements"]],
                         ["2026-10-04", "2026-10-11", "2026-10-18", "2026-10-25"])
        self.assertEqual(october["measurements"][0], {
            "sunday": "2026-10-04", "weight_kg": 72.5, "waist_cm": 84, "chest_cm": None,
            "bicep_cm": None, "thigh_cm": None, "forearm_cm": None,
        })
        self.assertEqual(self.editing["analytics"]["status"], 200)

    def test_editing_analytics(self):
        october = self.editing["data"]["2026-10"]["analytics"]
        self.assertEqual(october["exercise"]["counts"],
                         {"completed": 1, "partial": 1, "rest_day": 1, "missed": 16, "pending": 12})
        self.assertEqual((october["exercise"]["completion_percentage"], october["exercise"]["missed_not_entered"],
                          october["exercise"]["current_streak"], october["exercise"]["longest_streak"]),
                         (8.33, 15, 0, 1))
        self.assertEqual(october["junk_food"]["counts"],
                         {"none": 1, "controlled": 1, "had": 1, "missed": 16, "pending": 12})
        self.assertEqual((october["junk_food"]["success_percentage"], october["junk_food"]["longest_streak"]), (50.0, 1))
        self.assertEqual((october["cardio"]["total_display"], october["cardio"]["average_display"],
                          october["cardio"]["days_recorded"]), ("25:30", "12:45", 2))
        self.assertEqual((october["weight_lifted"]["total_kg"], october["weight_lifted"]["average_kg"],
                          october["weight_lifted"]["lowest_kg"]), (2450, 1225, 0))
        self.assertEqual((october["calories"]["total_kcal"], october["calories"]["average_kcal"],
                          october["calories"]["days_recorded"], october["calories"]["highest_kcal"],
                          october["calories"]["lowest_kcal"]), (4000, 1333.33, 3, 2200, 0))
        self.assertEqual((october["protein"]["total_g"], october["protein"]["average_g"],
                          october["protein"]["days_recorded"]), (140.5, 70.25, 2))
        self.assertEqual((october["measurements"]["weight_kg"]["current"],
                          october["measurements"]["weight_kg"]["change"]), (72, -0.5))
        november = self.editing["data"]["2026-11"]["analytics"]
        self.assertEqual(november["exercise"]["counts"]["pending"], 30)
        self.assertIsNone(november["calories"]["total_kcal"])
        self.assertEqual(november["cardio"]["days_recorded"], 0)

    def test_measurement_replies(self):
        replies = self.editing["measure"]["replies"]
        self.assertEqual({name: reply["status"] for name, reply in replies.items()},
                         {"invalid": 400, "not_sunday": 400, "future": 400, "missing": 404, "locked": 423,
                          "other_layout": 500})
        self.assertEqual({name: reply["body"]["error"]["code"] for name, reply in replies.items()},
                         {"invalid": "invalid_request", "not_sunday": "invalid_request", "future": "invalid_request",
                          "missing": "month_not_found", "locked": "workbook_locked",
                          "other_layout": "workbook_format"})
        self.assertEqual(sorted(replies["invalid"]["body"]["error"]["fields"]), ["waist_cm", "weight_kg"])
        self.assertEqual(list(replies["not_sunday"]["body"]["error"]["fields"]), ["day"])
        self.assertEqual(list(replies["future"]["body"]["error"]["fields"]), ["day"])
        edit = self.editing["measure"]["edit"]["reply"]
        self.assertEqual(edit["status"], 200)
        self.assertEqual(set(edit["body"]), {"year", "month", "label", "measurement", "workbook_modified"})
        self.assertEqual(edit["body"]["measurement"], {
            "sunday": "2026-10-04", "weight_kg": 70.5, "waist_cm": None, "chest_cm": 98,
            "bicep_cm": None, "thigh_cm": None, "forearm_cm": None,
        })

    def test_day_replies(self):
        replies = self.editing["replies"]
        self.assertEqual({name: reply["status"] for name, reply in replies.items()},
                         {"invalid": 400, "future": 400, "missing": 404, "locked": 423, "other_layout": 500})
        self.assertEqual({name: reply["body"]["error"]["code"] for name, reply in replies.items()},
                         {"invalid": "invalid_request", "future": "invalid_request", "missing": "month_not_found",
                          "locked": "workbook_locked", "other_layout": "workbook_format"})
        self.assertEqual(sorted(replies["invalid"]["body"]["error"]["fields"]), ["calories_kcal", "cardio", "exercise"])
        self.assertEqual(list(replies["future"]["body"]["error"]["fields"]), ["day"])
        edit = self.editing["edit"]["reply"]
        self.assertEqual(edit["status"], 200)
        self.assertEqual(edit["body"]["day"], {
            "date": "2026-10-02", "exercise": "Completed", "junk_food": "Controlled",
            "cardio_display": "30:00", "cardio_seconds": 1800,
            "calories_kcal": None, "protein_g": 120.5, "weight_lifted_kg": None,
        })
        self.assertEqual((self.editing["other_layout_month"]["status"],
                          self.editing["other_layout_month"]["body"]["error"]["code"]), (500, "workbook_format"))

    def test_october(self):
        october = self.fixtures["data"]["2026-10"]
        self.assertEqual(october["month"]["days"][0], {
            "date": "2026-10-01", "exercise": "Completed", "junk_food": "None",
            "cardio_display": "25:30", "cardio_seconds": 1530,
            "calories_kcal": None, "protein_g": None, "weight_lifted_kg": 2450,
        })
        self.assertEqual((october["month"]["days"][1]["exercise"], october["month"]["days"][1]["junk_food"]),
                         ("Completed", "Had"))
        self.assertEqual([m["weight_kg"] for m in october["month"]["measurements"]], [72.5, 72.0, 71.1, None])
        analytics = october["analytics"]
        # 8 days completed, today among them; the other 12 days so far were left blank, which is a miss; 11 to come.
        self.assertEqual(analytics["exercise"]["counts"],
                         {"completed": 8, "partial": 0, "rest_day": 0, "missed": 12, "pending": 11})
        self.assertEqual((analytics["exercise"]["current_streak"], analytics["exercise"]["longest_streak"]), (3, 3))
        self.assertEqual(analytics["junk_food"]["counts"],
                         {"none": 4, "controlled": 0, "had": 1, "missed": 14, "pending": 12})
        self.assertEqual((analytics["cardio"]["total_display"], analytics["cardio"]["average_display"],
                          analytics["weight_lifted"]["total_kg"], analytics["calories"]["total_kcal"]),
                         ("86:15", "28:45", 7270.5, None))
        self.assertEqual(analytics["measurements"]["weight_kg"]["change"], -0.9)

    def test_recorded_zero_month(self):
        january = self.fixtures["data"]["2027-1"]["analytics"]
        self.assertEqual(january["cardio"]["days_recorded"], 1)
        self.assertEqual(january["cardio"]["total_display"], "0:00")
        self.assertEqual(january["cardio"]["average_display"], "0:00")
        self.assertEqual(january["weight_lifted"]["days_recorded"], 1)
        self.assertEqual(january["weight_lifted"]["total_kg"], 0)
        self.assertEqual(january["weight_lifted"]["highest_kg"], 0)

    def test_other_months(self):
        november = self.fixtures["data"]["2026-11"]["analytics"]
        december = self.fixtures["data"]["2026-12"]["analytics"]
        self.assertIsNone(november["exercise"]["completion_percentage"])
        self.assertIsNone(november["measurements"]["weight_kg"]["previous"])
        self.assertEqual(november["measurements"]["weight_kg"]["current"], 80)
        self.assertEqual(december["cardio"]["days_recorded"], 0)
        self.assertEqual(december["weight_lifted"]["days_recorded"], 0)


@unittest.skipUnless(shutil.which("node"), "Node.js is not installed")
class FrontendJavaScriptTests(unittest.TestCase):
    def test_frontend_suite(self):
        tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)
        collected = build_fixtures(tmp)
        fixtures = tmp / "fixtures.json"
        fixtures.write_text(json.dumps(collected))

        result = subprocess.run(
            ["node", "--test", str(JS_TESTS)],
            env={**os.environ, "FIXTURES": str(fixtures)},
            capture_output=True, text=True, encoding="utf-8",
        )
        output = result.stdout + result.stderr
        self.assertEqual(result.returncode, 0, output)
        passed = re.search(r"pass (\d+)", output)
        failed = re.search(r"fail (\d+)", output)
        self.assertIsNotNone(passed, output)
        self.assertGreaterEqual(int(passed.group(1)), 80, output)
        self.assertEqual(int(failed.group(1)), 0, output)
        print(f"\n  JavaScript frontend tests: {passed.group(1)} passed", end="")


if __name__ == "__main__":
    unittest.main()

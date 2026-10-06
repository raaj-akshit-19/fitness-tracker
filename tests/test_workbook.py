"""Checks for the sheet layout and the clean template.

Everything is built in temporary folders. The template file is only read.

Run from the project root:  python -m unittest discover -s tests -v
"""

import hashlib
import shutil
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from backend import validation as v
from backend import workbook as wbk

PROJECT = Path(__file__).resolve().parent.parent
TEMPLATE = PROJECT / "Fitness_Tracker_Template.xlsx"


def as_date(value):
    return value.date() if isinstance(value, datetime) else value


def build(year=2026, month=10):
    wb = Workbook()
    wb.remove(wb.active)
    return wbk.add_month_sheet(wb, year, month)


def rules(ws):
    return {str(rule.sqref): rule for rule in ws.data_validations.dataValidation}


class LayoutTests(unittest.TestCase):
    def test_daily_headers_are_a_to_g(self):
        ws = build()
        self.assertEqual([ws.cell(1, col).value for col in range(1, 8)], [
            "Date", "Exercise", "Junk Food", "Cardio (min:sec)", "Calories (kcal)",
            "Protein (g)", "Weight Lifted (kg)",
        ])
        self.assertEqual(wbk.DAILY_HEADERS[1:3], ["Exercise", "Junk Food"])

    def test_measurements_are_j_to_p_with_two_blank_columns_before(self):
        ws = build()
        self.assertIsNone(ws.cell(1, 8).value)
        self.assertIsNone(ws.cell(1, 9).value)
        self.assertEqual([ws.cell(1, col).value for col in range(10, 17)], [
            "Sunday", "Weight (kg)", "Waist (cm)", "Chest (cm)", "Bicep (cm)", "Thigh (cm)", "Forearm (cm)",
        ])
        self.assertIsNone(ws.cell(1, 17).value)
        self.assertEqual(wbk.MEASUREMENT_FIRST_COL, 10)

    def test_tables_keep_their_names_and_cover_the_new_ranges(self):
        cases = {(2026, 10): ("A1:G32", "J1:P5"), (2026, 11): ("A1:G31", "J1:P6"),
                 (2027, 2): ("A1:G29", "J1:P5"), (2028, 2): ("A1:G30", "J1:P5")}
        for (year, month), (daily, measurements) in cases.items():
            ws = build(year, month)
            self.assertEqual({name: ws.tables[name].ref for name in ws.tables}, {
                f"Daily_{year}_{month:02d}": daily, f"Measurements_{year}_{month:02d}": measurements,
            })

    def test_dates_and_sundays_follow_the_calendar(self):
        for year, month in ((2026, 10), (2026, 11), (2027, 2), (2028, 2), (2027, 4)):
            ws = build(year, month)
            days, sundays = wbk.month_dates(year, month), wbk.month_sundays(year, month)
            self.assertEqual([as_date(ws.cell(row, 1).value) for row in range(2, len(days) + 2)], days)
            self.assertIsNone(ws.cell(len(days) + 2, 1).value)
            self.assertEqual([as_date(ws.cell(row, 10).value) for row in range(2, len(sundays) + 2)], sundays)
            self.assertIsNone(ws.cell(len(sundays) + 2, 10).value)
            self.assertTrue(all(day.weekday() == 6 for day in sundays))

    def test_one_sheet_per_month_named_as_before(self):
        wb = Workbook()
        wb.remove(wb.active)
        for year, month in ((2026, 10), (2026, 11), (2027, 1)):
            wbk.add_month_sheet(wb, year, month)
        self.assertEqual(wb.sheetnames, ["October 2026", "November 2026", "January 2027"])
        self.assertEqual(wbk.list_months(wb), [(2026, 10), (2026, 11), (2027, 1)])
        with self.assertRaises(ValueError):
            wbk.add_month_sheet(wb, 2026, 10)

    def test_every_data_cell_starts_blank(self):
        ws = build()
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.column not in (1, 10):
                    self.assertIsNone(cell.value, cell.coordinate)
        self.assertEqual((ws.max_row, ws.max_column), (32, 16))

    def test_cell_formats(self):
        ws = build()
        for row in range(2, 33):
            self.assertEqual([ws.cell(row, col).number_format for col in range(1, 8)],
                             ["ddd, dd mmm yyyy", "@", "@", "@", "0", "0.0", "0.0"])
        for row in range(2, 6):
            self.assertEqual(ws.cell(row, 10).number_format, "ddd, dd mmm yyyy")
            self.assertEqual({ws.cell(row, col).number_format for col in range(11, 17)}, {"0.0"})

    def test_no_formulas_goals_or_old_columns(self):
        ws = build()
        text = " ".join(str(cell.value) for row in ws.iter_rows() for cell in row if isinstance(cell.value, str))
        self.assertNotIn("=", text)
        for word in ("No Junk Food", "Total Weight Lifted", "Goal", "Target", "Meditation"):
            self.assertNotIn(word, text)


class ExcelValidationTests(unittest.TestCase):
    def setUp(self):
        self.ws = build()
        self.rules = rules(self.ws)

    def test_one_rule_per_column_group(self):
        self.assertEqual(sorted(self.rules), ["B2:B32", "C2:C32", "D2:D32", "E2:E32", "F2:F32", "G2:G32", "K2:P5"])
        for rule in self.rules.values():
            self.assertTrue(rule.allow_blank)
            self.assertTrue(rule.showErrorMessage)
            self.assertTrue(rule.error)

    def test_exercise_column_offers_exactly_the_four_statuses(self):
        rule = self.rules["B2:B32"]
        self.assertEqual(rule.type, "list")
        self.assertEqual(rule.formula1, '"Completed,Partial,Rest Day,Missed"')
        self.assertEqual(rule.formula1.strip('"').split(","), list(v.EXERCISE_STATUSES))

    def test_junk_food_column_offers_exactly_the_three_statuses(self):
        rule = self.rules["C2:C32"]
        self.assertEqual(rule.type, "list")
        self.assertEqual(rule.formula1, '"None,Controlled,Had"')
        self.assertEqual(rule.formula1.strip('"').split(","), list(v.JUNK_FOOD_STATUSES))

    def test_list_rules_stay_within_what_excel_allows(self):
        for cells in ("B2:B32", "C2:C32"):
            self.assertLessEqual(len(self.rules[cells].formula1), 255)

    def test_cardio_is_checked_as_min_sec_text(self):
        rule = self.rules["D2:D32"]
        self.assertEqual(rule.type, "custom")
        self.assertIn('FIND(":",D2)', rule.formula1)
        self.assertIn("<60", rule.formula1)

    def test_number_columns_use_the_limits_from_the_validation_module(self):
        calories, protein, lifted = self.rules["E2:E32"], self.rules["F2:F32"], self.rules["G2:G32"]
        self.assertEqual((calories.type, calories.operator, calories.formula1, calories.formula2),
                         ("whole", "between", "0", str(v.CALORIES_MAX)))
        self.assertEqual((protein.type, protein.formula1, protein.formula2), ("decimal", "0", str(v.PROTEIN_MAX)))
        self.assertEqual((lifted.type, lifted.formula1, lifted.formula2), ("decimal", "0", str(v.WEIGHT_LIFTED_MAX)))

    def test_measurements_must_be_above_zero_and_within_the_limit(self):
        rule = self.rules["K2:P5"]
        self.assertEqual(rule.type, "custom")
        self.assertEqual(rule.formula1, f"AND(ISNUMBER(K2),K2>0,K2<={v.MEASUREMENT_MAX})")

    def test_rules_cover_five_sunday_months(self):
        self.assertIn("K2:P6", rules(build(2026, 11)))
        self.assertIn("B2:B31", rules(build(2026, 11)))


class OneLayoutTests(unittest.TestCase):
    def test_the_layout_is_recognised_from_the_headers(self):
        self.assertIs(wbk.is_month_sheet(build()), True)
        wb = Workbook()
        self.assertIs(wbk.is_month_sheet(wb.active), False)
        changed = build()
        changed["C1"] = "No Junk Food"
        self.assertIs(wbk.is_month_sheet(changed), False)

    def test_there_is_one_layout_and_no_way_to_ask_for_another(self):
        import inspect
        for function in (wbk.add_month_sheet, wbk.create_workbook):
            self.assertNotIn("version", inspect.signature(function).parameters)
        with self.assertRaises(TypeError):
            wbk.add_month_sheet(Workbook(), 2026, 10, 1)
        source = inspect.getsource(wbk)
        for word in ("version", "V2", "No Junk Food", "Total Weight Lifted", "TRUE,FALSE"):
            self.assertNotIn(word, source, word)
        self.assertEqual(wbk.DAILY_HEADERS, ["Date", "Exercise", "Junk Food", "Cardio (min:sec)", "Calories (kcal)",
                                             "Protein (g)", "Weight Lifted (kg)"])
        self.assertEqual(wbk.MEASUREMENT_HEADERS, ["Sunday", "Weight (kg)", "Waist (cm)", "Chest (cm)", "Bicep (cm)",
                                                   "Thigh (cm)", "Forearm (cm)"])

    def test_a_saved_workbook_reads_back_the_same(self):
        tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)
        path = wbk.create_workbook(tmp / "book.xlsx")
        with self.assertRaises(FileExistsError):
            wbk.create_workbook(path)
        ws = load_workbook(path)["October 2026"]
        self.assertIs(wbk.is_month_sheet(ws), True)
        self.assertEqual({n: ws.tables[n].ref for n in ws.tables},
                         {"Daily_2026_10": "A1:G32", "Measurements_2026_10": "J1:P5"})
        self.assertEqual(len(ws.data_validations.dataValidation), 7)
        # Words, text durations and numbers survive a save in the form validation.py stores them.
        wb = load_workbook(path)
        ws = wb["October 2026"]
        ws["B2"], ws["C2"], ws["D2"] = v.validate_exercise("rest day"), v.validate_junk_food("None"), v.validate_cardio("05:30")
        ws["E2"], ws["F2"], ws["G2"], ws["K2"] = 2200, 140.5, 2450, 72.5
        wb.save(path)
        row = [c.value for c in load_workbook(path)["October 2026"][2]][1:7]
        self.assertEqual(row, ["Rest Day", "None", "5:30", 2200, 140.5, 2450])


class TemplateTests(unittest.TestCase):
    """The clean workbook kept in the repository."""

    @classmethod
    def setUpClass(cls):
        cls.before = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
        cls.wb = load_workbook(TEMPLATE)

    def test_is_the_trackers_layout_with_october_2026(self):
        self.assertEqual(self.wb.sheetnames, ["October 2026"])
        ws = self.wb["October 2026"]
        self.assertIs(wbk.is_month_sheet(ws), True)
        self.assertEqual({n: ws.tables[n].ref for n in ws.tables},
                         {"Daily_2026_10": "A1:G32", "Measurements_2026_10": "J1:P5"})
        self.assertEqual([as_date(ws.cell(row, 1).value) for row in range(2, 33)], wbk.month_dates(2026, 10))
        self.assertEqual([as_date(ws.cell(row, 10).value) for row in range(2, 6)],
                         [date(2026, 10, d) for d in (4, 11, 18, 25)])
        self.assertEqual(len(ws.data_validations.dataValidation), 7)

    def test_holds_no_data_and_no_personal_name(self):
        ws = self.wb["October 2026"]
        entered = [c.coordinate for row in ws.iter_rows(min_row=2) for c in row
                   if c.value is not None and c.column not in (1, 10)]
        self.assertEqual(entered, [])
        self.assertEqual(self.wb.properties.creator, "openpyxl")
        self.assertIn(self.wb.properties.lastModifiedBy, (None, "", "openpyxl"))

    def test_reading_it_did_not_change_it(self):
        self.assertEqual(hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(), self.before)


if __name__ == "__main__":
    unittest.main()

"""Checks for the HTTP API, run against a temporary workbook.

Run from the project root:  python -m unittest discover -s tests -v
"""

import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from backend import store as st
from backend import workbook as wbk
from backend.app import create_app


class ApiTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.path = wbk.create_workbook(self.tmp / "Fitness_Tracker.xlsx")
        self.client = create_app(self.path).test_client()

    def assert_error(self, response, status, code):
        self.assertEqual(response.status_code, status)
        body = response.get_json()
        self.assertEqual(body["error"]["code"], code)
        self.assertTrue(body["error"]["message"])

    def test_list_months(self):
        response = self.client.get("/api/months")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.get_json(), [{"year": 2026, "month": 10, "label": "October 2026"}]
        )

    def test_get_month(self):
        response = self.client.get("/api/months/2026/10")
        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertEqual(len(body["days"]), 31)
        self.assertEqual(len(body["measurements"]), 4)
        self.assertEqual(body["days"][0]["date"], "2026-10-01")
        self.assertIsNone(body["days"][0]["exercise"])
        self.assertIsNone(body["days"][0]["cardio_seconds"])
        self.assertIsNone(body["days"][0]["weight_lifted_kg"])
        self.assertEqual(set(body["days"][0]), {"date", "exercise", "junk_food", "cardio_display", "cardio_seconds",
                                               "calories_kcal", "protein_g", "weight_lifted_kg"})

    def test_get_missing_month(self):
        self.assert_error(self.client.get("/api/months/2026/11"), 404, "month_not_found")

    def test_get_invalid_month(self):
        self.assert_error(self.client.get("/api/months/2026/13"), 400, "invalid_month")

    def test_create_month(self):
        response = self.client.post("/api/months", json={"year": 2026, "month": 11})
        self.assertEqual(response.status_code, 201)
        body = response.get_json()
        self.assertEqual(body["label"], "November 2026")
        self.assertEqual(len(body["days"]), 30)
        self.assertEqual(len(body["measurements"]), 5)
        self.assertEqual(len(self.client.get("/api/months").get_json()), 2)

    def test_create_duplicate_month(self):
        self.assert_error(
            self.client.post("/api/months", json={"year": 2026, "month": 10}),
            409, "month_exists",
        )

    # ---- deleting a month

    def delete(self, year, month, confirm="?"):
        name = wbk.sheet_name(year, month) if confirm == "?" else confirm
        return self.client.delete(f"/api/months/{year}/{month}", json={"confirm": name})

    def sheets(self):
        return load_workbook(self.path).sheetnames

    def test_delete_month_removes_its_sheet_and_nothing_else(self):
        self.client.post("/api/months", json={"year": 2026, "month": 11})
        self.client.post("/api/months", json={"year": 2026, "month": 12})
        self.client.put("/api/months/2026/10/days/1", json={"exercise": "Completed", "calories_kcal": 2100})
        october = self.client.get("/api/months/2026/10").get_json()
        december = self.client.get("/api/months/2026/12").get_json()

        response = self.delete(2026, 11)
        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertEqual(body["deleted"], {"year": 2026, "month": 11, "label": "November 2026"})
        self.assertEqual([month["label"] for month in body["months"]], ["October 2026", "December 2026"])
        # The sheet is gone from the file itself, with its two tables; the others are as they were.
        wb = load_workbook(self.path)
        self.assertEqual(wb.sheetnames, ["October 2026", "December 2026"])
        self.assertEqual(sorted(name for ws in wb for name in ws.tables),
                         ["Daily_2026_10", "Daily_2026_12", "Measurements_2026_10", "Measurements_2026_12"])
        self.assertEqual(self.client.get("/api/months/2026/10").get_json(), october)
        self.assertEqual(self.client.get("/api/months/2026/12").get_json(), december)
        self.assert_error(self.client.get("/api/months/2026/11"), 404, "month_not_found")
        self.assertEqual(sorted(p.name for p in self.tmp.iterdir()), ["Fitness_Tracker.xlsx"])     # no copy is kept
        # It can be added again, empty, and deleted again.
        self.assertEqual(self.client.post("/api/months", json={"year": 2026, "month": 11}).status_code, 201)
        self.assertEqual(self.sheets(), ["October 2026", "November 2026", "December 2026"])
        self.assertEqual(self.delete(2026, 11).status_code, 200)

    def test_the_only_month_cannot_be_deleted(self):
        before = self.path.read_bytes()
        response = self.delete(2026, 10)
        self.assert_error(response, 409, "last_month")
        self.assertIn("only month", response.get_json()["error"]["message"])
        self.assertEqual(self.path.read_bytes(), before)
        # With two months either can go, and then the one left cannot.
        self.client.post("/api/months", json={"year": 2026, "month": 11})
        self.assertEqual(self.delete(2026, 10).status_code, 200)
        self.assert_error(self.delete(2026, 11), 409, "last_month")
        self.assertEqual(self.sheets(), ["November 2026"])

    def test_delete_needs_the_month_named_in_the_request(self):
        self.client.post("/api/months", json={"year": 2026, "month": 11})
        before = self.path.read_bytes()
        url = "/api/months/2026/11"
        for response in (self.client.delete(url), self.client.delete(url, data="nope"),
                         self.client.delete(url, json={}), self.client.delete(url, json={"confirm": True}),
                         self.client.delete(url, json={"confirm": "October 2026"}),
                         self.client.delete(url, json=["November 2026"])):
            self.assert_error(response, 400, "invalid_request")
        self.assertEqual(self.path.read_bytes(), before)

    def test_delete_a_month_that_is_not_there_or_not_a_month(self):
        self.client.post("/api/months", json={"year": 2026, "month": 11})
        before = self.path.read_bytes()
        self.assert_error(self.delete(2027, 1), 404, "month_not_found")
        self.assert_error(self.client.delete("/api/months/2026/13", json={"confirm": "x"}), 400, "invalid_month")
        self.assert_error(self.client.delete("/api/months/1800/1", json={"confirm": "January 1800"}), 400, "invalid_month")
        self.assertEqual(self.path.read_bytes(), before)

    def test_delete_with_the_workbook_locked_changes_nothing(self):
        self.client.post("/api/months", json={"year": 2026, "month": 11})
        before = self.path.read_bytes()
        with mock.patch("builtins.open", side_effect=PermissionError("locked")):
            self.assert_error(self.delete(2026, 11), 423, "workbook_locked")
        with mock.patch.object(st.os, "replace", side_effect=PermissionError("locked")):
            self.assert_error(self.delete(2026, 11), 423, "workbook_locked")
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(sorted(p.name for p in self.tmp.iterdir()), ["Fitness_Tracker.xlsx"])

    def test_delete_is_only_taken_from_the_trackers_own_page(self):
        self.client.post("/api/months", json={"year": 2026, "month": 11})
        response = self.client.delete("/api/months/2026/11", json={"confirm": "November 2026"},
                                      headers={"Origin": "https://example.com"})
        self.assert_error(response, 403, "forbidden_origin")
        self.assertEqual(self.sheets(), ["October 2026", "November 2026"])

    def test_delete_leaves_sheets_that_are_not_months_alone(self):
        wb = load_workbook(self.path)
        wb.create_sheet("Notes")["A1"] = "keep me"
        wb.save(self.path)
        # Another sheet does not count as a month: October is still the only one.
        self.assert_error(self.delete(2026, 10), 409, "last_month")
        self.client.post("/api/months", json={"year": 2026, "month": 11})
        self.assertEqual(self.delete(2026, 10).status_code, 200)
        wb = load_workbook(self.path)
        self.assertEqual(sorted(wb.sheetnames), ["Notes", "November 2026"])
        self.assertEqual(wb["Notes"]["A1"].value, "keep me")

    def test_create_with_bad_input(self):
        post = self.client.post
        self.assert_error(post("/api/months", json={"year": 2026}), 400, "invalid_request")
        self.assert_error(post("/api/months", data="nope"), 400, "invalid_request")
        self.assert_error(post("/api/months", json=[2026, 11]), 400, "invalid_request")
        self.assert_error(
            post("/api/months", json={"year": "2026", "month": 11}), 400, "invalid_month"
        )
        self.assert_error(
            post("/api/months", json={"year": 2026, "month": 13}), 400, "invalid_month"
        )
        self.assertEqual(len(self.client.get("/api/months").get_json()), 1)

    def test_goals_are_not_accepted_or_stored(self):
        response = self.client.post(
            "/api/months", json={"year": 2026, "month": 11, "goals": {"exercise": 20}}
        )
        self.assertEqual(response.status_code, 201)
        self.assertNotIn("goal", response.get_data(as_text=True).lower())

    def test_locked_workbook(self):
        with mock.patch.object(st.os, "replace", side_effect=PermissionError("locked")):
            response = self.client.post("/api/months", json={"year": 2026, "month": 11})
        self.assert_error(response, 423, "workbook_locked")
        self.assertEqual(len(self.client.get("/api/months").get_json()), 1)

    def test_missing_workbook(self):
        client = create_app(self.tmp / "nope.xlsx").test_client()
        self.assert_error(client.get("/api/months"), 500, "workbook_missing")

    def test_status(self):
        body = self.client.get("/api/status").get_json()
        self.assertEqual(set(body), {"workbook_modified", "locked", "unsaved_changes", "today"})
        self.assertEqual(body["today"], date.today().isoformat())
        self.assertIs(body["unsaved_changes"], False)
        self.assertIs(body["locked"], False)

    def test_unknown_route_and_method_return_json(self):
        self.assert_error(self.client.get("/api/nothing"), 404, "not_found")
        self.assert_error(self.client.delete("/api/months"), 405, "method_not_allowed")


if __name__ == "__main__":
    unittest.main()

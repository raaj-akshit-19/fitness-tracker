"""Checks for the packaged Windows application, release/Fitness Tracker/.

The program is copied into temporary folders, each with its own temporary
workbook, and run from there on a spare port. No browser window is opened: the
BROWSER environment variable points at a script that records the address.
The real workbook and the one in the release folder are only hashed.

The tests that run the program are skipped until it has been built with
windows_build/build.py.

Run from the project root:  python -m unittest discover -s tests -v
"""

import hashlib
import importlib
import json
import os
import shutil
import signal
import socket
import subprocess
import sys
import tempfile
import time
import types
import unittest
import urllib.error
import urllib.request
import zipfile
from datetime import date
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from backend import launch, paths
from backend import store as st
from backend import workbook as wbk

PROJECT = Path(__file__).resolve().parent.parent
RELEASE = PROJECT / "release" / "Fitness Tracker"
EXE = RELEASE / "Fitness Tracker.exe"
BUILT = os.name == "nt" and EXE.is_file()
TEMP = Path(tempfile.gettempdir())
PAGE_FILES = ["index.html", "styles.css", "format.js", "analytics.js", "app.js", "editor.js"]


def _source_changed_since_the_build():
    """True when the page files or the backend are newer than the built program.

    The program in release/ is only rebuilt for a release, so during
    development it can be older than the source it was made from. What it
    does then is what the older source did, and is not checked here.
    """
    if not BUILT:
        return False
    built = EXE.stat().st_mtime
    sources = [PROJECT / "frontend" / name for name in PAGE_FILES] + sorted((PROJECT / "backend").glob("*.py"))
    return any(path.stat().st_mtime > built for path in sources)


PAGE_CHANGED = _source_changed_since_the_build()
REBUILD = "the page or the backend has changed since the program was built; rebuild it to compare"


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def free_port():
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        return probe.getsockname()[1]


def port_open(port, host="127.0.0.1"):
    try:
        with socket.create_connection((host, port), timeout=1):
            return True
    except OSError:
        return False


def wait_for(condition, timeout=60):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if condition():
            return True
        time.sleep(0.2)
    return False


def get(url):
    with urllib.request.urlopen(url, timeout=10) as response:
        return response.read()


def post_month(base, year, month):
    request = urllib.request.Request(
        f"{base}api/months", data=json.dumps({"year": year, "month": month}).encode(),
        headers={"Content-Type": "application/json"}, method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=15) as response:
            return response.status, json.loads(response.read())
    except urllib.error.HTTPError as error:
        return error.code, json.loads(error.read())


def running_copies(folder):
    """Process ids of the program running from the given folder."""
    command = ("Get-CimInstance Win32_Process -Filter \"Name='Fitness Tracker.exe'\" | "
               "ForEach-Object { \"$($_.ProcessId)|$($_.ExecutablePath)\" }")
    output = subprocess.run(["powershell", "-NoProfile", "-Command", command],
                            capture_output=True, text=True).stdout
    found = []
    for line in output.splitlines():
        pid, _, path = line.partition("|")
        if path.lower().startswith(str(folder).lower()):
            found.append(int(pid))
    return found


def sockets_of(pids):
    """(local address, remote address, state) of every TCP socket the processes own."""
    output = subprocess.run(["netstat", "-ano", "-p", "TCP"], capture_output=True, text=True).stdout
    rows = []
    for line in output.splitlines():
        parts = line.split()
        if len(parts) == 5 and parts[0] == "TCP" and parts[4].isdigit() and int(parts[4]) in pids:
            rows.append((parts[1], parts[2], parts[3]))
    return rows


class PathTests(unittest.TestCase):
    """How the program finds its files, from source and when packaged."""

    def tearDown(self):
        importlib.reload(paths)

    def test_from_source_everything_is_in_the_project_folder(self):
        self.assertFalse(paths.FROZEN)
        self.assertEqual(paths.APP_DIR, PROJECT)
        self.assertEqual(paths.RESOURCE_DIR, PROJECT)

    def test_packaged_the_workbook_folder_is_where_the_program_file_is(self):
        folder = Path(tempfile.gettempdir()) / "Some Folder (2) & more"
        unpacked = Path(tempfile.gettempdir()) / "_MEIexample"
        with mock.patch.object(sys, "frozen", True, create=True), \
                mock.patch.object(sys, "_MEIPASS", str(unpacked), create=True), \
                mock.patch.object(sys, "executable", str(folder / "Fitness Tracker.exe")):
            importlib.reload(paths)
            self.assertTrue(paths.FROZEN)
            self.assertEqual(paths.APP_DIR, folder.resolve())
            self.assertEqual(paths.RESOURCE_DIR, unpacked)
            self.assertNotEqual(paths.APP_DIR, PROJECT)
            # The current directory plays no part.
            previous = os.getcwd()
            try:
                os.chdir(PROJECT)
                importlib.reload(paths)
                self.assertEqual(paths.APP_DIR, folder.resolve())
            finally:
                os.chdir(previous)

    def test_modules_take_their_paths_from_there(self):
        from backend import app
        self.assertEqual(app.FRONTEND_DIR, paths.RESOURCE_DIR / "frontend")
        self.assertEqual(launch.PROJECT_DIR, paths.APP_DIR)
        if not os.environ.get("FITNESS_TRACKER_WORKBOOK"):
            self.assertEqual(wbk.WORKBOOK_PATH, paths.APP_DIR / "Fitness_Tracker.xlsx")
        for name in ("paths.py", "workbook.py", "app.py", "launch.py"):
            text = (PROJECT / "backend" / name).read_text(encoding="utf-8")
            self.assertNotIn("getcwd", text)
            self.assertNotIn("chdir", text)


class StartupMessageTests(unittest.TestCase):
    """The checks made before the server starts, and what they say."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def test_missing_page_files(self):
        (self.tmp / "index.html").write_text("x")
        with self.assertRaises(launch.LaunchError) as caught:
            launch.check_page_files(self.tmp)
        message = str(caught.exception)
        self.assertIn("Files the dashboard needs are missing", message)
        self.assertIn("styles.css", message)
        self.assertNotIn("index.html", message)
        launch.check_page_files(PROJECT / "frontend")

    def test_missing_workbook_names_the_folder_to_put_it_in(self):
        missing = self.tmp / "Fitness_Tracker.xlsx"
        with self.assertRaises(launch.LaunchError) as caught:
            launch.check_workbook(missing)
        self.assertIn("The workbook was not found", str(caught.exception))
        self.assertIn(str(self.tmp), str(caught.exception))
        self.assertIn("never makes or replaces a workbook", str(caught.exception))
        self.assertFalse(missing.exists())

    def test_damaged_workbook_is_reported_and_left_alone(self):
        damaged = self.tmp / "Fitness_Tracker.xlsx"
        damaged.write_bytes(b"this is not a workbook")
        with self.assertRaises(launch.LaunchError) as caught:
            launch.check_workbook_opens(damaged)
        self.assertIn("The workbook could not be opened", str(caught.exception))
        self.assertIn("Nothing was changed", str(caught.exception))
        self.assertEqual(damaged.read_bytes(), b"this is not a workbook")

    def test_locked_workbook_does_not_stop_the_start(self):
        # Open in Excel is not damage: the dashboard loads it once it can be read.
        workbook = wbk.create_workbook(self.tmp / "Fitness_Tracker.xlsx")
        lines = []
        with mock.patch.object(st, "load_workbook", side_effect=PermissionError("locked")), \
                mock.patch.object(launch, "say", side_effect=lambda text="": lines.append(text)):
            launch.check_workbook_opens(workbook)
        self.assertTrue(any(line.startswith("Note: ") for line in lines))

    def test_good_workbook_passes_and_is_not_changed(self):
        workbook = wbk.create_workbook(self.tmp / "Fitness_Tracker.xlsx")
        before = workbook.read_bytes()
        launch.check_workbook(workbook)
        launch.check_workbook_opens(workbook)
        self.assertEqual(workbook.read_bytes(), before)

    def test_packaged_copy_missing_a_library_says_to_replace_the_program(self):
        with mock.patch.object(launch, "FROZEN", True), \
                mock.patch.object(launch.importlib.util, "find_spec", return_value=None):
            with self.assertRaises(launch.LaunchError) as caught:
                launch.check_dependencies()
        self.assertIn("This copy of Fitness Tracker is incomplete", str(caught.exception))
        self.assertNotIn("pip", str(caught.exception))

    def test_a_failed_start_is_never_silent(self):
        lines = []
        with mock.patch.object(launch, "say", side_effect=lambda text="": lines.append(text)), \
                mock.patch.object(launch, "main", return_value=1), \
                mock.patch.object(sys, "stdin", None):
            self.assertEqual(launch.run([]), 1)
        self.assertIn("Fitness Tracker did not start.", lines)

    def test_an_unexpected_problem_is_shown_not_swallowed(self):
        lines = []
        with mock.patch.object(launch, "say", side_effect=lambda text="": lines.append(text)), \
                mock.patch.object(launch, "main", side_effect=RuntimeError("boom")), \
                mock.patch.object(sys, "stdin", None):
            self.assertEqual(launch.run([]), 1)
        self.assertTrue(any("unexpected problem" in line and "boom" in line for line in lines))

    def test_window_waits_for_a_key_only_after_a_failure(self):
        keyboard = mock.Mock()
        keyboard.isatty.return_value = True
        with mock.patch.object(launch, "say"), mock.patch.object(sys, "stdin", keyboard), \
                mock.patch("builtins.input") as pressed, mock.patch.object(launch.time, "sleep"):
            with mock.patch.object(launch, "main", return_value=1):
                launch.run([])
            self.assertEqual(pressed.call_count, 1)
            with mock.patch.object(launch, "main", return_value=0):
                launch.run([])
            self.assertEqual(pressed.call_count, 1)


class ReleaseFolderTests(unittest.TestCase):
    """What is, and is not, handed to the end user."""

    @unittest.skipUnless(BUILT, "the application has not been built")
    def test_release_folder_holds_only_what_the_user_needs(self):
        # (Excel adds a "~$" lock file beside a workbook while it has it open.)
        names = sorted(p.name for p in RELEASE.iterdir() if not p.name.startswith("~$"))
        self.assertEqual(names, ["Fitness Tracker.exe", "Fitness_Tracker.xlsx", "How to use.txt"])
        self.assertEqual([p.name for p in (PROJECT / "release").iterdir()], ["Fitness Tracker"])
        self.assertGreater(EXE.stat().st_size, 5_000_000)

    @unittest.skipUnless(BUILT, "the application has not been built")
    def test_user_guide_covers_the_essentials(self):
        # The guide the build writes. (The copy in release/ is that text, as of the last build.)
        guide = importlib.import_module("windows_build.build").GUIDE
        for phrase in ("Double-click", "Fitness_Tracker.xlsx", "Save the workbook", "Select Month",
                       "Add New Month", "close the Fitness Tracker window", "BACK UP", "blank", "same folder"):
            self.assertIn(phrase, guide)
        # Entering data from the dashboard.
        for phrase in ('Press "Edit today"', "press Save", "Cancel throws your changes away",
                       "Completed, Partial, Rest Day or Missed", "None, Controlled or Had",
                       "min:sec", "kcal", "Days in the future cannot be edited yet",
                       "close it in Excel, then press Save again"):
            self.assertIn(phrase, guide, phrase)
        for word in ("pip", "python", "_MEI", "unpack", "—"):
            self.assertNotIn(word, guide.lower())
        # One tracker: it speaks of no earlier or later one, and of nothing to convert.
        for word in ("first version", "version 1", "version 2", "upgrade", "TRUE", "FALSE", "No Junk Food"):
            self.assertNotIn(word, guide, word)
        # It no longer says that data can only be entered in Excel.
        self.assertNotIn("Enter your data directly in Fitness_Tracker.xlsx", guide)

    @unittest.skipUnless(BUILT, "the application has not been built")
    @unittest.skipIf(PAGE_CHANGED, REBUILD)
    def test_release_guide_is_the_text_the_build_writes(self):
        guide = (RELEASE / "How to use.txt").read_text(encoding="utf-8")
        # It is the text the build writes, and nothing else.
        build = importlib.import_module("windows_build.build")
        self.assertEqual(guide.replace("\r\n", "\n"), build.GUIDE)

    @unittest.skipUnless(BUILT, "the application has not been built")
    def test_release_workbook_is_clean_and_in_the_trackers_layout(self):
        path = RELEASE / "Fitness_Tracker.xlsx"
        before = sha256(path)
        wb = load_workbook(path)                    # only read; never saved
        months = wbk.list_months(wb)
        self.assertTrue(months)
        self.assertEqual([wbk.is_month_sheet(ws) for ws in wb], [True] * len(wb.sheetnames))
        store = st.TrackerStore(path)
        for year, month in months:
            data = store.get_month(year, month)
            self.assertEqual(data["issues"], [])
            for day in data["days"]:
                self.assertEqual({key: value for key, value in day.items() if key != "date" and value is not None},
                                 {}, f"the release workbook has an entry on {day['date']}")
            for row in data["measurements"]:
                self.assertEqual({key: value for key, value in row.items() if key != "sunday" and value is not None},
                                 {}, f"the release workbook has a measurement on {row['sunday']}")
        # It is not the working workbook of whoever built it, and names nobody.
        if wbk.WORKBOOK_PATH.is_file():
            self.assertNotEqual(before, sha256(wbk.WORKBOOK_PATH))
        self.assertIn(wb.properties.creator, (None, "openpyxl"))
        self.assertIsNone(wb.properties.lastModifiedBy)
        user = Path.home().name.lower()
        with zipfile.ZipFile(path) as parts:
            for name in parts.namelist():
                self.assertNotIn(user.encode(), parts.read(name).lower(), name)
        self.assertEqual(sha256(path), before)

    @unittest.skipUnless(BUILT, "the application has not been built")
    @unittest.skipIf(PAGE_CHANGED, REBUILD)
    def test_program_contents(self):
        try:
            from PyInstaller.archive.readers import CArchiveReader
        except ImportError:
            self.skipTest("PyInstaller is not installed")
        archive = CArchiveReader(str(EXE))
        entries = {name.replace("\\", "/").lower() for name in archive.toc}
        for name in PAGE_FILES:
            self.assertIn(f"frontend/{name}", entries)
        # The workbook is not packed into the program, in any form.
        self.assertEqual([e for e in entries if ".xls" in e or "fitness_tracker" in e], [])
        self.assertEqual([e for e in entries if e.endswith((".log", ".png", ".bat")) or "__pycache__" in e], [])

        pyz_name = next(name for name in archive.toc if name.lower().endswith(".pyz"))
        modules = archive.open_embedded_archive(pyz_name)
        names = set(modules.toc)
        for needed in ("backend.launch", "backend.app", "backend.store", "backend.analytics",
                       "backend.workbook", "backend.paths", "backend.excel_live", "flask", "openpyxl",
                       "werkzeug.serving", "win32timezone"):
            self.assertIn(needed, names)
        # The field rules and the analytics.
        for needed in ("backend.validation", "backend.analytics"):
            self.assertIn(needed, names)
        # The by-hand converter is not part of the program, and neither is anything that copies workbooks.
        self.assertNotIn("backend.migrate", names)
        # No test tooling or browser automation went in with it.
        self.assertEqual([n for n in names if n.split(".")[0] in ("pytest", "playwright", "PyInstaller")], [])
        self.assertEqual([n for n in names if n.split(".")[0] in ("tests", "windows_build")], [])
        self.assertNotIn("backend.create_workbook", names)
        self.assertEqual([n for n in names if n.startswith("win32com")], [])     # not needed, not packed
        self.assertEqual([n for n in names if n.startswith(("test_", "tkinter"))], [])

        # Nothing inside mentions where it was built.
        user = Path.home().name.lower()
        cloud = ("One" + "Drive").lower()
        where = set()

        def walk(code):
            where.add(code.co_filename.lower())
            for const in code.co_consts:
                if isinstance(const, types.CodeType):
                    walk(const)

        for name in names:
            code = modules.extract(name)
            if isinstance(code, types.CodeType):
                walk(code)
        self.assertEqual([w for w in where if user in w or cloud in w or ":\\" in w], [])
        # pywin32's own libraries, needed to talk to Excel, are inside too.
        self.assertTrue(any("pythoncom" in e for e in entries))
        self.assertTrue(any("pywintypes" in e for e in entries))
        raw = EXE.read_bytes().lower()
        for text in (user, str(PROJECT).lower()):
            self.assertNotIn(text.encode(), raw)
            self.assertNotIn(text.encode("utf-16-le"), raw)

    def test_build_recipe_keeps_the_workbook_out(self):
        spec = (PROJECT / "windows_build" / "fitness_tracker.spec").read_text(encoding="utf-8")
        build = (PROJECT / "windows_build" / "build.py").read_text(encoding="utf-8")
        recipe = "\n".join(line for line in spec.splitlines() if not line.lstrip().startswith("#"))
        self.assertNotIn("xlsx", recipe)
        self.assertIn('"frontend"', recipe)
        self.assertIn("console=True", recipe)
        # The build only ever adds a workbook where there is none.
        self.assertIn("if workbook.exists():", build)
        self.assertNotIn("rmtree(RELEASE", build)
        self.assertNotIn("unlink", build)
        # What it adds then is the clean template, never the working workbook with personal data.
        self.assertIn("shutil.copyfile(ROOT / TEMPLATE_NAME, workbook)", build)
        self.assertNotIn("ROOT / WORKBOOK_NAME", build)
        self.assertEqual(build.count("shutil.copyfile("), 2)       # the program, and the template


@unittest.skipUnless(BUILT, "the application has not been built")
class PackagedCase(unittest.TestCase):
    """Runs the real Fitness Tracker.exe from temporary folders."""

    def setUp(self):
        self.base = Path(tempfile.mkdtemp(prefix="ft-packaged-"))
        self.addCleanup(shutil.rmtree, self.base, ignore_errors=True)
        self.addCleanup(self.clear_unpacked_copies)   # runs after every stop
        self.started = []
        self.port = free_port()
        self.opened = self.base / "opened.txt"
        recorder = self.base / "browser.cmd"
        recorder.write_bytes(b'@echo %~1>>"%~dp0opened.txt"\r\n')
        self.env = {k: v for k, v in os.environ.items() if not k.startswith("FITNESS_TRACKER")}
        self.env.update(FITNESS_TRACKER_PORT=str(self.port), BROWSER=str(recorder))
        self.url = f"http://127.0.0.1:{self.port}/"
        self.hashes = {path: sha256(path) for path in (wbk.WORKBOOK_PATH, RELEASE / "Fitness_Tracker.xlsx")}
        self.addCleanup(self.assert_real_workbooks_untouched)
        self.folder = self.install("Fitness Tracker")

    def assert_real_workbooks_untouched(self):
        for path, digest in self.hashes.items():
            self.assertEqual(sha256(path), digest, path)

    def install(self, name, workbook=True):
        """Copy the program into a new folder, with its own fresh workbook."""
        folder = self.base / name
        folder.mkdir(parents=True)
        shutil.copy(EXE, folder)
        if workbook:
            wbk.create_workbook(folder / "Fitness_Tracker.xlsx")
        return folder

    def start(self, folder=None, args=(), env=None):
        folder = folder or self.folder
        log = self.base / f"window-{len(self.started) + 1}.txt"
        handle = open(log, "w")
        self.addCleanup(handle.close)
        process = subprocess.Popen(
            [str(folder / "Fitness Tracker.exe"), *args], env=env or self.env, cwd=TEMP,
            stdin=subprocess.DEVNULL, stdout=handle, stderr=subprocess.STDOUT,
            creationflags=subprocess.CREATE_NEW_PROCESS_GROUP,
        )
        self.started.append(process)
        self.addCleanup(self.stop, process)
        return process, log

    def stop(self, process):
        """Stop with Ctrl+Break, as pressing it in the window would."""
        if process.poll() is not None:
            return process.returncode
        try:
            os.kill(process.pid, signal.CTRL_BREAK_EVENT)
            return process.wait(25)
        except (OSError, subprocess.TimeoutExpired):
            subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"], capture_output=True)
            process.wait(15)
            return None

    def unpacked_copies(self):
        """Temporary folders the started programs unpacked themselves into."""
        return [p for process in self.started for p in TEMP.glob(f"_MEI{process.pid:08x}*")]

    def clear_unpacked_copies(self):
        for folder in self.unpacked_copies():
            shutil.rmtree(folder, ignore_errors=True)

    def opened_urls(self):
        return self.opened.read_text().split() if self.opened.exists() else []

    def ready(self, log):
        self.assertTrue(wait_for(self.opened_urls, 90), log.read_text())

    def assert_local_only(self, folder=None):
        """The running program listens on this computer only and is connected to nothing outside it."""
        pids = running_copies(folder or self.folder)
        self.assertTrue(pids)
        rows = sockets_of(pids)
        self.assertEqual([local for local, _, state in rows if state == "LISTENING"], [f"127.0.0.1:{self.port}"])
        for local, remote, state in rows:
            self.assertTrue(local.startswith("127.0.0.1:"), (local, remote, state))
            if state != "LISTENING":
                self.assertTrue(remote.startswith("127.0.0.1:"), (local, remote, state))


@unittest.skipIf(PAGE_CHANGED, REBUILD)
class PackagedRunTests(PackagedCase):
    """The program started from its own folder, with a workbook beside it."""

    # ---- 1, 2, 3: starts, opens the browser once ready, finds its workbook

    def test_starts_and_opens_the_browser_once_ready(self):
        process, log = self.start()
        self.ready(log)
        self.assertEqual(self.opened_urls(), [self.url])
        self.assertIn("workbook_modified", json.loads(get(f"{self.url}api/status")))
        output = log.read_text()
        self.assertIn(f"Address:  {self.url}  (this computer only)", output)
        self.assertLess(output.index("Address:"), output.index("Ready."))
        self.assertIsNone(process.poll())

    def test_uses_the_workbook_beside_the_program(self):
        _, log = self.start()
        self.ready(log)
        workbook = self.folder / "Fitness_Tracker.xlsx"
        self.assertIn(f"Workbook: {workbook}", log.read_text())
        self.assertEqual(json.loads(get(f"{self.url}api/months")),
                         [{"year": 2026, "month": 10, "label": "October 2026"}])
        status, body = post_month(self.url, 2026, 11)
        self.assertEqual((status, body["label"], len(body["days"])), (201, "November 2026", 30))
        self.assertEqual(load_workbook(workbook).sheetnames, ["October 2026", "November 2026"])
        # Only that one file changed; nothing new appeared beside it.
        self.assertEqual(sorted(p.name for p in self.folder.iterdir()),
                         ["Fitness Tracker.exe", "Fitness_Tracker.xlsx"])

    @unittest.skipIf(PAGE_CHANGED, REBUILD)
    def test_serves_the_same_page_files_as_the_source(self):
        _, log = self.start()
        self.ready(log)
        for name in PAGE_FILES:
            served = get(self.url + ("" if name == "index.html" else name))
            self.assertEqual(served, (PROJECT / "frontend" / name).read_bytes(), name)

    def test_data_and_analytics_are_served(self):
        workbook = self.folder / "Fitness_Tracker.xlsx"
        wb = load_workbook(workbook)
        ws = wb["October 2026"]
        ws["B2"], ws["D2"], ws["G2"], ws["K2"] = "Completed", "25:30", 2450, 72.5
        wb.save(workbook)
        _, log = self.start()
        self.ready(log)
        day = json.loads(get(f"{self.url}api/months/2026/10"))["days"][0]
        self.assertEqual((day["exercise"], day["cardio_seconds"], day["weight_lifted_kg"]),
                         ("Completed", 1530, 2450))
        analytics = json.loads(get(f"{self.url}api/months/2026/10/analytics"))
        self.assertEqual(analytics["cardio"]["total_display"], "25:30")
        self.assertEqual(analytics["weight_lifted"]["total_kg"], 2450)
        self.assertEqual(analytics["measurements"]["weight_kg"]["current"], 72.5)
        self.assertEqual(len(analytics["weekly"]), 5)

    # ---- 4, 5: moved folders, awkward names

    def test_runs_from_moved_folders_each_with_its_own_workbook(self):
        first = self.install("Fitness Tracker Test")
        second = self.install("Moved Tracker (copy 2) & more")
        store = st.TrackerStore(second / "Fitness_Tracker.xlsx")
        store.create_month(2027, 3)

        _, log = self.start(first)
        self.ready(log)
        self.assertIn(f"Workbook: {first / 'Fitness_Tracker.xlsx'}", log.read_text())
        self.assertEqual([m["label"] for m in json.loads(get(f"{self.url}api/months"))],
                         ["October 2026"])
        self.assertEqual(self.stop(self.started[-1]), 0)
        self.assertTrue(wait_for(lambda: not port_open(self.port), 15))

        self.opened.unlink()
        _, log = self.start(second)
        self.ready(log)
        output = log.read_text()
        self.assertIn(f"Workbook: {second / 'Fitness_Tracker.xlsx'}", output)
        self.assertEqual([m["label"] for m in json.loads(get(f"{self.url}api/months"))],
                         ["October 2026", "March 2027"])
        self.assertIn(b"<title>Raft</title>", get(self.url))
        # No trace of where the program was built.
        self.assertNotIn(str(PROJECT), output)
        self.assertNotIn("One" + "Drive", output)

    # ---- 8, 9: month creation and locking

    def test_adding_a_month_is_refused_while_the_workbook_is_locked(self):
        _, log = self.start()
        self.ready(log)
        workbook = self.folder / "Fitness_Tracker.xlsx"
        before = sha256(workbook)
        # Hold the file the way Excel does: open, with no sharing for writers.
        import ctypes
        from ctypes import wintypes
        create = ctypes.windll.kernel32.CreateFileW
        create.restype = wintypes.HANDLE
        create.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                           wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        # Read access, other programs may read but not write, existing file only.
        handle = create(str(workbook), 0x80000000, 0x00000001, None, 3, 0x80, None)
        self.assertNotIn(handle, (None, wintypes.HANDLE(-1).value))
        try:
            self.assertIs(json.loads(get(f"{self.url}api/status"))["locked"], True)
            self.assertEqual(len(json.loads(get(f"{self.url}api/months/2026/10"))["days"]), 31)
            status, body = post_month(self.url, 2026, 11)
            self.assertEqual((status, body["error"]["code"]), (423, "workbook_locked"))
            self.assertIn("Nothing was changed", body["error"]["message"])
        finally:
            ctypes.windll.kernel32.CloseHandle(wintypes.HANDLE(handle))
        self.assertEqual(sha256(workbook), before)
        self.assertEqual([p.name for p in self.folder.iterdir() if p.name.startswith(".tracker-")], [])
        self.assertEqual(post_month(self.url, 2026, 11)[0], 201)
        self.assertEqual(post_month(self.url, 2026, 11)[0], 409)

    # ---- 10: readable errors, never a silent exit

    def test_missing_workbook_gives_a_readable_error(self):
        folder = self.install("No Workbook Here", workbook=False)
        process, log = self.start(folder)
        self.assertEqual(process.wait(60), 1)
        output = log.read_text()
        self.assertIn("The workbook was not found:", output)
        self.assertIn(str(folder / "Fitness_Tracker.xlsx"), output)
        self.assertIn("never makes or replaces a workbook", output)
        self.assertIn("Fitness Tracker did not start.", output)
        self.assertEqual([p.name for p in folder.iterdir()], ["Fitness Tracker.exe"])
        self.assertFalse(port_open(self.port))
        self.assertEqual(self.opened_urls(), [])

    def test_damaged_workbook_gives_a_readable_error(self):
        workbook = self.folder / "Fitness_Tracker.xlsx"
        workbook.write_bytes(b"not a real workbook")
        process, log = self.start()
        self.assertEqual(process.wait(60), 1)
        output = log.read_text()
        self.assertIn("The workbook could not be opened:", output)
        self.assertIn("Fitness Tracker did not start.", output)
        self.assertEqual(workbook.read_bytes(), b"not a real workbook")
        self.assertEqual(self.opened_urls(), [])

    def test_port_taken_by_another_program_gives_a_readable_error(self):
        other = socket.socket()
        other.bind(("127.0.0.1", self.port))
        other.listen()
        self.addCleanup(other.close)
        process, log = self.start()
        self.assertEqual(process.wait(60), 1)
        output = log.read_text()
        self.assertIn(f"Port {self.port} is already in use by another program", output)
        self.assertIn("Fitness Tracker did not start.", output)
        self.assertEqual(self.opened_urls(), [])

    def test_wrong_option_gives_a_readable_error(self):
        process, log = self.start(args=["--port", "12"])
        self.assertEqual(process.wait(60), 2)
        output = log.read_text()
        self.assertIn("The port must be between 1024 and 65535", output)
        self.assertIn("Fitness Tracker did not start.", output)

    def test_starting_it_twice_opens_the_one_already_running(self):
        first, log = self.start()
        self.ready(log)
        second, log = self.start()
        self.assertEqual(second.wait(60), 0)
        self.assertIn(f"The tracker is already running at {self.url}", log.read_text())
        self.assertEqual(self.opened_urls(), [self.url, self.url])
        self.assertIsNone(first.poll())

    # ---- 11: shutdown

    def test_stopping_leaves_nothing_running_and_nothing_behind(self):
        process, log = self.start()
        self.ready(log)
        self.assertTrue(running_copies(self.folder))
        self.assertTrue(self.unpacked_copies())
        self.assertEqual(self.stop(process), 0)
        self.assertIn("Stopped.", log.read_text())
        self.assertTrue(wait_for(lambda: not port_open(self.port), 15))
        self.assertTrue(wait_for(lambda: not running_copies(self.folder), 15))
        self.assertTrue(wait_for(lambda: not self.unpacked_copies(), 15), "temporary files left behind")
        self.assertEqual(sorted(p.name for p in self.folder.iterdir()),
                         ["Fitness Tracker.exe", "Fitness_Tracker.xlsx"])

    # ---- 12, 13: this computer only, no outside connections

    def test_listens_only_on_this_computer_and_calls_nothing_outside(self):
        _, log = self.start()
        self.ready(log)
        for path in ("api/months", "api/months/2026/10", "api/months/2026/10/analytics", "app.js"):
            get(self.url + path)
        pids = running_copies(self.folder)
        self.assertTrue(pids)
        rows = sockets_of(pids)
        listening = [local for local, _, state in rows if state == "LISTENING"]
        self.assertEqual(listening, [f"127.0.0.1:{self.port}"])
        for local, remote, state in rows:
            self.assertTrue(local.startswith("127.0.0.1:"), (local, remote, state))
            if state != "LISTENING":
                self.assertTrue(remote.startswith("127.0.0.1:"), (local, remote, state))
        try:
            lan = socket.gethostbyname(socket.gethostname())
        except OSError:
            lan = "127.0.0.1"
        if not lan.startswith("127."):
            self.assertFalse(port_open(self.port, lan), f"reachable at {lan}")


def past_months(count=2):
    """The last `count` calendar months, oldest first. They are wholly in the
    past whatever today is, so every day in them may be edited."""
    year, month = date.today().year, date.today().month
    found = []
    for _ in range(count):
        month -= 1
        if month == 0:
            year, month = year - 1, 12
        found.append((year, month))
    return found[::-1]


def send(method, url, body=None, headers=None):
    """A JSON request to the running program. Returns (status, body)."""
    data = None if body is None else json.dumps(body).encode()
    request = urllib.request.Request(url, data=data, method=method,
                                     headers={"Content-Type": "application/json", **(headers or {})})
    try:
        with urllib.request.urlopen(request, timeout=20) as response:
            return response.status, json.loads(response.read())
    except urllib.error.HTTPError as error:
        return error.code, json.loads(error.read())


def hold_like_excel(path):
    """Open the file the way Excel does: other programs may read it, none may write.
    Returns a function that lets go of it."""
    import ctypes
    from ctypes import wintypes
    create = ctypes.windll.kernel32.CreateFileW
    create.restype = wintypes.HANDLE
    create.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                       wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    handle = create(str(path), 0x80000000, 0x00000001, None, 3, 0x80, None)
    if handle in (None, wintypes.HANDLE(-1).value):
        raise OSError(f"could not hold {path}")
    return lambda: ctypes.windll.kernel32.CloseHandle(wintypes.HANDLE(handle))


def loaded_files(pids):
    """Every program file (the .exe and its libraries) the processes have loaded."""
    command = ("; ".join(f"(Get-Process -Id {pid} -ErrorAction SilentlyContinue).Modules | "
                         "ForEach-Object { $_.FileName }" for pid in pids))
    output = subprocess.run(["powershell", "-NoProfile", "-Command", command],
                            capture_output=True, text=True).stdout
    return [line.strip() for line in output.splitlines() if line.strip()]


@unittest.skipIf(PAGE_CHANGED, REBUILD)
class PackagedEditingTests(PackagedCase):
    """The program at work: the page, the analytics, and saving days and
    measurements into the workbook."""

    def setUp(self):
        super().setUp()
        self.before, self.month = past_months(2)
        self.workbook = self.folder / "Fitness_Tracker.xlsx"
        self.sheet = wbk.sheet_name(*self.month)
        self.api = f"{self.url}api/months/{self.month[0]}/{self.month[1]}"
        self.sunday = wbk.month_sundays(*self.month)[0].day

    def write_workbook(self, months, cells=None, folder=None):
        """A workbook with the given months, each (year, month)."""
        wb = Workbook()
        wb.remove(wb.active)
        for year, month in months:
            wbk.add_month_sheet(wb, year, month)
        for sheet, values in (cells or {}).items():
            for ref, value in values.items():
                wb[sheet][ref] = value
        path = (folder or self.folder) / "Fitness_Tracker.xlsx"
        wb.save(path)
        return path

    def cells(self, refs, sheet=None, folder=None):
        ws = load_workbook((folder or self.folder) / "Fitness_Tracker.xlsx")[sheet or self.sheet]
        return [ws[ref].value for ref in refs]

    def restart(self, process):
        self.assertEqual(self.stop(process), 0)
        self.assertTrue(wait_for(lambda: not port_open(self.port), 15))
        self.opened.unlink()
        process, log = self.start()
        self.ready(log)
        return process

    # ---- the page and the analytics are what is inside the program

    def test_final_page_and_analytics_are_in_the_program(self):
        self.write_workbook([self.month], {self.sheet: {
            "B2": "Completed", "C2": "None", "D2": "25:30", "E2": 2200, "F2": 140.5, "G2": 2450,
            "B3": "Partial", "C3": "Controlled", "B4": "Rest Day", "C4": "Had", "B5": "Missed", "E5": 0,
            "K2": 72.5,
        }})
        _, log = self.start()
        self.ready(log)
        month = json.loads(get(self.api))
        self.assertNotIn("version", month)
        self.assertEqual(month["days"][0], {
            "date": date(*self.month, 1).isoformat(), "exercise": "Completed", "junk_food": "None",
            "cardio_display": "25:30", "cardio_seconds": 1530,
            "calories_kcal": 2200, "protein_g": 140.5, "weight_lifted_kg": 2450,
        })
        analytics = json.loads(get(f"{self.api}/analytics"))
        days = len(month["days"])
        self.assertNotIn("version", analytics)
        self.assertEqual(analytics["exercise"]["counts"],
                         {"completed": 1, "partial": 1, "rest_day": 1, "missed": days - 3, "pending": 0})
        self.assertEqual(analytics["junk_food"]["counts"],
                         {"none": 1, "controlled": 1, "had": 1, "missed": days - 3, "pending": 0})
        self.assertEqual(analytics["junk_food"]["success_percentage"], 50.0)
        self.assertEqual((analytics["calories"]["total_kcal"], analytics["calories"]["days_recorded"],
                          analytics["calories"]["lowest_kcal"]), (2200, 2, 0))
        self.assertEqual(analytics["protein"]["total_g"], 140.5)
        self.assertEqual(analytics["cardio"]["total_display"], "25:30")
        self.assertEqual(analytics["weight_lifted"]["total_kg"], 2450)
        self.assertEqual(analytics["measurements"]["weight_kg"]["current"], 72.5)
        self.assertEqual(json.loads(get(f"{self.url}api/status"))["today"], date.today().isoformat())

        # The page it serves has today, the bar of links, both editors and its four scripts.
        page = get(self.url).decode("utf-8")
        for piece in ('id="today-strip"', 'id="month-bar"', 'id="edit-dialog"', 'id="measure-dialog"',
                      'class="stack-rows"', 'src="analytics.js"', 'src="editor.js"'):
            self.assertIn(piece, page)
        for gone in ("analytics_v2", "upgrade", "version"):
            self.assertNotIn(gone, page.lower(), gone)
        css = get(f"{self.url}styles.css").decode("utf-8")
        for piece in ("@media (prefers-reduced-motion: reduce)", "@keyframes rise", "@starting-style",
                      ".stack-rows tbody tr", ".month-bar {"):
            self.assertIn(piece, css)
        self.assert_local_only()

    @unittest.skipIf(PAGE_CHANGED, REBUILD)
    def test_the_page_in_the_program_is_the_final_one(self):
        """The page inside the program is the one in frontend/: written for a phone first."""
        self.write_workbook([self.month])
        _, log = self.start()
        self.ready(log)
        page = get(self.url).decode("utf-8")
        for piece in ('<meta name="color-scheme" content="dark">', 'href="#today-strip" class="today-link"',
                      'class="add-text"'):
            self.assertIn(piece, page)
        css = get(f"{self.url}styles.css").decode("utf-8")
        for piece in ("color-scheme: dark;", "@media (min-width: 360px)", "@media (min-width: 1000px)",
                      "#days-body tr.upcoming", ".chart.compact {"):
            self.assertIn(piece, css)
        self.assertNotIn("@media (max-width", css)
        for name in PAGE_FILES:
            served = get(self.url + ("" if name == "index.html" else name))
            self.assertEqual(served, (PROJECT / "frontend" / name).read_bytes(), name)
        self.assert_local_only()

    # ---- saving from the dashboard reaches the workbook, and is still there after a restart

    def test_days_and_measurements_are_saved_into_the_workbook_and_survive_a_restart(self):
        self.write_workbook([self.month, (2099, 1)])
        process, log = self.start()
        self.ready(log)

        status, body = send("PUT", f"{self.api}/days/2", {
            "exercise": "Completed", "junk_food": "None", "cardio": "25:30",
            "calories_kcal": 2200, "protein_g": 140.5, "weight_lifted_kg": 2450})
        self.assertEqual(status, 200, body)
        self.assertEqual(body["day"]["exercise"], "Completed")
        self.assertEqual(self.cells(["B3", "C3", "D3", "E3", "F3", "G3"]),
                         ["Completed", "None", "25:30", 2200, 140.5, 2450])
        self.assertEqual(body["workbook_modified"], json.loads(get(f"{self.url}api/status"))["workbook_modified"])

        row = 2                     # the first Sunday is the first row of the measurement table
        status, body = send("PUT", f"{self.api}/measurements/{self.sunday}", {"weight_kg": 70.5, "waist_cm": 80})
        self.assertEqual(status, 200, body)
        self.assertEqual(body["measurement"]["weight_kg"], 70.5)
        self.assertEqual(self.cells([f"K{row}", f"L{row}", f"M{row}"]), [70.5, 80, None])

        # A field left out is kept, and null clears one.
        self.assertEqual(send("PUT", f"{self.api}/days/2", {"cardio": None, "protein_g": 150})[0], 200)
        self.assertEqual(self.cells(["B3", "C3", "D3", "E3", "F3", "G3"]),
                         ["Completed", "None", None, 2200, 150, 2450])

        # Refused requests change nothing: bad values, a day to come, a day that is not a Sunday.
        saved = sha256(self.workbook)
        status, body = send("PUT", f"{self.api}/days/3", {"exercise": "Done", "calories_kcal": 1.5, "protein_g": 99})
        self.assertEqual((status, sorted(body["error"]["fields"])), (400, ["calories_kcal", "exercise"]))
        status, body = send("PUT", f"{self.url}api/months/2099/1/days/1", {"exercise": "Completed"})
        self.assertEqual((status, list(body["error"]["fields"])), (400, ["day"]))
        weekday = next(day for day in range(1, 8) if date(*self.month, day).weekday() != 6)
        status, body = send("PUT", f"{self.api}/measurements/{weekday}", {"weight_kg": 70})
        self.assertEqual(status, 400)
        self.assertIn("not a Sunday", body["error"]["fields"]["day"])
        self.assertEqual(send("PUT", f"{self.api}/days/3", {})[0], 400)
        self.assertEqual(sha256(self.workbook), saved)
        self.assertEqual(sorted(p.name for p in self.folder.iterdir()),
                         ["Fitness Tracker.exe", "Fitness_Tracker.xlsx"])     # no temporary file, no backup

        # The workbook is still a sound workbook that Excel's rules apply to.
        wb = load_workbook(self.workbook)
        ws = wb[self.sheet]
        self.assertIs(wbk.is_month_sheet(ws), True)
        self.assertEqual(len(ws.data_validations.dataValidation), 7)
        self.assertEqual(len(ws.tables), 2)

        # Close the program and start it again: what was saved is what it shows.
        self.restart(process)
        month = json.loads(get(self.api))
        self.assertEqual(month["days"][1], {
            "date": date(*self.month, 2).isoformat(), "exercise": "Completed", "junk_food": "None",
            "cardio_display": None, "cardio_seconds": None,
            "calories_kcal": 2200, "protein_g": 150, "weight_lifted_kg": 2450,
        })
        self.assertEqual((month["measurements"][0]["weight_kg"], month["measurements"][0]["waist_cm"]), (70.5, 80))
        self.assertEqual(month["issues"], [])
        analytics = json.loads(get(f"{self.api}/analytics"))
        self.assertEqual((analytics["calories"]["total_kcal"], analytics["protein"]["total_g"]), (2200, 150))
        self.assertEqual(analytics["measurements"]["weight_kg"]["current"], 70.5)

    def test_saving_is_refused_while_the_workbook_is_locked(self):
        self.write_workbook([self.month])
        _, log = self.start()
        self.ready(log)
        before = sha256(self.workbook)
        release = hold_like_excel(self.workbook)
        try:
            self.assertIs(json.loads(get(f"{self.url}api/status"))["locked"], True)
            status, body = send("PUT", f"{self.api}/days/1", {"exercise": "Completed"})
            self.assertEqual((status, body["error"]["code"]), (423, "workbook_locked"))
            self.assertIn("Nothing was changed", body["error"]["message"])
            status, body = send("PUT", f"{self.api}/measurements/{self.sunday}", {"weight_kg": 70})
            self.assertEqual((status, body["error"]["code"]), (423, "workbook_locked"))
            self.assertEqual(len(json.loads(get(self.api))["days"]), len(wbk.month_dates(*self.month)))   # it can still be read
        finally:
            release()
        self.assertEqual(sha256(self.workbook), before)
        self.assertEqual([p.name for p in self.folder.iterdir() if p.name.startswith(".tracker-")], [])
        # Once the file is let go, the same request goes through.
        self.assertEqual(send("PUT", f"{self.api}/days/1", {"exercise": "Completed"})[0], 200)
        self.assertEqual(self.cells(["B2"]), ["Completed"])

    def test_changes_are_only_taken_from_the_trackers_own_page(self):
        self.write_workbook([self.month])
        _, log = self.start()
        self.ready(log)
        before = sha256(self.workbook)
        for headers in ({"Origin": "http://evil.example"}, {"Origin": f"http://localhost:{self.port + 1}"},
                        {"Sec-Fetch-Site": "cross-site"}):
            for method, url, body in (("PUT", f"{self.api}/days/1", {"exercise": "Completed"}),
                                      ("PUT", f"{self.api}/measurements/{self.sunday}", {"weight_kg": 70}),
                                      ("POST", f"{self.url}api/months", {"year": 2098, "month": 1})):
                status, answer = send(method, url, body, headers)
                self.assertEqual((status, answer["error"]["code"]), (403, "forbidden_origin"), (headers, url))
        self.assertEqual(sha256(self.workbook), before)
        own = {"Origin": self.url.rstrip("/"), "Sec-Fetch-Site": "same-origin"}
        self.assertEqual(send("PUT", f"{self.api}/days/1", {"exercise": "Completed"}, own)[0], 200)

    # ---- a sheet that is not laid out the tracker's way is reported and left alone

    def test_a_month_sheet_in_another_layout_is_reported_and_never_changed(self):
        from tests.earlier_layout import add_month
        wb = Workbook()
        wb.remove(wb.active)
        add_month(wb, *self.before, earlier=True)
        wbk.add_month_sheet(wb, *self.month)
        wb[wbk.sheet_name(*self.before)]["B2"] = True
        wb.save(self.workbook)
        process, log = self.start()
        self.ready(log)
        other_api = f"{self.url}api/months/{self.before[0]}/{self.before[1]}"
        before = sha256(self.workbook)
        status, body = send("PUT", f"{other_api}/days/1", {"exercise": "Completed"})
        self.assertEqual((status, body["error"]["code"]), (500, "workbook_format"))
        for route in ("/api/upgrade", "/api/migrate"):
            self.assertIn(send("POST", f"{self.url.rstrip('/')}{route}", {"confirm": True})[0], (404, 405))
        self.assertEqual(sha256(self.workbook), before)
        # The month beside it works as usual.
        self.assertEqual(send("PUT", f"{self.api}/days/2", {"exercise": "Completed"})[0], 200)
        self.assertEqual(self.stop(process), 0)
        wb = load_workbook(self.workbook)
        self.assertEqual([wbk.is_month_sheet(ws) for ws in wb], [False, True])
        self.assertIs(wb[wbk.sheet_name(*self.before)]["B2"].value, True)
        self.assertEqual(sorted(p.name for p in self.folder.iterdir()),
                         ["Fitness Tracker.exe", "Fitness_Tracker.xlsx"])         # and no copy of the workbook was made

    # ---- the release stands on its own: no source tree, no Python, wherever the folder is put

    def test_runs_from_a_moved_folder_without_the_development_environment(self):
        folder = self.install("Moved somewhere else & renamed (copy)", workbook=False)
        self.write_workbook([self.month], folder=folder)
        shutil.copy(RELEASE / "How to use.txt", folder)
        keep = ("SystemRoot", "SystemDrive", "windir", "TEMP", "TMP", "USERPROFILE", "LOCALAPPDATA", "APPDATA",
                "ProgramData", "COMSPEC", "PATHEXT", "HOMEDRIVE", "HOMEPATH", "USERNAME", "COMPUTERNAME", "OS",
                "NUMBER_OF_PROCESSORS", "PROCESSOR_ARCHITECTURE")
        system = os.environ.get("SystemRoot", r"C:\Windows")
        bare = {name: os.environ[name] for name in keep if name in os.environ}
        bare.update(PATH=f"{system}\\System32;{system}", FITNESS_TRACKER_PORT=str(self.port),
                    BROWSER=self.env["BROWSER"])
        for name in bare:
            self.assertFalse(name.upper().startswith(("PYTHON", "VIRTUAL_ENV")), name)
        self.assertNotIn("python", bare["PATH"].lower())

        _, log = self.start(folder, env=bare)
        self.ready(log)
        output = log.read_text()
        self.assertIn(f"Workbook: {folder / 'Fitness_Tracker.xlsx'}", output)
        self.assertNotIn(str(PROJECT), output)
        self.assertEqual(len(json.loads(get(self.api))["days"]), len(wbk.month_dates(*self.month)))
        self.assertEqual(send("PUT", f"{self.api}/days/1", {"exercise": "Rest Day", "calories_kcal": 2000})[0], 200)
        self.assertEqual(self.cells(["B2", "E2"], folder=folder), ["Rest Day", 2000])
        self.assertIn(b'id="today-strip"', get(self.url))

        # Everything the running program has loaded comes from its own folder, the
        # folder it unpacked itself into, or Windows; nothing from the project or its Python.
        pids = running_copies(folder)
        self.assertTrue(pids)
        files = loaded_files(pids)
        self.assertTrue(any(name.lower().endswith("fitness tracker.exe") for name in files), files[:5])
        strays = [name for name in files
                  if name.lower().startswith(str(PROJECT).lower()) or "\\.venv\\" in name.lower()
                  or "site-packages" in name.lower()]
        self.assertEqual(strays, [])
        own = [name for name in files if "\\_mei" in name.lower()]
        self.assertTrue(any("python3" in Path(name).name.lower() for name in own), own[:5])
        # The page files it serves are the ones it carries, unpacked beside its own runtime.
        unpacked = self.unpacked_copies()
        self.assertTrue(unpacked)
        for name in PAGE_FILES:
            self.assertTrue(any((folder_ / "frontend" / name).is_file() for folder_ in unpacked), name)
        self.assert_local_only(folder)


if __name__ == "__main__":
    unittest.main()

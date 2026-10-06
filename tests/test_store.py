"""Checks for the backend data layer.

Anything that writes uses a temporary workbook. The real Fitness_Tracker.xlsx
is only ever read here.

Run from the project root:  python -m unittest discover -s tests -v
"""

import hashlib
import importlib
import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from backend import store as st
from backend import workbook as wbk


class TempWorkbookCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.path = wbk.create_workbook(self.tmp / "Fitness_Tracker.xlsx")
        self.store = st.TrackerStore(self.path)

    def set_cells(self, sheet, cells):
        wb = load_workbook(self.path)
        for ref, value in cells.items():
            wb[sheet][ref] = value
        wb.save(self.path)

    def file_hash(self):
        return hashlib.sha256(self.path.read_bytes()).hexdigest()


class RealWorkbookReadTests(unittest.TestCase):
    """Read-only checks that hold whatever the user has entered so far."""

    @classmethod
    def setUpClass(cls):
        cls.before = hashlib.sha256(wbk.WORKBOOK_PATH.read_bytes()).hexdigest()
        cls.store = st.TrackerStore()
        cls.october = cls.store.get_month(2026, 10)

    def test_october_is_listed(self):
        self.assertIn(
            {"year": 2026, "month": 10, "label": "October 2026"}, self.store.list_months()
        )

    def test_october_has_31_days(self):
        dates = [day["date"] for day in self.october["days"]]
        self.assertEqual(dates, [date(2026, 10, d).isoformat() for d in range(1, 32)])

    def test_october_has_4_sundays(self):
        sundays = [row["sunday"] for row in self.october["measurements"]]
        self.assertEqual(sundays, ["2026-10-04", "2026-10-11", "2026-10-18", "2026-10-25"])

    def test_field_types(self):
        for day in self.october["days"]:
            self.assertEqual(set(day), {
                "date", "exercise", "junk_food", "cardio_display", "cardio_seconds",
                "calories_kcal", "protein_g", "weight_lifted_kg",
            })
            self.assertIn(day["exercise"], (None, "Completed", "Partial", "Rest Day", "Missed"))
            self.assertIn(day["junk_food"], (None, "None", "Controlled", "Had"))
            self.assertIsInstance(day["cardio_seconds"], (int, type(None)))
            for key in ("calories_kcal", "protein_g", "weight_lifted_kg"):
                self.assertIsInstance(day[key], (int, float, type(None)))
        for row in self.october["measurements"]:
            self.assertEqual(list(row), st.MEASUREMENT_KEYS)

    def test_reading_does_not_change_the_file(self):
        self.store.list_months()
        self.store.get_month(2026, 10)
        self.store.status()
        after = hashlib.sha256(wbk.WORKBOOK_PATH.read_bytes()).hexdigest()
        self.assertEqual(after, self.before)


class ListMonthsTests(TempWorkbookCase):
    def test_months_come_from_sheets(self):
        self.assertEqual(
            self.store.list_months(), [{"year": 2026, "month": 10, "label": "October 2026"}]
        )

    def test_other_sheets_are_ignored(self):
        wb = load_workbook(self.path)
        for name in ["Notes", "October", "Oct 2026", "October 26", "october 2026 copy"]:
            wb.create_sheet(name)
        wb.save(self.path)
        self.assertEqual(len(self.store.list_months()), 1)

    def test_list_follows_the_workbook(self):
        wb = load_workbook(self.path)
        wbk.add_month_sheet(wb, 2027, 3)
        del wb["October 2026"]
        wb.save(self.path)
        self.assertEqual(
            self.store.list_months(), [{"year": 2027, "month": 3, "label": "March 2027"}]
        )


class ReadMonthTests(TempWorkbookCase):
    def test_empty_month(self):
        data = self.store.get_month(2026, 10)
        self.assertEqual((data["year"], data["month"], data["label"]), (2026, 10, "October 2026"))
        self.assertEqual(len(data["days"]), 31)
        self.assertEqual(len(data["measurements"]), 4)
        self.assertEqual(data["issues"], [])
        self.assertEqual(data["days"][0], {
            "date": "2026-10-01",
            "exercise": None,
            "junk_food": None,
            "cardio_display": None,
            "cardio_seconds": None,
            "calories_kcal": None,
            "protein_g": None,
            "weight_lifted_kg": None,
        })
        self.assertEqual(data["measurements"][0], {
            "sunday": "2026-10-04", "weight_kg": None, "waist_cm": None,
            "chest_cm": None, "bicep_cm": None, "thigh_cm": None, "forearm_cm": None,
        })

    def test_entered_values(self):
        self.set_cells("October 2026", {
            "B2": "Completed", "D2": "25:30", "E2": 2200, "F2": 140.5, "G2": 2450,
            "D3": "0:45", "G3": 1820.5,
            "D4": "60:00", "G4": 0, "E4": 0,
            "K2": 72.4, "P2": 28,
        })
        data = self.store.get_month(2026, 10)
        first, second, third = data["days"][:3]
        self.assertEqual(first["exercise"], "Completed")
        self.assertIsNone(first["junk_food"])
        self.assertEqual((first["cardio_display"], first["cardio_seconds"]), ("25:30", 1530))
        self.assertEqual((first["calories_kcal"], first["protein_g"], first["weight_lifted_kg"]), (2200, 140.5, 2450))
        self.assertEqual((second["cardio_display"], second["cardio_seconds"]), ("0:45", 45))
        self.assertEqual(second["weight_lifted_kg"], 1820.5)
        self.assertIsNone(second["calories_kcal"])              # blank is not a zero
        self.assertEqual((third["cardio_display"], third["cardio_seconds"]), ("60:00", 3600))
        self.assertEqual((third["weight_lifted_kg"], third["calories_kcal"]), (0, 0))      # a recorded zero is a value
        self.assertEqual(data["measurements"][0]["weight_kg"], 72.4)
        self.assertEqual(data["measurements"][0]["forearm_cm"], 28)
        self.assertIsNone(data["measurements"][0]["waist_cm"])
        self.assertEqual(data["issues"], [])

    def test_a_status_is_its_word_and_a_blank_is_not_entered(self):
        self.set_cells("October 2026", {"B2": "Completed", "B3": "Missed", "B4": None, "C2": "None", "C3": "Had"})
        days = self.store.get_month(2026, 10)["days"]
        self.assertEqual((days[0]["exercise"], days[0]["junk_food"]), ("Completed", "None"))
        self.assertEqual((days[1]["exercise"], days[1]["junk_food"]), ("Missed", "Had"))
        self.assertIsNone(days[2]["exercise"])      # not entered
        self.assertIsNone(days[2]["junk_food"])
        self.assertEqual(self.store.get_month(2026, 10)["issues"], [])

    def test_invalid_values_are_reported_not_converted(self):
        self.set_cells("October 2026", {
            "D2": 25.30, "D3": "25:75", "D4": "abc",
            "G5": "heavy", "G6": -10, "B7": "yes", "K2": 0,
        })
        data = self.store.get_month(2026, 10)
        for index in range(3):
            self.assertIsNone(data["days"][index]["cardio_seconds"])
        self.assertEqual(data["days"][0]["cardio_display"], "25.3")
        self.assertEqual(data["days"][1]["cardio_display"], "25:75")
        self.assertIsNone(data["days"][3]["weight_lifted_kg"])
        self.assertIsNone(data["days"][4]["weight_lifted_kg"])
        self.assertIsNone(data["days"][5]["exercise"])
        self.assertIsNone(data["measurements"][0]["weight_kg"])
        flagged = [(issue["date"], issue["field"]) for issue in data["issues"]]
        self.assertEqual(flagged, [
            ("2026-10-01", "cardio"), ("2026-10-02", "cardio"), ("2026-10-03", "cardio"),
            ("2026-10-04", "weight_lifted_kg"), ("2026-10-05", "weight_lifted_kg"),
            ("2026-10-06", "exercise"), ("2026-10-04", "weight_kg"),
        ])

    def test_reading_leaves_excel_values_untouched(self):
        self.set_cells("October 2026", {"D2": "25:30"})
        before = self.file_hash()
        self.store.get_month(2026, 10)
        self.assertEqual(self.file_hash(), before)
        self.assertEqual(load_workbook(self.path)["October 2026"]["D2"].value, "25:30")

    def test_response_has_no_cell_coordinates(self):
        data = self.store.get_month(2026, 10)
        self.assertEqual(
            set(data), {"year", "month", "label", "days", "measurements", "issues"}
        )

    def test_missing_month(self):
        with self.assertRaises(st.MonthNotFoundError):
            self.store.get_month(2026, 11)

    def test_invalid_month_arguments(self):
        for year, month in [(2026, 0), (2026, 13), (1999, 5), (2101, 5), ("2026", 10),
                            (2026, 10.0), (True, 1), (2026, None)]:
            with self.assertRaises(st.InvalidMonthError):
                self.store.get_month(year, month)


class CreateMonthTests(TempWorkbookCase):
    def test_november_can_be_created(self):
        data = self.store.create_month(2026, 11)
        self.assertEqual(data["label"], "November 2026")
        self.assertEqual(len(data["days"]), 30)
        self.assertEqual(
            [row["sunday"] for row in data["measurements"]],
            ["2026-11-01", "2026-11-08", "2026-11-15", "2026-11-22", "2026-11-29"],
        )
        self.assertEqual(self.store.get_month(2026, 11), data)
        self.assertEqual([m["label"] for m in self.store.list_months()],
                         ["October 2026", "November 2026"])

    def test_new_month_has_the_same_structure_and_leaves_the_others_alone(self):
        before = [[cell.value for cell in row] for row in load_workbook(self.path)["October 2026"].iter_rows()]
        self.store.create_month(2026, 11)
        wb = load_workbook(self.path)
        october, november = wb["October 2026"], wb["November 2026"]
        self.assertIs(wbk.is_month_sheet(november), True)
        for col in range(1, 17):
            self.assertEqual(november.cell(1, col).value, october.cell(1, col).value)
            self.assertEqual(november.cell(2, col).number_format, october.cell(2, col).number_format)
        self.assertEqual([november.cell(1, col).value for col in range(1, 8)], list(wbk.DAILY_HEADERS))
        self.assertEqual(november.tables["Daily_2026_11"].ref, "A1:G31")
        self.assertEqual(november.tables["Measurements_2026_11"].ref, "J1:P6")
        self.assertEqual(len(november.data_validations.dataValidation), 7)
        self.assertEqual(len(october.data_validations.dataValidation), 7)
        self.assertEqual([[cell.value for cell in row] for row in october.iter_rows()], before)
        text = " ".join(
            str(cell.value).lower() for row in november.iter_rows() for cell in row
            if isinstance(cell.value, str)
        )
        for word in ["goal", "meditation", "muscle", "bench", "squat"]:
            self.assertNotIn(word, text)

    def test_new_month_starts_empty(self):
        data = self.store.create_month(2026, 11)
        for day in data["days"]:
            for key in ("exercise", "junk_food", "cardio_seconds", "calories_kcal", "protein_g", "weight_lifted_kg"):
                self.assertIsNone(day[key], key)
        for row in data["measurements"]:
            self.assertTrue(all(row[key] is None for key in st.MEASUREMENT_KEYS[1:]))

    def test_duplicate_november_is_rejected(self):
        self.store.create_month(2026, 11)
        before = self.file_hash()
        with self.assertRaises(st.MonthExistsError):
            self.store.create_month(2026, 11)
        self.assertEqual(self.file_hash(), before)

    def test_month_lengths(self):
        for (year, month), days in {(2028, 2): 29, (2027, 2): 28, (2027, 4): 30}.items():
            data = self.store.create_month(year, month)
            self.assertEqual(len(data["days"]), days)
            self.assertEqual(data["days"][-1]["date"], date(year, month, days).isoformat())
            self.assertEqual(len(data["measurements"]), len(wbk.month_sundays(year, month)))

    def test_invalid_month_is_rejected_without_writing(self):
        before = self.file_hash()
        for year, month in [(2026, 13), (2026, 0), (1800, 1), ("2026", 11), (2026, 1.5)]:
            with self.assertRaises(st.InvalidMonthError):
                self.store.create_month(year, month)
        self.assertEqual(self.file_hash(), before)

    def test_existing_data_survives_creation(self):
        self.set_cells("October 2026", {"B2": "Completed", "D2": "25:30", "G2": 2450, "K2": 72.4})
        before = self.store.get_month(2026, 10)
        self.store.create_month(2026, 11)
        self.assertEqual(self.store.get_month(2026, 10), before)

    def test_sheets_are_kept_in_calendar_order(self):
        for year, month in [(2027, 1), (2026, 12), (2026, 9), (2026, 11)]:
            self.store.create_month(year, month)
        self.assertEqual(load_workbook(self.path).sheetnames, [
            "September 2026", "October 2026", "November 2026", "December 2026", "January 2027",
        ])


class MigrationRemovedTests(unittest.TestCase):
    """The one-time FALSE-to-blank conversion has been run and must stay gone:
    running it again would clear FALSE values the user entered on purpose."""

    def test_script_file_is_gone(self):
        backend = Path(st.__file__).resolve().parent
        self.assertFalse((backend / "clear_prefilled_habits.py").exists())
        self.assertEqual([p.name for p in backend.glob("*clear*")], [])

    def test_module_cannot_be_imported(self):
        with self.assertRaises(ModuleNotFoundError):
            importlib.import_module("backend.clear_prefilled_habits")

    def test_store_has_no_bulk_clearing_method(self):
        self.assertFalse(hasattr(st.TrackerStore, "clear_prefilled_habits"))
        writers = [name for name in vars(st.TrackerStore)
                   if not name.startswith("_")
                   and name not in ("status", "list_months", "get_month")]
        # The only things that write: editing one row, and adding a month.
        self.assertEqual(writers, ["update_day", "update_measurements", "create_month"])


class MonthIsolationTests(TempWorkbookCase):
    def test_each_month_returns_only_its_own_data(self):
        self.store.create_month(2026, 11)
        self.set_cells("October 2026", {"G2": 1000, "K2": 70})
        self.set_cells("November 2026", {"G2": 2000, "K2": 80})
        october = self.store.get_month(2026, 10)
        november = self.store.get_month(2026, 11)
        self.assertTrue(all(d["date"].startswith("2026-10-") for d in october["days"]))
        self.assertTrue(all(d["date"].startswith("2026-11-") for d in november["days"]))
        self.assertTrue(all(m["sunday"].startswith("2026-10-") for m in october["measurements"]))
        self.assertTrue(all(m["sunday"].startswith("2026-11-") for m in november["measurements"]))
        self.assertEqual(
            [d["weight_lifted_kg"] for d in october["days"] if d["weight_lifted_kg"]],
            [1000],
        )
        self.assertEqual(
            [d["weight_lifted_kg"] for d in november["days"] if d["weight_lifted_kg"]],
            [2000],
        )
        self.assertEqual(october["measurements"][0]["weight_kg"], 70)
        self.assertEqual(november["measurements"][0]["weight_kg"], 80)

    def test_foreign_date_in_a_sheet_is_an_error(self):
        self.set_cells("October 2026", {"A5": date(2026, 11, 4)})
        with self.assertRaises(st.WorkbookFormatError):
            self.store.get_month(2026, 10)

    def test_wrong_sunday_is_an_error(self):
        self.set_cells("October 2026", {"J2": date(2026, 10, 5)})
        with self.assertRaises(st.WorkbookFormatError):
            self.store.get_month(2026, 10)

    def test_changed_headers_are_an_error(self):
        self.set_cells("October 2026", {"C1": "Meditation"})
        with self.assertRaises(st.WorkbookFormatError):
            self.store.get_month(2026, 10)


class FileSafetyTests(TempWorkbookCase):
    def leftovers(self):
        return [p.name for p in self.tmp.iterdir() if p.name != self.path.name]

    def test_locked_on_save_leaves_workbook_untouched(self):
        before = self.file_hash()
        with mock.patch.object(st.os, "replace", side_effect=PermissionError("locked")):
            with self.assertRaises(st.WorkbookLockedError):
                self.store.create_month(2026, 11)
        self.assertEqual(self.file_hash(), before)
        self.assertEqual(self.leftovers(), [])

    def test_locked_before_write_is_reported(self):
        before = self.file_hash()
        with mock.patch("builtins.open", side_effect=PermissionError("locked")):
            with self.assertRaises(st.WorkbookLockedError):
                self.store.create_month(2026, 11)
        self.assertEqual(self.file_hash(), before)

    def test_failed_save_leaves_workbook_untouched(self):
        before = self.file_hash()
        with mock.patch("openpyxl.workbook.workbook.Workbook.save",
                        side_effect=RuntimeError("disk full")):
            with self.assertRaises(RuntimeError):
                self.store.create_month(2026, 11)
        self.assertEqual(self.file_hash(), before)
        self.assertEqual(self.leftovers(), [])

    def test_locked_on_read_is_reported(self):
        with mock.patch.object(st, "load_workbook", side_effect=PermissionError("locked")):
            with self.assertRaises(st.WorkbookLockedError):
                self.store.list_months()

    def test_missing_workbook(self):
        store = st.TrackerStore(self.tmp / "nope.xlsx")
        for call in (store.list_months, store.status, lambda: store.get_month(2026, 10),
                     lambda: store.create_month(2026, 11)):
            with self.assertRaises(st.WorkbookNotFoundError):
                call()
        self.assertFalse((self.tmp / "nope.xlsx").exists())

    def test_corrupt_workbook(self):
        self.path.write_bytes(b"this is not a workbook")
        with self.assertRaises(st.WorkbookUnreadableError):
            self.store.list_months()
        with self.assertRaises(st.WorkbookUnreadableError):
            self.store.create_month(2026, 11)
        self.assertEqual(self.path.read_bytes(), b"this is not a workbook")

    def test_status(self):
        status = self.store.status()
        self.assertIs(status["locked"], False)
        self.assertIsInstance(status["workbook_modified"], str)


if __name__ == "__main__":
    unittest.main()

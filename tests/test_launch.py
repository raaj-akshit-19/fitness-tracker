"""Checks for the Windows launcher and the startup code behind it.

Every server started here uses a temporary workbook and a spare port on
127.0.0.1. No browser window is opened: the BROWSER environment variable
points Python at a small script that only records the address it was given.
The real workbook is never used.

Run from the project root:  python -m unittest discover -s tests -v
"""

import hashlib
import json
import logging
import os
import re
import shutil
import socket
import subprocess
import tempfile
import threading
import time
import unittest
import urllib.request
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook
from werkzeug.serving import make_server

from backend import launch
from backend import workbook as wbk
from backend.app import create_app

PROJECT = Path(__file__).resolve().parent.parent
BAT = PROJECT / "Start Fitness Tracker.bat"
WINDOWS = os.name == "nt"
SOURCE_DIRS = ["backend", "frontend", "tests", "windows_build"]
SOURCE_FILES = ["Start Fitness Tracker.bat", "README.md", "requirements.txt",
                "requirements-build.txt", ".gitignore"]
RELEASE = PROJECT / "release" / "Fitness Tracker"


def setUpModule():
    # Servers started inside these tests would otherwise print every request.
    logging.getLogger("werkzeug").setLevel(logging.ERROR)


def found(pattern, text):
    """True when the pattern occurs. Keeps failure messages short."""
    return re.search(pattern, text) is not None


def free_port():
    with socket.socket() as probe:
        probe.bind((launch.HOST, 0))
        return probe.getsockname()[1]


def port_open(port, host=launch.HOST):
    try:
        with socket.create_connection((host, port), timeout=1):
            return True
    except OSError:
        return False


def get_json(url):
    with urllib.request.urlopen(url, timeout=5) as response:
        return json.loads(response.read())


def wait_for(condition, timeout=40):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if condition():
            return True
        time.sleep(0.2)
    return False


def source_files():
    for name in SOURCE_FILES:
        yield PROJECT / name
    if (RELEASE / "How to use.txt").exists():
        yield RELEASE / "How to use.txt"
    for folder in SOURCE_DIRS:
        for path in (PROJECT / folder).rglob("*"):
            if path.is_file() and "__pycache__" not in path.parts:
                yield path


class TempCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.workbook = wbk.create_workbook(self.tmp / "Fitness_Tracker.xlsx")

    def serve(self, port):
        """Run the tracker in this process on the given port until the test ends."""
        server = make_server(launch.HOST, port, create_app(self.workbook), threaded=True)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        self.addCleanup(thread.join, 5)
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        return server

    def listen(self, port):
        """Occupy a port with something that is not the tracker."""
        other = socket.socket()
        other.bind((launch.HOST, port))
        other.listen()
        self.addCleanup(other.close)
        return other


class StartupCheckTests(TempCase):
    def test_missing_workbook_is_reported_and_not_created(self):
        missing = self.tmp / "nowhere" / "Fitness_Tracker.xlsx"
        with self.assertRaises(launch.LaunchError) as caught:
            launch.check_workbook(missing)
        self.assertIn("was not found", str(caught.exception))
        self.assertIn("Nothing was created", str(caught.exception))
        self.assertFalse(missing.exists())
        self.assertFalse(missing.parent.exists())

    def test_existing_workbook_passes_and_is_untouched(self):
        before = self.workbook.read_bytes()
        launch.check_workbook(self.workbook)
        self.assertEqual(self.workbook.read_bytes(), before)

    def test_dependencies_present(self):
        launch.check_dependencies()

    def test_missing_dependency_names_the_package_and_the_fix(self):
        real = launch.importlib.util.find_spec
        fake = lambda name: None if name == "flask" else real(name)
        with mock.patch.object(launch.importlib.util, "find_spec", side_effect=fake):
            with self.assertRaises(launch.LaunchError) as caught:
                launch.check_dependencies()
        message = str(caught.exception)
        self.assertIn("Flask", message)
        self.assertNotIn("openpyxl,", message)
        self.assertIn("-m pip install -r", message)
        self.assertIn("requirements.txt", message)

    def test_only_listens_on_this_computer(self):
        self.assertEqual(launch.HOST, "127.0.0.1")
        for name in ("launch.py", "app.py"):
            text = (PROJECT / "backend" / name).read_text(encoding="utf-8")
            self.assertNotIn("0.0.0.0", text)
            self.assertFalse(found(r"host\s*=\s*[\"'](?!127\.0\.0\.1)", text), name)

    def test_launcher_never_writes_the_workbook(self):
        text = (PROJECT / "backend" / "launch.py").read_text(encoding="utf-8")
        for word in ("create_workbook", "create_month", ".save(", "write_bytes", "open("):
            self.assertNotIn(word, text.replace("urlopen(", "").replace("_open(", ""))


class PortTests(TempCase):
    def test_free_port(self):
        port = free_port()
        self.assertFalse(launch.port_is_busy(port))
        self.assertEqual(launch.choose_port(port, explicit=False), (port, False, None))
        self.assertEqual(launch.choose_port(port, explicit=True), (port, False, None))

    def test_busy_port_is_detected(self):
        port = free_port()
        self.listen(port)
        self.assertTrue(launch.port_is_busy(port))
        self.assertFalse(launch.tracker_is_running(port))

    def test_default_port_busy_moves_to_the_next_free_one_and_says_so(self):
        busy, also_busy, spare = free_port(), free_port(), free_port()
        self.listen(busy)
        self.listen(also_busy)
        port, running, note = launch.choose_port(busy, explicit=False, fallbacks=[also_busy, spare])
        self.assertEqual((port, running), (spare, False))
        self.assertIn(f"Port {busy} is in use by another program", note)
        self.assertIn(f"port {spare} instead", note)

    def test_explicit_port_busy_is_an_error(self):
        port = free_port()
        self.listen(port)
        with self.assertRaises(launch.LaunchError) as caught:
            launch.choose_port(port, explicit=True, fallbacks=[free_port()])
        self.assertIn(f"Port {port} is already in use", str(caught.exception))

    def test_no_port_available_is_an_error(self):
        first, second = free_port(), free_port()
        self.listen(first)
        self.listen(second)
        with self.assertRaises(launch.LaunchError) as caught:
            launch.choose_port(first, explicit=False, fallbacks=[second])
        self.assertIn("are all in use", str(caught.exception))

    def test_tracker_already_running_is_recognised(self):
        port = free_port()
        self.serve(port)
        self.assertTrue(launch.tracker_is_running(port))
        self.assertEqual(launch.choose_port(port, explicit=False), (port, True, None))
        self.assertEqual(launch.choose_port(port, explicit=True), (port, True, None))

    def test_tracker_running_on_a_fallback_port_is_reused(self):
        busy, tracker = free_port(), free_port()
        self.listen(busy)
        self.serve(tracker)
        self.assertEqual(launch.choose_port(busy, explicit=False, fallbacks=[tracker]),
                         (tracker, True, None))

    def test_port_arguments(self):
        self.assertEqual(launch._parse(["--port", "5050"]).port, 5050)
        self.assertIsNone(launch._parse([]).port)
        with mock.patch.dict(os.environ, {"FITNESS_TRACKER_PORT": "5051"}):
            self.assertEqual(launch._parse([]).port, 5051)
            self.assertEqual(launch._parse(["--port", "5052"]).port, 5052)
        with mock.patch("sys.stderr"):
            for bad in (["--port", "80"], ["--port", "70000"], ["--host", "0.0.0.0"]):
                with self.assertRaises(SystemExit):
                    launch._parse(bad)


class ReadinessTests(TempCase):
    def test_ready_once_the_status_endpoint_answers(self):
        port = free_port()
        self.assertFalse(launch.wait_until_ready(port, timeout=1))
        self.serve(port)
        self.assertTrue(launch.wait_until_ready(port, timeout=10))

    def test_waits_for_a_server_that_starts_late(self):
        port = free_port()
        timer = threading.Timer(1.0, self.serve, args=[port])
        timer.start()
        self.addCleanup(timer.cancel)
        started = time.monotonic()
        self.assertTrue(launch.wait_until_ready(port, timeout=15))
        self.assertGreaterEqual(time.monotonic() - started, 0.9)

    def test_something_else_on_the_port_is_not_ready(self):
        port = free_port()
        self.listen(port)
        self.assertFalse(launch.wait_until_ready(port, timeout=1.5))

    def test_gives_up_when_told_to(self):
        started = time.monotonic()
        self.assertFalse(launch.wait_until_ready(free_port(), timeout=30, keep_waiting=lambda: False))
        self.assertLess(time.monotonic() - started, 5)


class NoInternetTests(TempCase):
    """Nothing the backend or the launcher does may leave this computer."""

    def test_backend_and_launcher_only_connect_to_this_computer(self):
        real_connect = socket.socket.connect
        seen = []

        def guarded(sock, address):
            seen.append(address[0])
            if address[0] not in ("127.0.0.1", "::1", "localhost"):
                raise AssertionError(f"tried to reach {address}")
            return real_connect(sock, address)

        port = free_port()
        with mock.patch.object(socket.socket, "connect", guarded):
            self.serve(port)
            self.assertTrue(launch.wait_until_ready(port, timeout=10))
            self.assertTrue(launch.tracker_is_running(port))
            base = f"http://{launch.HOST}:{port}"
            self.assertEqual(len(get_json(f"{base}/api/months")), 1)
            self.assertEqual(len(get_json(f"{base}/api/months/2026/10")["days"]), 31)
            self.assertIn("exercise", get_json(f"{base}/api/months/2026/10/analytics"))
            request = urllib.request.Request(
                f"{base}/api/months", data=b'{"year": 2026, "month": 11}',
                headers={"Content-Type": "application/json"}, method="POST",
            )
            with urllib.request.urlopen(request, timeout=10) as response:
                self.assertEqual(response.status, 201)
            for name in ("", "styles.css", "format.js", "analytics.js", "app.js"):
                with urllib.request.urlopen(f"{base}/{name}", timeout=5) as response:
                    self.assertEqual(response.status, 200)
                    self.assertTrue(response.read())
        self.assertTrue(seen)
        self.assertEqual(set(seen), {"127.0.0.1"})

    def test_requirements_are_only_local_libraries(self):
        lines = (PROJECT / "requirements.txt").read_text(encoding="utf-8").splitlines()
        names = sorted(line.split("==")[0].strip().lower() for line in lines if line.strip())
        # Flask serves the page, openpyxl reads the file, pywin32 talks to Excel on this computer.
        self.assertEqual(names, ["flask", "openpyxl", "pywin32"])

    def test_backend_imports_no_network_client(self):
        for path in (PROJECT / "backend").glob("*.py"):
            text = path.read_text(encoding="utf-8")
            for module in ("requests", "httpx", "aiohttp", "smtplib", "ftplib", "boto", "firebase"):
                self.assertFalse(found(rf"(?m)^\s*(import|from)\s+{module}", text), path.name)
        # The launcher's only web request goes to its own status endpoint.
        text = (PROJECT / "backend" / "launch.py").read_text(encoding="utf-8")
        self.assertEqual(re.findall(r"https?://[^\s\"'{]*", text), ["http://", "http://"])


class ProjectFileTests(unittest.TestCase):
    def test_launcher_file(self):
        data = BAT.read_bytes()
        text = data.decode("ascii")
        # Windows line endings throughout: batch labels misbehave without them.
        self.assertEqual(data.count(b"\n"), data.count(b"\r\n"))
        self.assertIn('cd /d "%~dp0"', text)
        self.assertIn(r"%~dp0.venv\Scripts\python.exe", text)
        self.assertIn("-m backend.launch", text)
        self.assertFalse(found(r"[A-Za-z]:\\", text), "a drive letter path")
        self.assertNotIn("0.0.0.0", text)
        self.assertNotIn("http", text.lower())
        for word in ("create_workbook", "pip install -", "curl", "powershell", "del ", "rmdir"):
            self.assertNotIn(word, text.lower().replace("-m pip install -r requirements.txt", ""))

    def test_no_machine_specific_paths_in_the_project(self):
        # Built in pieces so this file does not contain what it looks for.
        needles = {
            "the user name": re.escape(Path.home().name),
            "a cloud folder path": r"[\\/]" + "One" + "Drive" + r"[\\/]",
            "a Windows profile path": r"[A-Za-z]:\\" + "Users" + r"\\",
            "a profile path": "/" + "Users" + "/|/" + "home" + r"/\w+/",
        }
        for path in source_files():
            text = path.read_text(encoding="utf-8", errors="replace")
            name = path.relative_to(PROJECT).as_posix()
            for what, pattern in needles.items():
                self.assertFalse(found(pattern, text), f"{name} contains {what}")

    def test_paths_come_from_the_file_locations(self):
        self.assertEqual(launch.PROJECT_DIR, PROJECT)
        self.assertEqual(Path(wbk.__file__).resolve().parent.parent, PROJECT)
        default = Path(wbk.__file__).resolve().parent.parent / "Fitness_Tracker.xlsx"
        if not os.environ.get("FITNESS_TRACKER_WORKBOOK"):
            self.assertEqual(wbk.WORKBOOK_PATH, default)

    def test_no_leftover_files(self):
        allowed_root = {
            ".gitignore", ".gitattributes", ".venv", "backend", "frontend", "tests", "Fitness_Tracker.xlsx",
            "README.md", "requirements.txt", "Start Fitness Tracker.bat",
            "windows_build", "requirements-build.txt", "Fitness_Tracker_Template.xlsx",
        }
        # The built application, once made, lives in release/. personal/ is a
        # private folder that is never committed.
        self.assertEqual({p.name for p in PROJECT.iterdir()} - {".git", "release", "personal"},
                         allowed_root)
        workbooks = {p for p in PROJECT.rglob("*.xls*")
                     if ".venv" not in p.parts and "personal" not in p.parts
                     and not p.name.startswith("~$")}
        self.assertEqual(workbooks - {RELEASE / "Fitness_Tracker.xlsx"},
                         {PROJECT / "Fitness_Tracker.xlsx", PROJECT / "Fitness_Tracker_Template.xlsx"})
        for leftover in ("build", "dist"):
            self.assertFalse((PROJECT / leftover).exists(), f"{leftover} folder left in the project")
        self.assertEqual(sorted(p.name for p in (PROJECT / "windows_build").iterdir()
                                if p.name != "__pycache__"),
                         ["app_entry.py", "build.py", "fitness_tracker.spec"])
        for path in source_files():
            self.assertFalse(found(r"\.(log|tmp|bak|png|exe)$|^~\$|^\.tracker-", path.name),
                             path.name)

    def test_no_credentials(self):
        pattern = re.compile(r"(api[_-]?key|secret|password|token)\s*[=:]\s*[\"'][^\"']+[\"']", re.I)
        for path in source_files():
            text = path.read_text(encoding="utf-8", errors="replace")
            self.assertIsNone(pattern.search(text), path.name)


@unittest.skipUnless(WINDOWS, "the launcher is a Windows batch file")
class LauncherRunTests(unittest.TestCase):
    """Runs the real .bat file."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.workbook = wbk.create_workbook(self.tmp / "Fitness_Tracker.xlsx")
        self.port = free_port()
        self.opened = self.tmp / "opened.txt"
        # Stands in for the browser: writes down the address it is asked to open.
        recorder = self.tmp / "browser.cmd"
        recorder.write_bytes(b'@echo %~1>>"%~dp0opened.txt"\r\n')
        self.env = dict(os.environ)
        self.env.update({
            "FITNESS_TRACKER_WORKBOOK": str(self.workbook),
            "FITNESS_TRACKER_PORT": str(self.port),
            "BROWSER": str(recorder),
        })
        self.real_hash = hashlib.sha256(wbk.WORKBOOK_PATH.read_bytes()).hexdigest()
        self.addCleanup(self.assert_real_workbook_untouched)
        self.runs = 0

    def assert_real_workbook_untouched(self):
        self.assertEqual(hashlib.sha256(wbk.WORKBOOK_PATH.read_bytes()).hexdigest(), self.real_hash)

    def start(self, bat=BAT, args="", env=None):
        """Start the .bat the way a double-click does. Returns (process, log file)."""
        self.runs += 1
        log = self.tmp / f"launcher-{self.runs}.log"
        handle = open(log, "w")
        self.addCleanup(handle.close)
        process = subprocess.Popen(
            f'cmd /s /c ""{bat}" {args}"', env=env or self.env, cwd=self.tmp,
            stdin=subprocess.DEVNULL, stdout=handle, stderr=subprocess.STDOUT,
            creationflags=subprocess.CREATE_NEW_PROCESS_GROUP,
        )
        self.addCleanup(self.stop, process)
        return process, log

    def stop(self, process):
        """What closing the launcher's window does: end it and everything it started."""
        if process.poll() is None:
            subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"], capture_output=True)
            process.wait(15)

    def opened_urls(self):
        return self.opened.read_text().split() if self.opened.exists() else []

    def test_starts_waits_until_ready_and_opens_the_real_address(self):
        process, log = self.start()
        self.assertTrue(wait_for(self.opened_urls), log.read_text())
        url = f"http://127.0.0.1:{self.port}/"
        self.assertEqual(self.opened_urls(), [url])
        # The address was only handed to the browser once the server answered.
        self.assertIn("workbook_modified", get_json(f"{url}api/status"))
        with urllib.request.urlopen(url, timeout=5) as response:
            self.assertIn("<title>Raft</title>", response.read().decode())
        self.assertEqual(get_json(f"{url}api/months"),
                         [{"year": 2026, "month": 10, "label": "October 2026"}])
        output = log.read_text()
        self.assertIn(f"Workbook: {self.workbook}", output)
        self.assertIn(f"Address:  {url}  (this computer only)", output)
        self.assertLess(output.index("Address:"), output.index("Ready."))
        self.assertIsNone(process.poll())   # still running while in use

        # Not reachable through this computer's network address.
        try:
            lan = socket.gethostbyname(socket.gethostname())
        except OSError:
            lan = "127.0.0.1"
        if not lan.startswith("127."):
            self.assertFalse(port_open(self.port, lan), f"reachable at {lan}")

        self.stop(process)
        self.assertTrue(wait_for(lambda: not port_open(self.port), timeout=15))

    def test_startup_leaves_the_workbook_alone(self):
        before = self.workbook.read_bytes()
        self.start()
        self.assertTrue(wait_for(self.opened_urls))
        self.assertEqual(self.workbook.read_bytes(), before)
        self.assertEqual(sorted(p.name for p in self.tmp.iterdir() if p.suffix == ".xlsx"),
                         ["Fitness_Tracker.xlsx"])

    def test_missing_workbook_stops_with_a_message(self):
        missing = self.tmp / "gone" / "Fitness_Tracker.xlsx"
        env = dict(self.env, FITNESS_TRACKER_WORKBOOK=str(missing))
        process, log = self.start(env=env)
        self.assertEqual(process.wait(40), 1)
        output = log.read_text()
        self.assertIn("The workbook was not found", output)
        self.assertIn("Fitness Tracker did not start", output)
        self.assertFalse(missing.exists())
        self.assertFalse(port_open(self.port))
        self.assertEqual(self.opened_urls(), [])

    def test_port_in_use_by_another_program_stops_with_a_message(self):
        other = socket.socket()
        other.bind(("127.0.0.1", self.port))
        other.listen()
        self.addCleanup(other.close)
        process, log = self.start()
        self.assertEqual(process.wait(40), 1)
        output = log.read_text()
        self.assertIn(f"Port {self.port} is already in use by another program", output)
        self.assertIn("Fitness Tracker did not start", output)
        self.assertEqual(self.opened_urls(), [])

    def test_second_start_reuses_the_running_tracker(self):
        first, _ = self.start()
        self.assertTrue(wait_for(self.opened_urls))
        second, log = self.start()
        self.assertEqual(second.wait(40), 0)
        url = f"http://127.0.0.1:{self.port}/"
        self.assertIn(f"The tracker is already running at {url}", log.read_text())
        self.assertEqual(self.opened_urls(), [url, url])
        self.assertIsNone(first.poll())

    def test_port_given_on_the_command_line(self):
        port = free_port()
        env = {k: v for k, v in self.env.items() if k != "FITNESS_TRACKER_PORT"}
        self.start(args=f"--port {port}", env=env)
        self.assertTrue(wait_for(self.opened_urls))
        self.assertEqual(self.opened_urls(), [f"http://127.0.0.1:{port}/"])

    def test_works_from_a_moved_and_renamed_folder(self):
        try:
            from _winapi import CreateJunction
        except ImportError:
            self.skipTest("cannot link the Python environment into the copy")
        moved = self.tmp / "Moved Tracker (copy 2) & more"
        moved.mkdir()
        shutil.copy(BAT, moved)
        shutil.copy(PROJECT / "requirements.txt", moved)
        for folder in ("backend", "frontend"):
            shutil.copytree(PROJECT / folder, moved / folder,
                            ignore=shutil.ignore_patterns("__pycache__"))
        workbook = wbk.create_workbook(moved / "Fitness_Tracker.xlsx")
        # The copy shares the real environment through a link, removed again
        # before the temporary folder is deleted.
        link = moved / ".venv"
        CreateJunction(str(PROJECT / ".venv"), str(link))
        self.addCleanup(lambda: link.exists() and os.rmdir(link))

        env = {k: v for k, v in self.env.items() if k != "FITNESS_TRACKER_WORKBOOK"}
        process, log = self.start(bat=moved / BAT.name, env=env)
        self.assertTrue(wait_for(self.opened_urls), log.read_text())
        url = f"http://127.0.0.1:{self.port}/"
        self.assertEqual(self.opened_urls(), [url])
        # It found the workbook next to itself, not the one in the original folder.
        self.assertIn(f"Workbook: {workbook}", log.read_text())
        request = urllib.request.Request(
            f"{url}api/months", data=b'{"year": 2026, "month": 11}',
            headers={"Content-Type": "application/json"}, method="POST",
        )
        with urllib.request.urlopen(request, timeout=10) as response:
            self.assertEqual(response.status, 201)
        self.assertEqual(load_workbook(workbook).sheetnames, ["October 2026", "November 2026"])
        with urllib.request.urlopen(f"{url}app.js", timeout=5) as response:
            self.assertIn("pollLoop", response.read().decode())
        self.stop(process)
        os.rmdir(link)
        self.assertTrue((PROJECT / ".venv" / "Scripts" / "python.exe").exists())


if __name__ == "__main__":
    unittest.main()

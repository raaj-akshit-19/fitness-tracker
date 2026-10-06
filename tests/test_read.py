"""Checks for reading a month, and for telling a sheet the tracker laid out from any other.

Temporary workbooks only.

Run from the project root:  python -m unittest discover -s tests -v
"""

import hashlib
import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from backend import excel_live as xl
from backend import store as st
from backend import workbook as wbk
from backend.app import create_app
from tests.earlier_layout import add_month

DAY_KEYS = {"date", "exercise", "junk_food", "cardio_display", "cardio_seconds",
            "calories_kcal", "protein_g", "weight_lifted_kg"}
MEASUREMENT_KEYS = {"sunday", "weight_kg", "waist_cm", "chest_cm", "bicep_cm", "thigh_cm", "forearm_cm"}


def make_workbook(path, months, other=()):
    """A workbook with the given months. Those also listed in other are laid out another way."""
    wb = Workbook()
    wb.remove(wb.active)
    for year, month in months:
        add_month(wb, year, month, earlier=(year, month) in other)
    wb.save(path)
    return path


class Case(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.path = self.tmp / "Fitness_Tracker.xlsx"

    def workbook(self, months, other=()):
        make_workbook(self.path, months, other)
        return st.TrackerStore(self.path)

    def write(self, sheet, cells):
        wb = load_workbook(self.path)
        for ref, value in cells.items():
            wb[sheet][ref] = value
        wb.save(self.path)

    def digest(self):
        return hashlib.sha256(self.path.read_bytes()).hexdigest()


class LayoutTests(Case):
    """The tracker has one layout, and knows a sheet in it from any other."""

    def sheet(self):
        wb = Workbook()
        wb.remove(wb.active)
        return wbk.add_month_sheet(wb, 2026, 10)

    def test_a_sheet_the_tracker_made_is_recognised(self):
        self.assertIs(wbk.is_month_sheet(self.sheet()), True)

    def test_a_changed_header_is_not(self):
        for cell, text in (("C1", "No Junk Food"), ("E1", "Calories"), ("K1", "Weight"), ("A1", "Day"), ("G1", None)):
            ws = self.sheet()
            ws[cell] = text
            self.assertIs(wbk.is_month_sheet(ws), False, cell)

    def test_a_missing_table_or_a_sheet_that_is_not_a_month_is_not(self):
        ws = self.sheet()
        del ws.tables["Measurements_2026_10"]
        self.assertIs(wbk.is_month_sheet(ws), False)
        notes = Workbook().active
        notes.title = "Notes"
        self.assertIs(wbk.is_month_sheet(notes), False)

    def test_it_goes_by_table_and_header_names_not_by_position(self):
        # A sheet whose measurements table sits somewhere else is still the tracker's.
        ws = self.sheet()
        table = ws.tables["Measurements_2026_10"]
        for row in range(1, 6):
            for offset in range(7):
                ws.cell(row, 20 + offset, ws.cell(row, 10 + offset).value)
                ws.cell(row, 10 + offset).value = None
        table.ref = "T1:Z5"
        self.assertIs(wbk.is_month_sheet(ws), True)

    def test_a_month_laid_out_another_way_is_not_and_is_never_read_as_one(self):
        store = self.workbook([(2026, 10), (2026, 11)], other=[(2026, 10)])
        self.write("October 2026", {"B2": True, "E2": 1000})
        before = self.digest()
        wb = load_workbook(self.path)
        self.assertEqual([wbk.is_month_sheet(ws) for ws in wb], [False, True])
        # It is listed, since the sheet is there, but reading it says plainly that it cannot be read.
        self.assertEqual([m["label"] for m in store.list_months()], ["October 2026", "November 2026"])
        with self.assertRaises(st.WorkbookFormatError) as caught:
            store.get_month(2026, 10)
        self.assertIn("unexpected columns", str(caught.exception))
        self.assertEqual(len(store.get_month(2026, 11)["days"]), 30)       # the other month is not affected
        self.assertEqual(self.digest(), before)


class ReadTests(Case):
    def test_month_reads_with_its_field_names(self):
        store = self.workbook([(2026, 10)])
        self.write("October 2026", {
            "B2": "Completed", "C2": "None", "D2": "25:30", "E2": 2200, "F2": 140.5, "G2": 2450,
            "B3": "Rest Day", "C3": "Controlled", "B4": "Partial", "C4": "Had", "B5": "Missed", "G5": 0,
            "K2": 72.5, "L2": 84, "P5": 28,
        })
        data = store.get_month(2026, 10)
        self.assertEqual(set(data), {"year", "month", "label", "days", "measurements", "issues"})
        self.assertEqual((data["label"], len(data["days"]), len(data["measurements"])),
                         ("October 2026", 31, 4))
        self.assertEqual(data["days"][0], {
            "date": "2026-10-01", "exercise": "Completed", "junk_food": "None",
            "cardio_display": "25:30", "cardio_seconds": 1530,
            "calories_kcal": 2200, "protein_g": 140.5, "weight_lifted_kg": 2450,
        })
        self.assertEqual([d["exercise"] for d in data["days"][:5]],
                         ["Completed", "Rest Day", "Partial", "Missed", None])
        self.assertEqual([d["junk_food"] for d in data["days"][:4]], ["None", "Controlled", "Had", None])
        self.assertEqual(data["days"][3]["weight_lifted_kg"], 0)              # a recorded zero stays
        self.assertTrue(all(set(day) == DAY_KEYS for day in data["days"]))
        self.assertEqual(data["measurements"][0], {
            "sunday": "2026-10-04", "weight_kg": 72.5, "waist_cm": 84, "chest_cm": None,
            "bicep_cm": None, "thigh_cm": None, "forearm_cm": None,
        })
        self.assertEqual(data["measurements"][3]["forearm_cm"], 28)
        self.assertTrue(all(set(row) == MEASUREMENT_KEYS for row in data["measurements"]))
        self.assertEqual(data["issues"], [])

    def test_empty_month_is_all_blank(self):
        data = self.workbook([(2026, 11)]).get_month(2026, 11)
        self.assertEqual(len(data["days"]), 30)
        self.assertEqual([m["sunday"] for m in data["measurements"]],
                         ["2026-11-01", "2026-11-08", "2026-11-15", "2026-11-22", "2026-11-29"])
        for day in data["days"]:
            self.assertEqual([day[key] for key in sorted(DAY_KEYS - {"date"})], [None] * 7)

    def test_blank_and_the_word_none_are_different(self):
        store = self.workbook([(2026, 10)])
        self.write("October 2026", {"C2": "None"})
        days = store.get_month(2026, 10)["days"]
        self.assertEqual(days[0]["junk_food"], "None")
        self.assertIsNone(days[1]["junk_food"])

    def test_bad_cells_are_reported_not_converted(self):
        store = self.workbook([(2026, 10)])
        self.write("October 2026", {
            "B2": "Done", "C2": True, "D2": 25.30, "D3": "25:75", "E2": 2200.5, "F2": "lots",
            "G2": -5, "K2": 0, "B3": " rest  day ",
        })
        data = store.get_month(2026, 10)
        first = data["days"][0]
        self.assertEqual([first[k] for k in ("exercise", "junk_food", "cardio_seconds", "calories_kcal",
                                             "protein_g", "weight_lifted_kg")], [None] * 6)
        self.assertEqual(first["cardio_display"], "25.3")
        self.assertEqual(data["days"][1]["cardio_display"], "25:75")
        self.assertEqual(data["days"][1]["exercise"], "Rest Day")           # spelling tidied, not an issue
        self.assertIsNone(data["measurements"][0]["weight_kg"])
        self.assertEqual(sorted((i["date"], i["field"]) for i in data["issues"]), [
            ("2026-10-01", "calories_kcal"), ("2026-10-01", "cardio"), ("2026-10-01", "exercise"),
            ("2026-10-01", "junk_food"), ("2026-10-01", "protein_g"), ("2026-10-01", "weight_lifted_kg"),
            ("2026-10-02", "cardio"), ("2026-10-04", "weight_kg"),
        ])

    def test_another_layout_gives_a_clear_error(self):
        store = self.workbook([(2026, 10)])
        self.write("October 2026", {"C1": "Sweets"})
        with self.assertRaises(st.WorkbookFormatError) as caught:
            store.get_month(2026, 10)
        self.assertIn("unexpected columns", str(caught.exception))
        self.assertIn("Sweets", str(caught.exception))

    def test_wrong_dates_are_an_error(self):
        store = self.workbook([(2026, 10)])
        self.write("October 2026", {"A5": date(2026, 11, 4)})
        with self.assertRaises(st.WorkbookFormatError):
            store.get_month(2026, 10)
        self.write("October 2026", {"A5": date(2026, 10, 4), "J2": date(2026, 10, 5)})
        with self.assertRaises(st.WorkbookFormatError):
            store.get_month(2026, 10)

    @unittest.skipUnless(xl.AVAILABLE, "pywin32 is not installed")
    def test_month_reads_the_same_through_excel(self):
        from tests.test_excel_live import FakeBook, FakeExcel, locked_file
        store = self.workbook([(2026, 10)])
        self.write("October 2026", {"B2": "Rest Day", "C2": "None", "D2": "25:30", "E2": 2200,
                                    "F2": 140.5, "G2": 2450, "K2": 72.5})
        from_file = store.get_month(2026, 10)
        excel = FakeExcel()
        excel.open(FakeBook(self.path))
        with mock.patch.object(xl, "_running_documents", side_effect=excel.running), locked_file():
            from_excel = st.TrackerStore(self.path).get_month(2026, 10)
        self.assertEqual(from_excel, from_file)

    def test_reading_never_changes_the_file(self):
        store = self.workbook([(2026, 10), (2026, 11)])
        self.write("October 2026", {"B2": "Partial", "G2": 2000, "E2": 1800})
        before = self.digest()
        store.list_months()
        store.get_month(2026, 10)
        store.get_month(2026, 11)
        store.status()
        self.assertEqual(self.digest(), before)
        self.assertEqual(sorted(p.name for p in self.tmp.iterdir()), ["Fitness_Tracker.xlsx"])

    def test_a_new_month_is_in_the_one_layout(self):
        data = self.workbook([(2026, 10)]).create_month(2026, 11)
        self.assertEqual(set(data), {"year", "month", "label", "days", "measurements", "issues"})
        self.assertEqual(set(data["days"][0]), DAY_KEYS)
        wb = load_workbook(self.path)
        self.assertEqual([wbk.is_month_sheet(ws) for ws in wb], [True, True])
        self.assertEqual([wb["November 2026"].cell(1, col).value for col in range(1, 8)], wbk.DAILY_HEADERS)
        self.assertEqual([wb["November 2026"].cell(1, col).value for col in range(10, 17)], wbk.MEASUREMENT_HEADERS)

    @unittest.skipUnless(xl.AVAILABLE, "pywin32 is not installed")
    def test_month_reads_the_same_through_excel(self):
        from tests.test_excel_live import FakeBook, FakeExcel, locked_file
        store = self.workbook({(2026, 10): 2})
        self.write("October 2026", {"B2": "Rest Day", "C2": "None", "D2": "25:30", "E2": 2200,
                                    "F2": 140.5, "G2": 2450, "K2": 72.5})
        from_file = store.get_month(2026, 10)
        excel = FakeExcel()
        excel.open(FakeBook(self.path))
        with mock.patch.object(xl, "_running_documents", side_effect=excel.running), locked_file():
            from_excel = st.TrackerStore(self.path).get_month(2026, 10)
        self.assertEqual(from_excel, from_file)


class ApiTests(Case):
    def test_a_month_is_served_with_its_analytics(self):
        self.workbook([(2026, 10), (2026, 11)])
        self.write("November 2026", {"B2": "Completed", "E2": 2200})
        client = create_app(self.path, today=lambda: date(2026, 12, 1)).test_client()
        self.assertEqual(client.get("/api/months").get_json(), [
            {"year": 2026, "month": 10, "label": "October 2026"},
            {"year": 2026, "month": 11, "label": "November 2026"}])
        month = client.get("/api/months/2026/11")
        self.assertEqual(month.status_code, 200)
        self.assertEqual(month.get_json()["days"][0]["calories_kcal"], 2200)
        analytics = client.get("/api/months/2026/11/analytics")
        self.assertEqual(analytics.status_code, 200)
        self.assertEqual(analytics.get_json()["calories"]["total_kcal"], 2200)
        self.assertEqual(analytics.get_json()["exercise"]["counts"]["completed"], 1)

    def test_nothing_in_what_is_served_names_a_version(self):
        self.workbook([(2026, 10)])
        client = create_app(self.path).test_client()
        for url in ("/api/status", "/api/months", "/api/months/2026/10", "/api/months/2026/10/analytics"):
            text = client.get(url).get_data(as_text=True).lower()
            for word in ("version", "upgrade", "migrat", "legacy", '"v1"', '"v2"'):
                self.assertNotIn(word, text, (url, word))

    def test_there_is_no_way_to_convert_a_workbook_through_the_api(self):
        self.workbook([(2026, 10)], other=[(2026, 10)])
        client = create_app(self.path).test_client()
        before = self.digest()
        for method, url in (("post", "/api/upgrade"), ("put", "/api/upgrade"), ("post", "/api/migrate"),
                            ("post", "/api/convert")):
            response = getattr(client, method)(url, json={"confirm": True})
            self.assertIn(response.status_code, (404, 405), url)
        self.assertEqual(self.digest(), before)
        self.assertEqual(sorted(p.name for p in self.path.parent.iterdir()), [self.path.name])
        self.assertEqual(client.get("/api/months/2026/10").get_json()["error"]["code"], "workbook_format")


if __name__ == "__main__":
    unittest.main()

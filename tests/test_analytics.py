"""Checks for the month analytics.

The calculations are pure functions, so most tests hand them a month built
here. The API tests use temporary workbooks.

Run from the project root:  python -m unittest discover -s tests -v
"""

import calendar
import copy
import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from backend import analytics as an
from backend import store as st
from backend import workbook as wbk
from backend.app import create_app
from tests.test_migration import sha256

TODAY = date(2026, 10, 20)          # a Tuesday. 1 Oct 2026 is a Thursday.
MEASUREMENT_KEYS = ("weight_kg", "waist_cm", "chest_cm", "bicep_cm", "thigh_cm", "forearm_cm")


def month(year=2026, number=10, days=None, sundays=None):
    """A month as the store returns it, blank except for what is given.

    days is {day: {field: value}}; cardio is given as "m:ss" under "cardio".
    sundays is {day: {field: value}}.
    """
    records = []
    for day in range(1, calendar.monthrange(year, number)[1] + 1):
        entry = dict((days or {}).get(day, {}))
        cardio = entry.pop("cardio", None)
        record = {"date": date(year, number, day).isoformat(), "exercise": None, "junk_food": None,
                  "cardio_display": cardio, "cardio_seconds": None,
                  "calories_kcal": None, "protein_g": None, "weight_lifted_kg": None}
        if cardio is not None:
            minutes, seconds = cardio.split(":")
            record["cardio_seconds"] = int(minutes) * 60 + int(seconds)
        record.update(entry)
        records.append(record)
    measurements = []
    for sunday in wbk.month_sundays(year, number):
        row = {"sunday": sunday.isoformat(), **dict.fromkeys(MEASUREMENT_KEYS)}
        row.update((sundays or {}).get(sunday.day, {}))
        measurements.append(row)
    return {"year": year, "month": number, "label": wbk.sheet_name(year, number),
            "days": records, "measurements": measurements, "issues": []}


def exercise(values, today=TODAY, **kwargs):
    """Analytics for October with the given exercise statuses by day."""
    return an.month_analytics(month(days={day: {"exercise": value} for day, value in values.items()}, **kwargs), today)


def junk(values, today=TODAY):
    return an.month_analytics(month(days={day: {"junk_food": value} for day, value in values.items()}), today)


# The exercise month used by several tests. 20 Oct is today and blank.
EXERCISE = {1: "Completed", 2: "Completed", 3: "Rest Day", 4: "Completed", 5: "Partial", 6: "Completed",
            7: "Missed", 9: "Rest Day", 10: "Completed", 11: "Completed", 18: "Completed", 19: "Completed",
            22: "Completed"}
# 8 Oct and 12 to 17 Oct are past and blank; 22 Oct is in the future.

JUNK = {1: "None", 2: "None", 3: "Controlled", 4: "None", 5: "Had", 6: "None", 7: "None", 8: "None",
        10: "None", 11: "None", 12: "Controlled", **{day: "None" for day in range(13, 20)}}
# 9 Oct is past and blank.


class ExerciseTests(unittest.TestCase):
    def test_each_day_is_classified(self):
        daily = {entry["date"][-2:]: entry for entry in exercise(EXERCISE)["daily"]}
        self.assertEqual([daily[day]["exercise_status"] for day in ("01", "05", "03", "07")],
                         ["completed", "partial", "rest_day", "missed"])
        self.assertEqual(daily["08"]["exercise_status"], "missed")          # past, left blank
        self.assertEqual(daily["20"]["exercise_status"], "pending")         # today, left blank
        self.assertEqual(daily["21"]["exercise_status"], "pending")         # future, blank
        self.assertEqual(daily["22"]["exercise_status"], "pending")         # future, whatever it holds
        self.assertEqual(daily["31"]["exercise_status"], "pending")

    def test_credit_for_each_status(self):
        daily = {entry["date"][-2:]: entry for entry in exercise(EXERCISE)["daily"]}
        self.assertEqual([daily[day]["exercise_credit"] for day in ("01", "05", "07", "08")], [1.0, 0.5, 0.0, 0.0])
        self.assertIsNone(daily["03"]["exercise_credit"])                   # rest day: not counted
        self.assertIsNone(daily["20"]["exercise_credit"])                   # pending: not counted

    def test_stored_value_stays_visible_beside_how_it_was_read(self):
        daily = {entry["date"][-2:]: entry for entry in exercise(EXERCISE)["daily"]}
        self.assertEqual((daily["07"]["exercise"], daily["07"]["exercise_status"]), ("Missed", "missed"))
        self.assertEqual((daily["08"]["exercise"], daily["08"]["exercise_status"]), (None, "missed"))
        self.assertEqual((daily["22"]["exercise"], daily["22"]["exercise_status"]), ("Completed", "pending"))
        self.assertEqual((daily["03"]["exercise"], daily["03"]["exercise_status"]), ("Rest Day", "rest_day"))

    def test_counts(self):
        result = exercise(EXERCISE)["exercise"]
        self.assertEqual(result["counts"], {"completed": 8, "partial": 1, "rest_day": 2, "missed": 8, "pending": 12})
        self.assertEqual(sum(result["counts"].values()), 31)
        self.assertEqual(result["missed_not_entered"], 7)                   # 8 Oct and 12 to 17 Oct
        self.assertEqual(list(result["counts"]), ["completed", "partial", "rest_day", "missed", "pending"])

    def test_completion_percentage(self):
        result = exercise(EXERCISE)["exercise"]
        self.assertEqual((result["eligible_days"], result["credits"]), (17, 8.5))
        self.assertEqual(result["completion_percentage"], 50.0)

    def test_the_worked_example(self):
        # Completed, Partial, Missed, Rest Day => (1 + 0.5 + 0) / 3 = 50%
        result = exercise({1: "Completed", 2: "Partial", 3: "Missed", 4: "Rest Day"}, today=date(2026, 10, 4))
        self.assertEqual(result["exercise"]["completion_percentage"], 50.0)
        self.assertEqual(result["exercise"]["eligible_days"], 3)

    def test_rest_days_do_not_lower_or_raise_the_percentage(self):
        without = exercise({1: "Completed", 2: "Missed"}, today=date(2026, 10, 2))["exercise"]
        with_rest = exercise({1: "Completed", 2: "Missed", 3: "Rest Day", 4: "Rest Day"},
                             today=date(2026, 10, 4))["exercise"]
        self.assertEqual(without["completion_percentage"], 50.0)
        self.assertEqual(with_rest["completion_percentage"], 50.0)
        self.assertEqual(with_rest["counts"]["rest_day"], 2)

    def test_percentage_is_none_when_nothing_is_eligible(self):
        only_rest = exercise({1: "Rest Day", 2: "Rest Day"}, today=date(2026, 10, 2))["exercise"]
        self.assertIsNone(only_rest["completion_percentage"])
        self.assertEqual(only_rest["eligible_days"], 0)
        not_started = exercise({}, today=date(2026, 9, 30))["exercise"]
        self.assertIsNone(not_started["completion_percentage"])
        self.assertEqual(not_started["counts"]["pending"], 31)

    def test_percentage_extremes(self):
        self.assertEqual(exercise({1: "Completed", 2: "Completed"}, today=date(2026, 10, 2))
                         ["exercise"]["completion_percentage"], 100.0)
        self.assertEqual(exercise({1: "Missed"}, today=date(2026, 10, 2))["exercise"]["completion_percentage"], 0.0)
        self.assertEqual(exercise({1: "Partial", 2: "Partial"}, today=date(2026, 10, 2))
                         ["exercise"]["completion_percentage"], 50.0)
        thirds = exercise({1: "Completed", 2: "Partial", 3: "Partial"}, today=date(2026, 10, 3))["exercise"]
        self.assertEqual(thirds["completion_percentage"], 66.67)

    def test_blank_today_is_pending_but_a_blank_past_day_is_missed(self):
        blank = exercise({1: "Completed"}, today=date(2026, 10, 2))["exercise"]
        self.assertEqual((blank["counts"]["missed"], blank["counts"]["pending"], blank["completion_percentage"]),
                         (0, 30, 100.0))
        next_day = exercise({1: "Completed"}, today=date(2026, 10, 3))["exercise"]
        self.assertEqual((next_day["counts"]["missed"], next_day["counts"]["pending"],
                          next_day["completion_percentage"]), (1, 29, 50.0))

    def test_a_status_entered_today_counts_today(self):
        for status, key in (("Completed", "completed"), ("Partial", "partial"), ("Rest Day", "rest_day"),
                            ("Missed", "missed")):
            result = exercise({20: status})
            self.assertEqual(result["daily"][19]["exercise_status"], key)
            self.assertEqual(result["exercise"]["counts"]["pending"], 11)

    def test_streaks(self):
        result = exercise(EXERCISE)["exercise"]
        # 1, 2, (rest), 4 is the longest run; 18, 19 is the current one, today being pending.
        self.assertEqual((result["current_streak"], result["longest_streak"]), (2, 3))

    def test_rest_day_is_neutral_in_a_streak(self):
        result = exercise({1: "Completed", 2: "Rest Day", 3: "Rest Day", 4: "Completed", 5: "Rest Day"},
                          today=date(2026, 10, 5))["exercise"]
        self.assertEqual((result["current_streak"], result["longest_streak"]), (2, 2))
        # A rest day alone starts nothing and adds nothing.
        only = exercise({1: "Rest Day", 2: "Rest Day"}, today=date(2026, 10, 2))["exercise"]
        self.assertEqual((only["current_streak"], only["longest_streak"]), (0, 0))

    def test_partial_breaks_a_streak_and_does_not_start_one(self):
        result = exercise({1: "Completed", 2: "Completed", 3: "Partial", 4: "Completed"},
                          today=date(2026, 10, 4))["exercise"]
        self.assertEqual((result["current_streak"], result["longest_streak"]), (1, 2))
        partials = exercise({1: "Partial", 2: "Partial", 3: "Partial"}, today=date(2026, 10, 3))["exercise"]
        self.assertEqual((partials["current_streak"], partials["longest_streak"]), (0, 0))

    def test_missed_and_past_blank_break_a_streak(self):
        missed = exercise({1: "Completed", 2: "Completed", 3: "Missed", 4: "Completed"},
                          today=date(2026, 10, 4))["exercise"]
        self.assertEqual((missed["current_streak"], missed["longest_streak"]), (1, 2))
        blank = exercise({1: "Completed", 2: "Completed", 4: "Completed"}, today=date(2026, 10, 4))["exercise"]
        self.assertEqual((blank["current_streak"], blank["longest_streak"]), (1, 2))

    def test_today_changes_the_current_streak_only_once_it_is_entered(self):
        self.assertEqual(exercise(EXERCISE)["exercise"]["current_streak"], 2)                       # blank
        self.assertEqual(exercise({**EXERCISE, 20: "Completed"})["exercise"]["current_streak"], 3)
        self.assertEqual(exercise({**EXERCISE, 20: "Rest Day"})["exercise"]["current_streak"], 2)
        self.assertEqual(exercise({**EXERCISE, 20: "Partial"})["exercise"]["current_streak"], 0)
        self.assertEqual(exercise({**EXERCISE, 20: "Missed"})["exercise"]["current_streak"], 0)

    def test_future_entries_do_not_affect_counts_or_streaks(self):
        with_future = exercise({**EXERCISE, 21: "Completed", 23: "Missed", 30: "Partial"})["exercise"]
        without = exercise({day: value for day, value in EXERCISE.items() if day <= 20})["exercise"]
        self.assertEqual(with_future, without)

    def test_streak_does_not_reach_outside_the_month(self):
        every = exercise({day: "Completed" for day in range(1, 32)}, today=date(2026, 11, 15))["exercise"]
        self.assertEqual((every["current_streak"], every["longest_streak"], every["counts"]["completed"]),
                         (31, 31, 31))


class JunkFoodTests(unittest.TestCase):
    def test_each_day_is_classified(self):
        daily = {entry["date"][-2:]: entry for entry in junk({**JUNK, 25: "Had"})["daily"]}
        self.assertEqual([daily[day]["junk_food_status"] for day in ("01", "03", "05")], ["none", "controlled", "had"])
        self.assertEqual((daily["09"]["junk_food"], daily["09"]["junk_food_status"]), (None, "missed"))
        self.assertEqual(daily["20"]["junk_food_status"], "pending")        # today, blank
        self.assertEqual(daily["21"]["junk_food_status"], "pending")        # future, blank
        self.assertEqual((daily["25"]["junk_food"], daily["25"]["junk_food_status"]), ("Had", "pending"))

    def test_credit_for_each_status(self):
        daily = {entry["date"][-2:]: entry for entry in junk(JUNK)["daily"]}
        self.assertEqual([daily[day]["junk_food_credit"] for day in ("01", "03", "05")], [1.0, 0.5, 0.0])
        self.assertIsNone(daily["09"]["junk_food_credit"])                  # no entry: not in the percentage
        self.assertIsNone(daily["20"]["junk_food_credit"])

    def test_the_text_none_is_a_status_and_a_blank_is_not(self):
        result = junk({1: "None"}, today=date(2026, 10, 2))
        self.assertEqual((result["daily"][0]["junk_food"], result["daily"][0]["junk_food_status"]), ("None", "none"))
        self.assertEqual((result["daily"][1]["junk_food"], result["daily"][1]["junk_food_status"]), (None, "pending"))

    def test_counts(self):
        result = junk(JUNK)["junk_food"]
        self.assertEqual(result["counts"], {"none": 15, "controlled": 2, "had": 1, "missed": 1, "pending": 12})
        self.assertEqual(sum(result["counts"].values()), 31)
        self.assertEqual(list(result["counts"]), ["none", "controlled", "had", "missed", "pending"])

    def test_success_percentage_is_over_the_days_with_an_entry(self):
        result = junk(JUNK)["junk_food"]
        self.assertEqual((result["eligible_days"], result["credits"]), (18, 16.0))
        self.assertEqual(result["success_percentage"], 88.89)
        # A past blank is counted as missed but is not one of the recorded days.
        self.assertEqual(junk({1: "None", 3: "Had"}, today=date(2026, 10, 3))["junk_food"]["success_percentage"], 50.0)
        self.assertEqual(junk({1: "None", 2: "Controlled", 3: "Had"}, today=date(2026, 10, 3))
                         ["junk_food"]["success_percentage"], 50.0)
        self.assertEqual(junk({1: "Controlled"}, today=date(2026, 10, 1))["junk_food"]["success_percentage"], 50.0)
        self.assertEqual(junk({1: "Had"}, today=date(2026, 10, 1))["junk_food"]["success_percentage"], 0.0)

    def test_percentage_is_none_with_no_entries(self):
        self.assertIsNone(junk({}, today=date(2026, 10, 5))["junk_food"]["success_percentage"])
        self.assertEqual(junk({}, today=date(2026, 10, 5))["junk_food"]["counts"]["missed"], 4)
        self.assertIsNone(junk({}, today=date(2026, 9, 1))["junk_food"]["success_percentage"])

    def test_streaks(self):
        result = junk(JUNK)["junk_food"]
        self.assertEqual((result["current_streak"], result["longest_streak"]), (7, 7))

    def test_controlled_had_and_past_blank_each_break_the_clean_streak(self):
        for breaker in ("Controlled", "Had", None):
            days = {1: "None", 2: "None", 3: "None", 5: "None"}
            if breaker:
                days[4] = breaker
            result = junk(days, today=date(2026, 10, 5))["junk_food"]
            self.assertEqual((result["current_streak"], result["longest_streak"]), (1, 3), breaker)

    def test_today_and_future_blanks_do_not_affect_the_streak(self):
        self.assertEqual(junk(JUNK)["junk_food"]["current_streak"], 7)                          # today blank
        self.assertEqual(junk({**JUNK, 20: "None"})["junk_food"]["current_streak"], 8)
        self.assertEqual(junk({**JUNK, 20: "Controlled"})["junk_food"]["current_streak"], 0)
        self.assertEqual(junk({**JUNK, 20: "Had"})["junk_food"]["current_streak"], 0)
        self.assertEqual(junk({**JUNK, 21: "Had", 22: "None"})["junk_food"], junk(JUNK)["junk_food"])

    def test_the_two_habits_are_worked_out_separately(self):
        data = month(days={1: {"exercise": "Completed", "junk_food": "Had"}, 2: {"exercise": "Missed"}})
        result = an.month_analytics(data, date(2026, 10, 2))
        self.assertEqual(result["exercise"]["completion_percentage"], 50.0)
        self.assertEqual(result["junk_food"]["success_percentage"], 0.0)
        self.assertEqual(result["junk_food"]["counts"]["pending"], 30)      # 2 Oct is today and blank


NUTRITION = {
    1: {"calories_kcal": 2200, "protein_g": 140.5, "cardio": "25:30", "weight_lifted_kg": 2450},
    2: {"calories_kcal": 1800, "protein_g": 120, "cardio": "0:45", "weight_lifted_kg": 1820},
    4: {"calories_kcal": 0},
    6: {"calories_kcal": 2500, "protein_g": 0, "cardio": "60:00", "weight_lifted_kg": 3000},
    12: {"calories_kcal": 2000, "protein_g": 99.5, "cardio": "0:00"},
    13: {"weight_lifted_kg": 0},
}


class NumberTests(unittest.TestCase):
    def setUp(self):
        self.result = an.month_analytics(month(days=NUTRITION), TODAY)

    def test_calories(self):
        self.assertEqual(self.result["calories"], {
            "days_recorded": 5, "total_kcal": 8500, "average_kcal": 1700,
            "highest_kcal": 2500, "highest_date": "2026-10-06", "lowest_kcal": 0, "lowest_date": "2026-10-04",
        })

    def test_protein(self):
        self.assertEqual(self.result["protein"], {
            "days_recorded": 4, "total_g": 360, "average_g": 90,
            "highest_g": 140.5, "highest_date": "2026-10-01", "lowest_g": 0, "lowest_date": "2026-10-06",
        })

    def test_cardio(self):
        self.assertEqual(self.result["cardio"], {
            "days_recorded": 4, "total_seconds": 5175, "total_display": "86:15",
            "average_seconds": 1294, "average_display": "21:34",
            "longest_seconds": 3600, "longest_display": "60:00", "longest_date": "2026-10-06",
        })

    def test_weight_lifted(self):
        self.assertEqual(self.result["weight_lifted"], {
            "days_recorded": 4, "total_kg": 7270, "average_kg": 1817.5,
            "highest_kg": 3000, "highest_date": "2026-10-06", "lowest_kg": 0, "lowest_date": "2026-10-13",
        })

    def test_a_recorded_zero_is_an_entry_and_a_blank_is_not(self):
        zero = an.month_analytics(month(days={1: {"calories_kcal": 0, "protein_g": 0}}), TODAY)
        self.assertEqual(zero["calories"], {
            "days_recorded": 1, "total_kcal": 0, "average_kcal": 0,
            "highest_kcal": 0, "highest_date": "2026-10-01", "lowest_kcal": 0, "lowest_date": "2026-10-01",
        })
        self.assertEqual((zero["protein"]["days_recorded"], zero["protein"]["average_g"]), (1, 0))

    def test_nothing_recorded_gives_none_not_zero(self):
        empty = an.month_analytics(month(), TODAY)
        for name, unit in (("calories", "kcal"), ("protein", "g"), ("weight_lifted", "kg")):
            self.assertEqual(empty[name], {
                "days_recorded": 0, f"total_{unit}": None, f"average_{unit}": None,
                f"highest_{unit}": None, "highest_date": None, f"lowest_{unit}": None, "lowest_date": None,
            }, name)
        self.assertEqual(set(empty["cardio"].values()), {0, None})
        self.assertEqual(empty["cardio"]["days_recorded"], 0)

    def test_average_is_over_recorded_days_only(self):
        result = an.month_analytics(month(days={1: {"calories_kcal": 2000}, 9: {"calories_kcal": 1000},
                                                   17: {"calories_kcal": 3100}}), TODAY)
        self.assertEqual((result["calories"]["average_kcal"], result["calories"]["days_recorded"]), (2033.33, 3))
        protein = an.month_analytics(month(days={1: {"protein_g": 100.5}, 2: {"protein_g": 99.7}}), TODAY)
        self.assertEqual((protein["protein"]["total_g"], protein["protein"]["average_g"]), (200.2, 100.1))

    def test_ties_name_the_earliest_day(self):
        result = an.month_analytics(month(days={3: {"calories_kcal": 2000}, 5: {"calories_kcal": 2000},
                                                   9: {"calories_kcal": 1500}, 11: {"calories_kcal": 1500}}), TODAY)
        self.assertEqual((result["calories"]["highest_date"], result["calories"]["lowest_date"]),
                         ("2026-10-03", "2026-10-09"))

    def test_numbers_do_not_depend_on_the_habit_statuses_or_on_today(self):
        early = an.month_analytics(month(days=NUTRITION), date(2026, 10, 1))
        for name in ("calories", "protein", "cardio", "weight_lifted"):
            self.assertEqual(early[name], self.result[name], name)

    def test_daily_entries_carry_the_values_as_stored(self):
        daily = self.result["daily"]
        self.assertEqual(len(daily), 31)
        self.assertEqual(daily[0], {
            "date": "2026-10-01", "exercise": None, "exercise_status": "missed", "exercise_credit": 0.0,
            "junk_food": None, "junk_food_status": "missed", "junk_food_credit": None,
            "cardio_seconds": 1530, "cardio_display": "25:30",
            "calories_kcal": 2200, "protein_g": 140.5, "weight_lifted_kg": 2450,
        })
        self.assertEqual([daily[3][key] for key in ("calories_kcal", "protein_g", "cardio_seconds", "cardio_display",
                                                   "weight_lifted_kg")], [0, None, None, None, None])
        self.assertEqual((daily[11]["cardio_seconds"], daily[11]["cardio_display"]), (0, "0:00"))
        self.assertEqual([daily[30][key] for key in ("calories_kcal", "protein_g", "cardio_seconds",
                                                    "weight_lifted_kg")], [None] * 4)

    def test_no_targets_or_goals_anywhere(self):
        text = str(self.result).lower()
        for word in ("goal", "target", "score", "remaining", "deficit", "surplus"):
            self.assertNotIn(word, text)


class WeeklyTests(unittest.TestCase):
    def test_weeks_run_monday_to_sunday_and_are_cut_at_the_month(self):
        weekly = an.month_analytics(month(), TODAY)["weekly"]
        self.assertEqual([(w["week"], w["start"], w["end"], w["days"]) for w in weekly], [
            (1, "2026-10-01", "2026-10-04", 4), (2, "2026-10-05", "2026-10-11", 7),
            (3, "2026-10-12", "2026-10-18", 7), (4, "2026-10-19", "2026-10-25", 7),
            (5, "2026-10-26", "2026-10-31", 6),
        ])
        self.assertEqual(sum(w["days"] for w in weekly), 31)
        self.assertEqual([w["due_days"] for w in weekly], [4, 7, 7, 1, 0])     # 19 Oct; today is blank

    def test_other_month_shapes(self):
        november = an.month_analytics(month(2026, 11), date(2026, 12, 1))["weekly"]      # starts on a Sunday
        self.assertEqual([(w["start"][-2:], w["end"][-2:]) for w in november],
                         [("01", "01"), ("02", "08"), ("09", "15"), ("16", "22"), ("23", "29"), ("30", "30")])
        february = an.month_analytics(month(2027, 2), date(2027, 3, 1))["weekly"]        # four whole weeks
        self.assertEqual([w["days"] for w in february], [7, 7, 7, 7])
        june = an.month_analytics(month(2026, 6), date(2026, 7, 1))["weekly"]            # starts on a Monday
        self.assertEqual([(w["start"][-2:], w["end"][-2:]) for w in june],
                         [("01", "07"), ("08", "14"), ("15", "21"), ("22", "28"), ("29", "30")])

    def test_weekly_exercise(self):
        weekly = exercise(EXERCISE)["weekly"]
        self.assertEqual(weekly[0]["exercise"], {"completed": 3, "partial": 0, "rest_day": 1, "missed": 0,
                                                 "pending": 0, "completion_percentage": 100.0})
        # 5 Partial, 6 Completed, 7 Missed, 8 blank, 9 Rest Day, 10 and 11 Completed: 3.5 of 6.
        self.assertEqual(weekly[1]["exercise"], {"completed": 3, "partial": 1, "rest_day": 1, "missed": 2,
                                                 "pending": 0, "completion_percentage": 58.33})
        self.assertEqual(weekly[2]["exercise"], {"completed": 1, "partial": 0, "rest_day": 0, "missed": 6,
                                                 "pending": 0, "completion_percentage": 14.29})
        self.assertEqual(weekly[3]["exercise"], {"completed": 1, "partial": 0, "rest_day": 0, "missed": 0,
                                                 "pending": 6, "completion_percentage": 100.0})
        self.assertEqual(weekly[4]["exercise"], {"completed": 0, "partial": 0, "rest_day": 0, "missed": 0,
                                                 "pending": 6, "completion_percentage": None})

    def test_weekly_junk_food(self):
        weekly = junk(JUNK)["weekly"]
        self.assertEqual(weekly[0]["junk_food"], {"none": 3, "controlled": 1, "had": 0, "missed": 0, "pending": 0,
                                                  "success_percentage": 87.5})
        self.assertEqual(weekly[1]["junk_food"], {"none": 5, "controlled": 0, "had": 1, "missed": 1, "pending": 0,
                                                  "success_percentage": 83.33})
        self.assertEqual(weekly[2]["junk_food"], {"none": 6, "controlled": 1, "had": 0, "missed": 0, "pending": 0,
                                                  "success_percentage": 92.86})
        self.assertEqual(weekly[3]["junk_food"], {"none": 1, "controlled": 0, "had": 0, "missed": 0, "pending": 6,
                                                  "success_percentage": 100.0})
        self.assertIsNone(weekly[4]["junk_food"]["success_percentage"])

    def test_weekly_counts_add_up_to_the_month(self):
        for result, name in ((exercise(EXERCISE), "exercise"), (junk(JUNK), "junk_food")):
            for status, total in result[name]["counts"].items():
                self.assertEqual(sum(week[name][status] for week in result["weekly"]), total, (name, status))

    def test_weekly_calories_and_protein(self):
        weekly = an.month_analytics(month(days=NUTRITION), TODAY)["weekly"]
        self.assertEqual(weekly[0]["calories"], {"days_recorded": 3, "total_kcal": 4000, "average_kcal": 1333.33,
                                                 "highest_kcal": 2200, "lowest_kcal": 0})
        self.assertEqual(weekly[0]["protein"], {"days_recorded": 2, "total_g": 260.5, "average_g": 130.25,
                                                "highest_g": 140.5, "lowest_g": 120})
        self.assertEqual(weekly[1]["calories"], {"days_recorded": 1, "total_kcal": 2500, "average_kcal": 2500,
                                                 "highest_kcal": 2500, "lowest_kcal": 2500})
        self.assertEqual(weekly[1]["protein"], {"days_recorded": 1, "total_g": 0, "average_g": 0,
                                                "highest_g": 0, "lowest_g": 0})                # a recorded zero
        self.assertEqual(weekly[2]["protein"]["total_g"], 99.5)
        for week in weekly[3:]:
            self.assertEqual(week["calories"], {"days_recorded": 0, "total_kcal": None, "average_kcal": None,
                                                "highest_kcal": None, "lowest_kcal": None})
            self.assertEqual(week["protein"]["total_g"], None)

    def test_weekly_cardio_and_weight_lifted(self):
        weekly = an.month_analytics(month(days=NUTRITION), TODAY)["weekly"]
        self.assertEqual([week["cardio"] for week in weekly], [
            {"days_recorded": 2, "total_seconds": 1575, "total_display": "26:15"},
            {"days_recorded": 1, "total_seconds": 3600, "total_display": "60:00"},
            {"days_recorded": 1, "total_seconds": 0, "total_display": "0:00"},
            {"days_recorded": 0, "total_seconds": None, "total_display": None},
            {"days_recorded": 0, "total_seconds": None, "total_display": None},
        ])
        self.assertEqual([week["weight_lifted"] for week in weekly], [
            {"days_recorded": 2, "total_kg": 4270, "average_kg": 2135},
            {"days_recorded": 1, "total_kg": 3000, "average_kg": 3000},
            {"days_recorded": 1, "total_kg": 0, "average_kg": 0},
            {"days_recorded": 0, "total_kg": None, "average_kg": None},
            {"days_recorded": 0, "total_kg": None, "average_kg": None},
        ])

    def test_weekly_totals_add_up_to_the_month(self):
        result = an.month_analytics(month(days=NUTRITION), TODAY)
        for name, key in (("calories", "total_kcal"), ("protein", "total_g"), ("weight_lifted", "total_kg"),
                          ("cardio", "total_seconds")):
            self.assertEqual(sum(week[name][key] or 0 for week in result["weekly"]), result[name][key], name)
            self.assertEqual(sum(week[name]["days_recorded"] for week in result["weekly"]),
                             result[name]["days_recorded"], name)

    def test_week_keys(self):
        week = an.month_analytics(month(), TODAY)["weekly"][0]
        self.assertEqual(set(week), {"week", "start", "end", "days", "due_days", "exercise", "junk_food",
                                     "cardio", "weight_lifted", "calories", "protein"})


class MeasurementTests(unittest.TestCase):
    SUNDAYS = {4: {"weight_kg": 72.5, "waist_cm": 84}, 11: {"weight_kg": 72}, 18: {"waist_cm": 82.5}}

    def test_latest_recorded_sunday_against_the_one_before(self):
        result = an.month_analytics(month(sundays=self.SUNDAYS), TODAY)["measurements"]
        self.assertEqual(set(result), set(MEASUREMENT_KEYS))
        weight = result["weight_kg"]
        self.assertEqual((weight["current"], weight["current_date"], weight["previous"], weight["previous_date"]),
                         (72, "2026-10-11", 72.5, "2026-10-04"))
        self.assertEqual((weight["change"], weight["change_percentage"]), (-0.5, -0.69))
        self.assertEqual([p["value"] for p in weight["trend"]], [72.5, 72, None, None])
        # Waist was not measured on the Sunday before the latest one: no change to show.
        waist = result["waist_cm"]
        self.assertEqual((waist["current"], waist["current_date"], waist["previous"], waist["previous_date"],
                          waist["change"]), (82.5, "2026-10-18", None, "2026-10-11", None))

    def test_blank_measurements(self):
        result = an.month_analytics(month(sundays=self.SUNDAYS), TODAY)["measurements"]
        chest = result["chest_cm"]
        self.assertEqual([chest[key] for key in ("current", "current_date", "previous", "previous_date", "change",
                                                 "change_percentage")], [None] * 6)
        self.assertEqual([p["date"] for p in chest["trend"]], ["2026-10-04", "2026-10-11", "2026-10-18", "2026-10-25"])
        empty = an.month_analytics(month(), TODAY)["measurements"]
        self.assertTrue(all(empty[key]["current"] is None for key in MEASUREMENT_KEYS))

    def test_each_measurement_is_summarised_on_its_own(self):
        from backend.analytics import measurement_summary
        data = month(sundays=self.SUNDAYS)
        result = an.month_analytics(data, TODAY)["measurements"]
        for key in MEASUREMENT_KEYS:
            self.assertEqual(result[key], measurement_summary(data["measurements"], key), key)

    def test_first_sunday_has_no_previous_value_from_another_month(self):
        result = an.month_analytics(month(2026, 11, sundays={1: {"weight_kg": 80}}), date(2026, 11, 30))
        weight = result["measurements"]["weight_kg"]
        self.assertEqual((weight["current"], weight["previous"], weight["previous_date"], weight["change"]),
                         (80, None, None, None))


class PayloadTests(unittest.TestCase):
    def test_top_level_shape(self):
        result = an.month_analytics(month(days=NUTRITION), TODAY)
        self.assertEqual(list(result), ["month", "exercise", "junk_food", "daily", "weekly", "cardio",
                                        "weight_lifted", "calories", "protein", "measurements", "issues"])
        self.assertEqual(result["month"], {"year": 2026, "month": 10, "label": "October 2026", "days_in_month": 31,
                                           "due_days": 19, "pending_days": 12, "today": "2026-10-20"})
        self.assertEqual(set(result["exercise"]), {"counts", "missed_not_entered", "eligible_days", "credits",
                                                   "completion_percentage", "current_streak", "longest_streak"})
        self.assertEqual(set(result["junk_food"]), {"counts", "eligible_days", "credits", "success_percentage",
                                                    "current_streak", "longest_streak"})

    def test_it_is_a_pure_function(self):
        data = month(days={**{day: {"exercise": value} for day, value in EXERCISE.items()}, 1: NUTRITION[1]},
                     sundays={4: {"weight_kg": 72.5}})
        before = copy.deepcopy(data)
        first = an.month_analytics(data, TODAY)
        second = an.month_analytics(data, TODAY)
        self.assertEqual(data, before)                 # the input is not changed
        self.assertEqual(first, second)
        import inspect
        source = inspect.getsource(an)
        for word in ("flask", "request", "load_workbook", "date.today", "datetime.now", "open("):
            self.assertNotIn(word, source)

    def test_issues_are_passed_through_and_unreadable_cells_count_as_blank(self):
        data = month(days={1: {"exercise": "Completed"}})
        data["issues"] = [{"date": "2026-10-02", "field": "exercise", "value": "Done", "message": "x"}]
        result = an.month_analytics(data, date(2026, 10, 3))
        self.assertEqual(result["issues"], data["issues"])
        self.assertEqual(result["daily"][1]["exercise_status"], "missed")       # nothing readable was entered

    def test_month_fully_in_the_past_and_fully_in_the_future(self):
        past = an.month_analytics(month(), date(2026, 11, 5))
        self.assertEqual(past["exercise"]["counts"], {"completed": 0, "partial": 0, "rest_day": 0, "missed": 31,
                                                      "pending": 0})
        self.assertEqual((past["exercise"]["completion_percentage"], past["month"]["pending_days"]), (0.0, 0))
        self.assertEqual(past["junk_food"]["counts"]["missed"], 31)
        self.assertIsNone(past["junk_food"]["success_percentage"])
        future = an.month_analytics(month(), date(2026, 9, 5))
        self.assertEqual(future["exercise"]["counts"]["pending"], 31)
        self.assertEqual(future["junk_food"]["counts"]["pending"], 31)
        self.assertEqual(future["month"]["due_days"], 0)


class ApiTests(unittest.TestCase):
    """A workbook with September, October and November."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.path = self.tmp / "Fitness_Tracker.xlsx"
        wb = Workbook()
        wb.remove(wb.active)
        september = wbk.add_month_sheet(wb, 2026, 9)
        october = wbk.add_month_sheet(wb, 2026, 10)
        november = wbk.add_month_sheet(wb, 2026, 11)
        for ref, value in {"B2": "Completed", "C2": "Had", "B3": "Completed", "D2": "10:00", "G2": 900, "K2": 73}.items():
            september[ref] = value
        for ref, value in {
            "B2": "Completed", "C2": "None", "D2": "25:30", "E2": 2200, "F2": 140.5, "G2": 2450,
            "B3": "Partial", "C3": "Controlled", "E3": 1800,
            "B4": "Rest Day", "C4": "Had",
            "B5": "Missed", "E5": 0,
            "B7": "Done", "E7": "lots",                        # typed wrongly in Excel
            "K2": 72.5, "L2": 84, "K3": 72,
        }.items():
            october[ref] = value
        for ref, value in {"B2": "Completed", "E2": 9999, "F2": 999, "D2": "99:00", "G2": 99999, "K2": 80}.items():
            november[ref] = value
        wb.save(self.path)
        self.before = sha256(self.path)
        self.client = create_app(self.path, today=lambda: TODAY).test_client()

    def analytics(self, month_number):
        response = self.client.get(f"/api/months/2026/{month_number}/analytics")
        self.assertEqual(response.status_code, 200)
        return response.get_json()

    def test_a_month_gets_its_analytics(self):
        result = self.analytics(10)
        self.assertNotIn("version", result)
        self.assertEqual(result["month"]["today"], "2026-10-20")
        self.assertEqual(result["exercise"]["counts"],
                         {"completed": 1, "partial": 1, "rest_day": 1, "missed": 16, "pending": 12})
        self.assertEqual(result["exercise"]["missed_not_entered"], 15)
        # 1.5 credits over 18 eligible days.
        self.assertEqual(result["exercise"]["completion_percentage"], 8.33)
        self.assertEqual(result["junk_food"]["counts"],
                         {"none": 1, "controlled": 1, "had": 1, "missed": 16, "pending": 12})
        self.assertEqual(result["junk_food"]["success_percentage"], 50.0)
        self.assertEqual((result["calories"]["total_kcal"], result["calories"]["days_recorded"],
                          result["calories"]["lowest_kcal"]), (4000, 3, 0))
        self.assertEqual(result["protein"]["total_g"], 140.5)
        self.assertEqual(result["cardio"]["total_display"], "25:30")
        self.assertEqual(result["weight_lifted"]["total_kg"], 2450)
        self.assertEqual((result["measurements"]["weight_kg"]["current"], result["measurements"]["weight_kg"]["change"]),
                         (72, -0.5))
        self.assertEqual(len(result["daily"]), 31)
        self.assertEqual(len(result["weekly"]), 5)

    def test_it_is_the_calculation_applied_to_that_month_as_stored(self):
        store = st.TrackerStore(self.path)
        self.assertEqual(self.analytics(10), an.month_analytics(store.get_month(2026, 10), TODAY))

    def test_unreadable_cells_are_reported_and_not_counted(self):
        result = self.analytics(10)
        self.assertEqual(sorted((i["date"], i["field"]) for i in result["issues"]),
                         [("2026-10-06", "calories_kcal"), ("2026-10-06", "exercise")])
        day = result["daily"][5]
        self.assertEqual((day["exercise"], day["exercise_status"], day["calories_kcal"]), (None, "missed", None))

    def test_month_isolation(self):
        october = self.analytics(10)
        november = self.analytics(11)
        self.assertEqual(november["month"]["label"], "November 2026")
        self.assertEqual(november["exercise"]["counts"]["pending"], 30)          # the whole month is to come
        self.assertEqual((november["calories"]["total_kcal"], november["protein"]["total_g"],
                          november["weight_lifted"]["total_kg"]), (9999, 999, 99999))
        self.assertEqual(november["measurements"]["weight_kg"]["previous"], None)
        self.assertTrue(all(day["date"].startswith("2026-11-") for day in november["daily"]))
        self.assertTrue(all(day["date"].startswith("2026-10-") for day in october["daily"]))
        self.assertNotIn("9999", str(october))
        # Changing November changes nothing in October's analytics.
        wb = load_workbook(self.path)
        wb["November 2026"]["E3"], wb["November 2026"]["B3"], wb["November 2026"]["K2"] = 5000, "Missed", 60
        wb["September 2026"]["B4"] = "Completed"
        wb.save(self.path)
        self.assertEqual(self.analytics(10), october)

    def test_every_month_gets_the_one_set_of_analytics(self):
        keys = [list(self.analytics(number)) for number in (9, 10, 11)]
        self.assertEqual(keys[0], keys[1])
        self.assertEqual(keys[1], keys[2])
        september = self.analytics(9)
        self.assertEqual((september["exercise"]["counts"]["completed"], september["exercise"]["counts"]["missed"]), (2, 28))
        self.assertEqual((september["junk_food"]["counts"]["had"], september["cardio"]["total_display"],
                          september["weight_lifted"]["total_kg"]), (1, "10:00", 900))
        self.assertIsNone(september["calories"]["total_kcal"])

    def test_analytics_never_change_the_workbook(self):
        for _ in range(2):
            for number in (9, 10, 11):
                self.analytics(number)
        self.assertEqual(sha256(self.path), self.before)
        self.assertEqual([p.name for p in self.tmp.iterdir()], ["Fitness_Tracker.xlsx"])
        wb = load_workbook(self.path)
        self.assertEqual([wbk.is_month_sheet(ws) for ws in wb], [True, True, True])

    def test_missing_month_and_bad_month(self):
        missing = self.client.get("/api/months/2026/12/analytics")
        self.assertEqual((missing.status_code, missing.get_json()["error"]["code"]), (404, "month_not_found"))
        bad = self.client.get("/api/months/2026/13/analytics")
        self.assertEqual((bad.status_code, bad.get_json()["error"]["code"]), (400, "invalid_month"))

    def test_month_in_an_unknown_layout_is_an_error_not_a_guess(self):
        wb = load_workbook(self.path)
        wb["October 2026"]["C1"] = "Sweets"
        wb.save(self.path)
        response = self.client.get("/api/months/2026/10/analytics")
        self.assertEqual((response.status_code, response.get_json()["error"]["code"]), (500, "workbook_format"))
        self.assertEqual(self.analytics(11)["month"]["label"], "November 2026")  # the other months still work

    def test_analytics_follow_an_edit(self):
        saved = self.client.put("/api/months/2026/10/days/20", json={"exercise": "Completed", "calories_kcal": 2100})
        self.assertEqual(saved.status_code, 200)
        result = self.analytics(10)
        self.assertEqual(result["exercise"]["counts"],
                         {"completed": 2, "partial": 1, "rest_day": 1, "missed": 16, "pending": 11})
        self.assertEqual(result["exercise"]["current_streak"], 1)
        self.assertEqual(result["calories"]["total_kcal"], 6100)
        self.assertEqual(result["daily"][19]["junk_food_status"], "pending")    # today, still blank

    def test_no_new_routes(self):
        rules = sorted(f"{sorted(rule.methods - {'HEAD', 'OPTIONS'})[0]} {rule.rule}"
                       for rule in self.client.application.url_map.iter_rules() if rule.rule.startswith("/api/"))
        self.assertEqual(rules, [
            "GET /api/months",
            "GET /api/months/<int:year>/<int:month>",
            "GET /api/months/<int:year>/<int:month>/analytics",
            "GET /api/status",
            "POST /api/months",
            "PUT /api/months/<int:year>/<int:month>/days/<int:day>",
            "PUT /api/months/<int:year>/<int:month>/measurements/<int:day>",
        ])


if __name__ == "__main__":
    unittest.main()

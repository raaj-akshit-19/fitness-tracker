"""Checks for reading the saved workbook out of Excel while Excel has it open.

Three layers:

  Logic      Stand-ins for Excel's objects. Any attempt to change something
             through them fails the test. No Excel needed.
  Real Excel A private, hidden Excel started by the tests opens a temporary
             workbook. Skipped when Excel is not installed.
  Packaged   The built program is asked what it sees. Skipped until built.

Excel only refuses to let the file be read when it lives in a OneDrive folder.
These tests use ordinary temporary folders, so they make the file reader fail
on purpose to stand in for that. The real workbook is never used.

Run from the project root:  python -m unittest discover -s tests -v
"""

import ast
import datetime
import decimal
import hashlib
import json
import os
import shutil
import subprocess
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from backend import excel_live as xl
from backend import launch
from backend import store as st
from backend import workbook as wbk
from backend.analytics import month_analytics
from backend.app import create_app

PROJECT = Path(__file__).resolve().parent.parent
EXE = PROJECT / "release" / "Fitness Tracker" / "Fitness Tracker.exe"
SOURCE = (PROJECT / "backend" / "excel_live.py").read_text(encoding="utf-8")
NEEDS_READER = unittest.skipUnless(xl.AVAILABLE, "pywin32 is not installed")

# Members of Excel's objects that change, save, close or take control.
FORBIDDEN = {
    "Save", "SaveAs", "SaveCopyAs", "Close", "Quit", "Open", "Add", "Delete", "Clear",
    "ClearContents", "Activate", "Select", "Run", "Calculate", "Copy", "Cut", "Paste",
    "PasteSpecial", "Insert", "Protect", "Unprotect", "Visible", "DisplayAlerts",
    "AutoSaveOn", "ScreenUpdating", "EnableEvents", "SendKeys", "Undo", "Refresh",
}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def locked_file():
    """Make the file reader fail the way it does when Excel holds a OneDrive file."""
    return mock.patch.object(st, "load_workbook", side_effect=PermissionError("held by Excel"))


# ---------------------------------------------------------------- stand-ins for Excel

class ReadOnly:
    """Stands in for one of Excel's objects.

    Fails the test on any write, on anything that is not a plain property read,
    and on any name outside the reader's fixed list.
    """

    asked = []      # every name the reader requested, across all stand-ins

    def GetIDsOfNames(self, name):
        return name

    def Invoke(self, name, locale, flags, want_result, *args):
        ReadOnly.asked.append(name)
        if flags != xl.pythoncom.DISPATCH_PROPERTYGET:
            raise AssertionError(f"{name} was not asked for as a plain read")
        if args:
            raise AssertionError(f"the reader passed {args!r} to Excel")
        if name == xl.pythoncom.DISPID_NEWENUM:
            return FakeListing(self._items)         # stepping through a collection
        if name in FORBIDDEN or name not in xl.READS:
            raise AssertionError(f"the reader asked Excel for {name}")
        return getattr(self, name)

    def __setattr__(self, name, value):
        if name.startswith("_") or name in getattr(self, "_settable", ()):
            object.__setattr__(self, name, value)
        else:
            raise AssertionError(f"the reader tried to change {name} in Excel")

    def __getattr__(self, name):
        if name in FORBIDDEN:
            raise AssertionError(f"the reader used {name} on Excel")
        raise AttributeError(name)


class FakeRange(ReadOnly):
    def __init__(self, ref, values, formulas, book):
        self._ref, self._values, self._formulas, self._book = ref, values, formulas, book

    @property
    def Address(self):
        col_row = self._ref.split(":")
        return ":".join("$" + part[0] + "$" + part[1:] for part in col_row)

    @property
    def Value(self):
        self._book._cell_reads += 1
        if self._book._edit_during_read:
            object.__setattr__(self._book, "Saved", False)
        return self._values

    @property
    def Formula(self):
        return self._formulas


class FakeListing:
    """What Excel hands over for stepping through a collection."""

    def __init__(self, items):
        self._left = list(items)

    def QueryInterface(self, interface):
        return self

    def Next(self, count=1):
        return (self._left.pop(0),) if self._left else ()


class FakeCollection(ReadOnly):
    def __init__(self, items):
        self._items = items

    @property
    def Count(self):
        return len(self._items)

    def Item(self, key):
        if isinstance(key, int):
            return self._items[key - 1]
        return next(item for item in self._items if item.Name == key)


class FakeTable(ReadOnly):
    def __init__(self, name, area):
        self._name, self._area = name, area

    Name = property(lambda self: self._name)
    Range = property(lambda self: self._area)


class FakeSheet(ReadOnly):
    def __init__(self, name, tables):
        self._name, self._tables = name, tables

    Name = property(lambda self: self._name)
    ListObjects = property(lambda self: FakeCollection(self._tables))


class FakeBook(ReadOnly):
    """What Excel would hold in memory for the workbook file at source."""

    _settable = ("Saved", "FullName")

    def __init__(self, source, full_name=None, saved=True):
        self._cell_reads = 0
        self._edit_during_read = False
        self._busy = False
        self._sheets = []
        self.FullName = str(full_name or source)
        self.Saved = saved
        wb = load_workbook(source)
        for ws in wb.worksheets:
            tables = []
            for name in ws.tables:
                ref = ws.tables[name].ref
                rows = [[cell.value for cell in row] for row in ws[ref]]
                values = tuple(tuple(as_excel(v) for v in row) for row in rows)
                formulas = tuple(tuple("" if v is None else str(v) for v in row) for row in rows)
                tables.append(FakeTable(name, FakeRange(ref, values, formulas, self)))
            self._sheets.append(FakeSheet(ws.title, tables))

    def __getattribute__(self, name):
        if name == "Saved" and object.__getattribute__(self, "_busy"):
            raise xl.pythoncom.com_error(-2147418111, "Call was rejected by callee.", None, None)
        return object.__getattribute__(self, name)

    Name = property(lambda self: os.path.basename(self.FullName))
    Worksheets = property(lambda self: FakeCollection(self._sheets))


def as_excel(value):
    """A value as Excel hands it over: every number a float, dates with a time zone."""
    if isinstance(value, bool) or value is None or isinstance(value, str):
        return value
    if isinstance(value, datetime.datetime):
        return value.replace(tzinfo=datetime.timezone.utc)
    if isinstance(value, datetime.date):
        return datetime.datetime(value.year, value.month, value.day, tzinfo=datetime.timezone.utc)
    return float(value)


class FakeExcel:
    """The documents registered as running, across every Excel that is open."""

    def __init__(self):
        self.documents = []     # (name it is registered under, object)

    def open(self, book, name=None):
        self.documents.append((name or book.FullName, book))
        return book

    def close(self, book):
        self.documents = [(n, b) for n, b in self.documents if b is not book]

    def running(self):
        for name, book in list(self.documents):
            yield name, lambda book=book: book


@NEEDS_READER
class LogicCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.path = wbk.create_workbook(self.tmp / "Fitness_Tracker.xlsx")
        self.store = st.TrackerStore(self.path)
        self.excel = FakeExcel()
        ReadOnly.asked.clear()
        patcher = mock.patch.object(xl, "_running_documents", side_effect=self.excel.running)
        self.running = patcher.start()
        self.addCleanup(patcher.stop)

    def write(self, cells, path=None, sheet="October 2026"):
        """Save new values into the workbook file (what Excel saving does)."""
        path = path or self.path
        wb = load_workbook(path)
        for ref, value in cells.items():
            wb[sheet][ref] = value
        wb.save(path)

    def in_excel(self, cells=None, saved=True, path=None):
        """A copy of the workbook as Excel holds it, optionally with other values."""
        path = path or self.path
        scratch = self.tmp / f"in-excel-{len(list(self.tmp.iterdir()))}.xlsx"
        shutil.copy(path, scratch)
        if cells:
            self.write(cells, scratch)
        return FakeBook(scratch, full_name=path, saved=saved)

    def day(self, number, store=None):
        return (store or self.store).get_month(2026, 10)["days"][number - 1]


class FallbackTests(LogicCase):
    """1, 11, 15: the file reader stays the normal path and the fallback."""

    def test_closed_workbook_is_read_from_the_file_without_asking_excel(self):
        self.write({"B2": "Completed", "D2": "25:30", "G2": 2450})
        with mock.patch.object(xl, "inspect", side_effect=AssertionError("Excel was asked")):
            self.assertEqual(self.day(1)["cardio_display"], "25:30")
            self.assertEqual(self.store.list_months(), [{"year": 2026, "month": 10, "label": "October 2026"}])
            self.assertEqual(self.store.status()["unsaved_changes"], False)
            self.assertIs(self.store.status()["locked"], False)
        self.running.assert_not_called()

    def test_readable_file_is_used_even_when_excel_has_it_open(self):
        self.write({"G2": 100})
        self.excel.open(self.in_excel({"G2": 999}, saved=False))
        self.assertEqual(self.day(1)["weight_lifted_kg"], 100)
        self.running.assert_not_called()

    def test_locked_and_not_open_in_excel_keeps_the_lock_error(self):
        with locked_file():
            with self.assertRaises(st.WorkbookLockedError) as caught:
                self.store.get_month(2026, 10)
        self.assertEqual(str(caught.exception),
                         "The workbook could not be read because another program has it locked.")

    def test_without_the_excel_reader_behaviour_is_as_before(self):
        self.excel.open(self.in_excel())
        with locked_file(), mock.patch.object(xl, "AVAILABLE", False):
            self.assertEqual(xl.inspect(self.path).state, xl.UNAVAILABLE)
            with self.assertRaises(st.WorkbookLockedError) as caught:
                self.store.list_months()
        self.assertIn("another program has it locked", str(caught.exception))
        self.running.assert_not_called()

    def test_unreadable_file_is_still_reported_as_unreadable(self):
        self.path.write_bytes(b"not a workbook")
        self.excel.open(FakeBook(wbk.create_workbook(self.tmp / "other.xlsx"), full_name=self.path))
        with self.assertRaises(st.WorkbookUnreadableError):
            self.store.list_months()
        self.running.assert_not_called()


class SavedStateTests(LogicCase):
    """2, 4, 5, 6: what is read, and when."""

    def test_saved_workbook_open_in_excel_is_read_through_excel(self):
        book = self.excel.open(self.in_excel({"B2": "Completed", "C2": "Had", "D2": "25:30", "G2": 2450}))
        with locked_file():
            self.assertEqual(xl.inspect(self.path, read=False).state, xl.SAVED)
            day = self.day(1)
            self.assertEqual(self.store.list_months(), [{"year": 2026, "month": 10, "label": "October 2026"}])
        self.assertEqual(day, {
            "date": "2026-10-01", "exercise": "Completed", "junk_food": "Had",
            "cardio_display": "25:30", "cardio_seconds": 1530,
            "calories_kcal": None, "protein_g": None, "weight_lifted_kg": 2450,
        })
        self.assertGreater(book._cell_reads, 0)
        self.assertLessEqual(set(ReadOnly.asked), set(xl.READS) | {xl.pythoncom.DISPID_NEWENUM})

    def test_result_is_identical_to_reading_the_file(self):
        cells = {
            "B2": "Completed", "C2": "Had", "D2": "25:30", "G2": 2450, "B3": "Missed", "D3": "0:45",
            "G3": 1820.5, "G4": 0, "D5": "abc", "G6": -5, "B7": "yes",
            "K2": 72.5, "L2": 84, "M3": 101.25, "P5": 28, "K4": 0,
        }
        self.write(cells)
        from_file = self.store.get_month(2026, 10)
        self.excel.open(self.in_excel())
        with locked_file():
            fresh = st.TrackerStore(self.path)
            from_excel = fresh.get_month(2026, 10)
        self.assertEqual(from_excel, from_file)
        self.assertEqual(json.dumps(from_excel), json.dumps(from_file))   # same types too
        today = datetime.date(2026, 10, 20)
        self.assertEqual(month_analytics(from_excel, today), month_analytics(from_file, today))
        self.assertEqual(len(from_excel["issues"]), 5)

    def test_unsaved_changes_are_never_shown(self):
        self.write({"G2": 100})
        self.assertEqual(self.day(1)["weight_lifted_kg"], 100)     # read before Excel took it
        book = self.excel.open(self.in_excel({"G2": 999, "B2": "Completed"}, saved=False))
        with locked_file():
            self.assertEqual(xl.inspect(self.path).state, xl.UNSAVED)
            self.assertIsNone(xl.inspect(self.path).workbook)
            for _ in range(3):
                day = self.day(1)
                self.assertEqual(day["weight_lifted_kg"], 100)
                self.assertIsNone(day["exercise"])
        self.assertEqual(book._cell_reads, 0)       # the unsaved cells were not even looked at

    def test_unsaved_with_nothing_loaded_yet_says_to_save(self):
        self.excel.open(self.in_excel({"G2": 999}, saved=False))
        with locked_file():
            with self.assertRaises(st.WorkbookLockedError) as caught:
                self.store.get_month(2026, 10)
        self.assertIn("changes that have not been saved", str(caught.exception))
        self.assertIn("Save it in Excel (Ctrl+S)", str(caught.exception))
        self.assertNotIn("999", str(caught.exception))

    def test_change_appears_once_it_is_saved(self):
        self.write({"G2": 100})
        self.assertEqual(self.day(1)["weight_lifted_kg"], 100)
        book = self.excel.open(self.in_excel({"G2": 1250}, saved=False))
        with locked_file():
            self.assertEqual(self.day(1)["weight_lifted_kg"], 100)     # typed, not saved
            book.Saved = True                                             # Ctrl+S ...
            os.utime(self.path, ns=(time.time_ns(), time.time_ns()))      # ... rewrites the file
            self.assertEqual(self.day(1)["weight_lifted_kg"], 1250)
            book2 = self.in_excel({"G2": 1300}, saved=False)              # typed again, not saved
            self.excel.close(book)
            self.excel.open(book2)
            self.assertEqual(self.day(1)["weight_lifted_kg"], 1250)    # still the saved value

    def test_cells_are_not_read_again_until_something_is_saved(self):
        book = self.excel.open(self.in_excel({"G2": 1250}))
        with locked_file():
            self.day(1)
            reads = book._cell_reads
            for _ in range(5):
                self.day(1)
                self.store.list_months()
            self.assertEqual(book._cell_reads, reads)
            os.utime(self.path, ns=(time.time_ns(), time.time_ns()))
            self.day(1)
            self.assertGreater(book._cell_reads, reads)

    def test_a_change_typed_while_reading_discards_what_was_read(self):
        book = self.excel.open(self.in_excel({"G2": 999}))
        book._edit_during_read = True
        with locked_file():
            found = xl.inspect(self.path)
            self.assertEqual((found.state, found.workbook), (xl.UNSAVED, None))
            with self.assertRaises(st.WorkbookLockedError):
                self.store.get_month(2026, 10)

    def test_excel_busy_keeps_the_last_saved_state_and_then_recovers(self):
        self.write({"G2": 100})
        self.assertEqual(self.day(1)["weight_lifted_kg"], 100)
        book = self.excel.open(self.in_excel({"G2": 1250}))
        book._busy = True                               # a cell is being edited
        with locked_file():
            self.assertEqual(xl.inspect(self.path).state, xl.BUSY)
            self.assertEqual(self.day(1)["weight_lifted_kg"], 100)
            fresh = st.TrackerStore(self.path)          # nothing loaded yet: explain, do not guess
            with self.assertRaises(st.WorkbookLockedError) as caught:
                fresh.get_month(2026, 10)
            self.assertIn("Excel is busy", str(caught.exception))
            self.assertIn("try again", str(caught.exception))
            book._busy = False
            os.utime(self.path, ns=(time.time_ns(), time.time_ns()))
            self.assertEqual(self.day(1)["weight_lifted_kg"], 1250)
            self.assertEqual(self.day(1, fresh)["weight_lifted_kg"], 1250)

    def test_excel_not_answering_at_all_is_busy_not_closed(self):
        # While a cell is being typed in, Excel will not even say which file it has.
        self.write({"G2": 100})
        self.assertEqual(self.day(1)["weight_lifted_kg"], 100)

        def rejected():
            raise xl.pythoncom.com_error(-2147418111, "Call was rejected by callee.", None, None)

        self.excel.documents.append((str(self.path), None))
        self.running.side_effect = lambda: iter([(str(self.path), rejected)])
        with locked_file():
            self.assertEqual(xl.inspect(self.path, read=False).state, xl.BUSY)
            self.assertEqual(self.day(1)["weight_lifted_kg"], 100)       # last saved state kept
            with self.assertRaises(st.WorkbookLockedError) as caught:
                st.TrackerStore(self.path).get_month(2026, 10)
            self.assertIn("Excel is busy", str(caught.exception))
            # pywin32 can also report a refused question as a missing attribute.
            def refused():
                raise AttributeError("<unknown>.FullName")

            self.running.side_effect = lambda: iter([(str(self.path), refused)])
            self.assertEqual(xl.inspect(self.path, read=False).state, xl.BUSY)
            self.assertEqual(self.day(1)["weight_lifted_kg"], 100)
            # A different workbook that will not answer says nothing about this one.
            other = self.tmp / "elsewhere"
            other.mkdir()
            self.running.side_effect = lambda: iter([(str(other / "Fitness_Tracker.xlsx"), rejected)])
            self.assertEqual(xl.inspect(self.path, read=False).state, xl.NOT_OPEN)

    def test_new_month_sheet_saved_in_excel_is_listed(self):
        scratch = self.tmp / "with-november.xlsx"
        shutil.copy(self.path, scratch)
        st.TrackerStore(scratch).create_month(2026, 11)
        self.excel.open(FakeBook(scratch, full_name=self.path))
        with locked_file():
            self.assertEqual([m["label"] for m in self.store.list_months()],
                             ["October 2026", "November 2026"])
            self.assertEqual(len(self.store.get_month(2026, 11)["days"]), 30)
            with self.assertRaises(st.MonthNotFoundError):
                self.store.get_month(2026, 12)


class WorkbookSelectionTests(LogicCase):
    """3: the workbook is picked by its full path, never by position."""

    def test_same_file_name_in_other_folders_is_not_mistaken_for_it(self):
        for index, value in enumerate((111, 222), start=1):
            folder = self.tmp / f"elsewhere {index}"
            folder.mkdir()
            other = wbk.create_workbook(folder / "Fitness_Tracker.xlsx")
            self.write({"G2": value}, other)
            self.excel.open(FakeBook(other))                    # registered first
        self.excel.open(self.in_excel({"G2": 333}))
        with locked_file():
            self.assertEqual(self.day(1)["weight_lifted_kg"], 333)

    def test_only_other_workbooks_open_means_not_open(self):
        folder = self.tmp / "elsewhere"
        folder.mkdir()
        self.excel.open(FakeBook(wbk.create_workbook(folder / "Fitness_Tracker.xlsx")))
        self.excel.open(FakeBook(wbk.create_workbook(self.tmp / "Budget.xlsx")))
        self.assertEqual(xl.inspect(self.path).state, xl.NOT_OPEN)
        with locked_file():
            with self.assertRaises(st.WorkbookLockedError) as caught:
                self.store.list_months()
        self.assertIn("another program has it locked", str(caught.exception))

    def test_no_excel_running_means_not_open(self):
        self.assertEqual(xl.inspect(self.path).state, xl.NOT_OPEN)

    def test_path_is_compared_without_regard_to_case_or_spelling(self):
        book = self.in_excel({"G2": 5})
        book.FullName = str(self.path).upper()
        self.excel.open(book, name=str(self.path.parent / "." / "FITNESS_TRACKER.XLSX"))
        with locked_file():
            self.assertEqual(self.day(1)["weight_lifted_kg"], 5)

    def test_registered_name_and_workbook_must_both_be_this_file(self):
        impostor = FakeBook(wbk.create_workbook(self.tmp / "Other.xlsx"))   # says it is another file
        self.excel.open(impostor, name=str(self.path))
        self.assertEqual(xl.inspect(self.path).state, xl.NOT_OPEN)

    def test_cloud_address_of_a_onedrive_file_is_matched_to_the_local_path(self):
        mounts = [("https://d.docs.live.net", str(self.tmp))]
        target = xl._normal(self.path)
        cloud = "https://d.docs.live.net/0123456789abcdef/Fitness_Tracker.xlsx"
        self.assertTrue(xl._is_workbook(cloud, target, mounts))
        self.assertTrue(xl._is_workbook(cloud.upper().replace("HTTPS://D.DOCS.LIVE.NET", "https://d.docs.live.net"),
                                        target, mounts))
        # A different folder, a different file, a different service: not this workbook.
        for other in (
            "https://d.docs.live.net/0123456789abcdef/Elsewhere/Fitness_Tracker.xlsx",
            "https://d.docs.live.net/0123456789abcdef/Fitness_Tracker (1).xlsx",
            "https://example.sharepoint.com/sites/x/Fitness_Tracker.xlsx",
        ):
            self.assertFalse(xl._is_workbook(other, target, mounts), other)

    def test_cloud_address_with_folders_spaces_and_escapes(self):
        nested = self.tmp / "Desktop" / "gym tracker (2) & more"
        nested.mkdir(parents=True)
        path = wbk.create_workbook(nested / "Fitness_Tracker.xlsx")
        target = xl._normal(path)
        personal = [("https://d.docs.live.net", str(self.tmp))]
        business = [("https://contoso-my.sharepoint.com/personal/me_contoso_com/Documents/", str(self.tmp))]
        self.assertTrue(xl._is_workbook(
            "https://d.docs.live.net/0123abc/Desktop/gym tracker (2) & more/Fitness_Tracker.xlsx", target, personal))
        self.assertTrue(xl._is_workbook(
            "https://d.docs.live.net/0123abc/Desktop/gym%20tracker%20(2)%20&%20more/Fitness_Tracker.xlsx",
            target, personal))
        self.assertTrue(xl._is_workbook(
            "https://contoso-my.sharepoint.com/personal/me_contoso_com/Documents/Desktop/"
            "gym tracker (2) & more/Fitness_Tracker.xlsx", target, business))
        self.assertFalse(xl._is_workbook(
            "https://d.docs.live.net/0123abc/Desktop/gym tracker/Fitness_Tracker.xlsx", target, personal))

    def test_workbook_opened_through_its_cloud_address_is_read(self):
        cloud = "https://d.docs.live.net/0123456789abcdef/Fitness_Tracker.xlsx"
        book = self.in_excel({"G2": 777})
        book.FullName = cloud
        self.excel.open(book)
        with locked_file(), mock.patch.object(
                xl, "_onedrive_mounts", return_value=[("https://d.docs.live.net", str(self.tmp))]):
            self.assertEqual(self.day(1)["weight_lifted_kg"], 777)

    def test_onedrive_folders_come_from_this_computer_not_from_the_code(self):
        for prefix, folder in xl._onedrive_mounts():
            self.assertTrue(prefix.startswith("https://"))
            self.assertTrue(os.path.isabs(folder))
        self.assertNotIn(Path.home().name, SOURCE)
        self.assertNotIn("Users", SOURCE)


class ReadOnlyTests(LogicCase):
    """7, 8, 9: nothing is changed, saved or closed."""

    def test_reading_changes_nothing_in_excel(self):
        book = self.excel.open(self.in_excel({"B2": "Completed", "G2": 2450, "K2": 70}))
        with locked_file():
            for _ in range(3):
                self.store.status()
                self.store.list_months()
                self.store.get_month(2026, 10)
        # Any write or forbidden member would have raised inside the stand-ins.
        self.assertIs(book.Saved, True)
        self.assertEqual(self.excel.documents[0][1], book)      # still open

    def test_the_stand_ins_really_do_catch_writes(self):
        book = self.in_excel()
        with self.assertRaises(AssertionError):
            book.Worksheets.Item(1).ListObjects.Item(1).Range.Value = 5
        get, put = xl.pythoncom.DISPATCH_PROPERTYGET, xl.pythoncom.DISPATCH_PROPERTYPUT
        for name in ("Save", "Close", "Quit", "Application"):
            with self.assertRaises(AssertionError):
                book.Invoke(name, 0, get, True)
        with self.assertRaises(AssertionError):
            book.Invoke("Saved", 0, put, False, True)
        with self.assertRaises(AssertionError):
            book.Invoke("Saved", 0, get | xl.pythoncom.DISPATCH_METHOD, True)
        with self.assertRaises(AssertionError):
            book.Invoke("Name", 0, get, True, "an argument")

    def test_only_listed_names_can_be_asked_for(self):
        self.assertEqual(xl.READS, {"FullName", "Saved", "Worksheets", "ListObjects",
                                    "Name", "Range", "Address", "Value", "Formula"})
        self.assertEqual(xl.READS & FORBIDDEN, set())

        class Trap:
            def __getattr__(self, name):
                raise AssertionError("Excel was asked")

        for name in sorted(FORBIDDEN) + ["Application", "Workbooks", "Sheets", "Cells"]:
            with self.assertRaises(ValueError):
                xl._get(Trap(), name)           # refused before Excel is touched

    def test_source_only_reads(self):
        tree = ast.parse(SOURCE)
        used = {node.attr for node in ast.walk(tree) if isinstance(node, ast.Attribute)}
        self.assertEqual(used & FORBIDDEN, set())
        # Excel is only ever asked through _get, always for a name written out in full,
        # and _get only ever makes a plain property read.
        asked = [node for node in ast.walk(tree)
                 if isinstance(node, ast.Call) and getattr(node.func, "id", "") == "_get"]
        self.assertGreater(len(asked), 10)
        for call in asked:
            self.assertEqual(len(call.args), 2)                 # the object and a name: nothing is passed in
            self.assertIsInstance(call.args[1], ast.Constant)
            self.assertIn(call.args[1].value, xl.READS)
        # The two requests ever made of Excel: read a named property, and step through a collection.
        invokes = [node for node in ast.walk(tree)
                   if isinstance(node, ast.Call) and getattr(node.func, "attr", "") == "Invoke"]
        self.assertEqual(len(invokes), 2)
        for call in invokes:
            self.assertEqual(len(call.args), 4)                 # which, locale, kind, want the answer
            self.assertEqual(ast.unparse(call.args[2]), "pythoncom.DISPATCH_PROPERTYGET")
        for phrase in ("DISPATCH_PROPERTYPUT", "DISPATCH_METHOD", "win32com", "InvokeTypes"):
            self.assertNotIn(phrase, SOURCE)
        strings = {node.value for node in ast.walk(tree)
                   if isinstance(node, ast.Constant) and isinstance(node.value, str)}
        self.assertEqual(strings & FORBIDDEN, set())
        # Excel's own members start with a capital letter. None is ever assigned to.
        for node in ast.walk(tree):
            targets = []
            if isinstance(node, ast.Assign):
                targets = node.targets
            elif isinstance(node, (ast.AugAssign, ast.AnnAssign)):
                targets = [node.target]
            for target in targets:
                for part in (target.elts if isinstance(target, ast.Tuple) else [target]):
                    if isinstance(part, ast.Attribute):
                        self.assertFalse(part.attr[:1].isupper(), f"assigns to .{part.attr}")
        self.assertEqual([node for node in ast.walk(tree) if isinstance(node, ast.Delete)], [])
        # It never starts Excel or opens a file: only objects already running.
        for phrase in ("Excel.Application", "DispatchEx", "EnsureDispatch", "client.GetObject",
                       "CoCreateInstance", "startfile", "subprocess", "Workbooks"):
            self.assertNotIn(phrase, SOURCE)
        self.assertEqual(SOURCE.count("GetObject("), 1)        # the look-up of a running document
        self.assertIn("table.GetObject(moniker)", SOURCE)

    def test_reader_does_not_load_the_scripting_helper_library(self):
        # That library writes a cache folder into the temp directory when loaded.
        import sys
        code = ("import sys, tempfile, os; from backend import excel_live, store; "
                "found = excel_live.inspect(os.path.join(tempfile.gettempdir(), 'not-open.xlsx')); "
                "print(found.state, 'win32com' in sys.modules)")
        done = subprocess.run([sys.executable, "-c", code], cwd=PROJECT, capture_output=True, text=True,
                              env=dict(os.environ, PYTHONDONTWRITEBYTECODE="1"))
        self.assertEqual(done.stdout.strip(), "not_open False", done.stderr)

    def test_the_store_only_asks_excel_when_reading(self):
        text = (PROJECT / "backend" / "store.py").read_text(encoding="utf-8")
        create = text[text.index("def create_month"):text.index("def _load")]
        self.assertNotIn("excel_live", create)
        self.assertNotIn("_read(", create)
        self.assertIn("self._ensure_writable()", create)


class WriteStaysBlockedTests(LogicCase):
    """10: creating a month needs the file itself."""

    def test_month_creation_is_blocked_while_excel_holds_the_workbook(self):
        book = self.excel.open(self.in_excel())
        before = sha256(self.path)
        real_open = open

        def held(file, mode="r", *args, **kwargs):
            if "+" in mode and isinstance(file, (str, os.PathLike)) and Path(file) == self.path:
                raise PermissionError("held by Excel")
            return real_open(file, mode, *args, **kwargs)

        with locked_file(), mock.patch("builtins.open", side_effect=held):
            self.assertEqual(len(self.store.get_month(2026, 10)["days"]), 31)    # reading works
            with self.assertRaises(st.WorkbookLockedError) as caught:
                self.store.create_month(2026, 11)
        self.assertEqual(str(caught.exception),
                         "The workbook is open in Excel or locked by another program. "
                         "Close it and try again. Nothing was changed.")
        self.assertEqual(sha256(self.path), before)
        self.assertIs(book.Saved, True)
        self.assertEqual([p.name for p in self.tmp.iterdir() if p.name.startswith(".tracker-")], [])
        # With Excel closed it works again.
        self.excel.close(book)
        self.assertEqual(self.store.create_month(2026, 11)["label"], "November 2026")


class StatusAndApiTests(LogicCase):
    def held(self):
        real_open = open

        def refuse(file, mode="r", *args, **kwargs):
            if "+" in mode and isinstance(file, (str, os.PathLike)) and Path(file) == self.path:
                raise PermissionError("held by Excel")
            return real_open(file, mode, *args, **kwargs)

        return mock.patch("builtins.open", side_effect=refuse)

    def test_status_reports_unsaved_changes_only_when_excel_has_them(self):
        self.assertEqual(self.store.status()["unsaved_changes"], False)
        book = self.excel.open(self.in_excel(saved=False))
        with self.held():
            status = self.store.status()
            self.assertEqual((status["locked"], status["unsaved_changes"]), (True, True))
            book.Saved = True
            status = self.store.status()
            self.assertEqual((status["locked"], status["unsaved_changes"]), (True, False))
            self.excel.close(book)                  # locked by something that is not Excel
            self.assertEqual(self.store.status()["unsaved_changes"], False)
        self.assertEqual(set(status), {"workbook_modified", "locked", "unsaved_changes"})

    def test_api_answers_normally_while_excel_holds_a_saved_workbook(self):
        self.excel.open(self.in_excel({"B2": "Completed", "D2": "15:30", "G2": 1250, "K2": 72.5}))
        client = create_app(self.path, today=lambda: datetime.date(2026, 10, 3)).test_client()
        with locked_file(), self.held():
            self.assertEqual(client.get("/api/status").get_json()["locked"], True)
            self.assertEqual(client.get("/api/months").status_code, 200)
            month = client.get("/api/months/2026/10")
            analytics = client.get("/api/months/2026/10/analytics")
            blocked = client.post("/api/months", json={"year": 2026, "month": 11})
        self.assertEqual(month.status_code, 200)
        self.assertEqual(set(month.get_json()), {"year", "month", "label", "days", "measurements", "issues"})
        self.assertEqual(month.get_json()["days"][0]["cardio_display"], "15:30")
        self.assertEqual(analytics.get_json()["cardio"]["total_display"], "15:30")
        self.assertEqual(analytics.get_json()["weight_lifted"]["total_kg"], 1250)
        self.assertEqual(analytics.get_json()["measurements"]["weight_kg"]["current"], 72.5)
        self.assertEqual(analytics.get_json()["exercise"]["counts"]["completed"], 1)
        self.assertEqual((blocked.status_code, blocked.get_json()["error"]["code"]), (423, "workbook_locked"))

    def test_api_explains_unsaved_changes_when_there_is_nothing_to_show(self):
        self.excel.open(self.in_excel({"G2": 999}, saved=False))
        client = create_app(self.path).test_client()
        with locked_file(), self.held():
            response = client.get("/api/months/2026/10")
            status = client.get("/api/status").get_json()
        self.assertEqual(response.status_code, 423)
        self.assertEqual(response.get_json()["error"]["code"], "workbook_locked")
        self.assertIn("Save it in Excel", response.get_json()["error"]["message"])
        self.assertNotIn("999", response.get_data(as_text=True))
        self.assertIs(status["unsaved_changes"], True)

    def test_startup_check_does_not_refuse_a_workbook_that_is_only_locked(self):
        lines = []
        with locked_file(), mock.patch.object(launch, "say", side_effect=lambda text="": lines.append(text)):
            launch.check_workbook_opens(self.path)                  # no error
        self.assertTrue(any(line.startswith("Note: ") and "locked" in line for line in lines))
        self.path.write_bytes(b"damaged")
        with self.assertRaises(launch.LaunchError):
            launch.check_workbook_opens(self.path)


class ValueTests(unittest.TestCase):
    def test_values_come_out_as_the_file_reader_gives_them(self):
        aware = datetime.datetime(2026, 10, 1, tzinfo=datetime.timezone.utc)
        cases = [
            ((None, ""), None), ((True, "TRUE"), True), ((False, "FALSE"), False),
            ((2450.0, "2450"), 2450), ((1820.5, "1820.5"), 1820.5), ((0.0, "0"), 0),
            (("25:30", "25:30"), "25:30"), (("", ""), None),
            ((aware, "46296"), datetime.datetime(2026, 10, 1)),
            ((decimal.Decimal("12.50"), "12.5"), 12.5), ((decimal.Decimal("3"), "3"), 3),
            ((7.0, "=3+4"), "=3+4"), ((-2146826281, "=1/0"), "=1/0"),
            ((-2146826281, "#DIV/0!"), "#DIV/0!"), ((-2146826246, "#N/A"), "#N/A"),
        ]
        for (value, formula), expected in cases:
            result = xl._plain(value, formula)
            self.assertEqual(result, expected, (value, formula))
            self.assertIs(type(result), type(expected), (value, formula))
        self.assertIsNone(xl._plain(aware, "46296").tzinfo)

    def test_one_cell_and_many_cells(self):
        self.assertEqual(xl._grid(5.0), ((5.0,),))
        self.assertEqual(xl._grid(((1.0, 2.0), (3.0, 4.0))), ((1.0, 2.0), (3.0, 4.0)))

    def test_snapshot_reads_like_a_workbook(self):
        sheet = xl.SheetSnapshot("October 2026", {"Daily_2026_10": "A1:B3"},
                                 {(1, 1): "Date", (1, 2): "Exercise", (2, 2): True})
        book = xl.WorkbookSnapshot(["Notes", "October 2026"], {"October 2026": sheet})
        self.assertEqual(book.sheetnames, ["Notes", "October 2026"])
        self.assertEqual(book["October 2026"].tables["Daily_2026_10"].ref, "A1:B3")
        self.assertEqual(list(sheet.iter_rows(1, 3, 1, 2, values_only=True)),
                         [("Date", "Exercise"), (None, True), (None, None)])
        self.assertEqual(wbk.list_months(book), [(2026, 10)])


@NEEDS_READER
class WorkerTests(unittest.TestCase):
    """12: one thread talks to Excel, with a time limit, and holds nothing afterwards."""

    def test_excel_not_answering_gives_busy_and_does_not_pile_up(self):
        release = threading.Event()

        def stuck(path, read):
            release.wait(20)
            return xl.Inspection(xl.NOT_OPEN)

        with mock.patch.object(xl, "_inspect", side_effect=stuck), \
                mock.patch.object(xl, "STATE_TIMEOUT", 0.3):
            started = time.monotonic()
            first = xl.inspect("x.xlsx", read=False)
            second = xl.inspect("x.xlsx", read=False)
            third = xl.inspect("x.xlsx", read=False)
            elapsed = time.monotonic() - started
            release.set()
        self.assertEqual([first.state, second.state, third.state], [xl.BUSY] * 3)
        self.assertLess(elapsed, 6)
        self.assertEqual(len([t for t in threading.enumerate() if t.name == "excel-reader"]), 1)
        # Once Excel answers again, so does the reader.
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            if xl.inspect(Path(tempfile.gettempdir()) / "nothing-here.xlsx", read=False).state == xl.NOT_OPEN:
                break
            time.sleep(0.2)
        else:
            self.fail("the reader did not recover")

    def test_an_unexpected_failure_is_busy_not_a_crash(self):
        with mock.patch.object(xl, "_inspect", side_effect=RuntimeError("boom")):
            found = xl.inspect("x.xlsx")
        self.assertEqual(found.state, xl.BUSY)
        self.assertIn("boom", found.detail)
        self.assertEqual(xl.inspect(Path(tempfile.gettempdir()) / "nothing-here.xlsx").state, xl.NOT_OPEN)

    def test_reader_thread_is_in_the_background(self):
        xl.inspect(Path(tempfile.gettempdir()) / "nothing-here.xlsx")
        threads = [t for t in threading.enumerate() if t.name == "excel-reader"]
        self.assertEqual(len(threads), 1)
        self.assertTrue(threads[0].daemon)      # never keeps the program from closing


# ---------------------------------------------------------------- real Excel

def excel_installed():
    if not xl.AVAILABLE:
        return False
    try:
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Excel.Application"))
        return True
    except OSError:
        return False


def process_running(pid):
    out = subprocess.run(["tasklist", "/FI", f"PID eq {pid}", "/FO", "CSV", "/NH"],
                         capture_output=True, text=True).stdout
    return str(pid) in out


@unittest.skipUnless(excel_installed(), "Microsoft Excel is not installed")
class RealExcelTests(unittest.TestCase):
    """A private, hidden Excel holds a temporary workbook open, as a user would."""

    @classmethod
    def setUpClass(cls):
        import ctypes
        from win32com.client import DispatchEx
        cls.tmp = Path(tempfile.mkdtemp(prefix="ft-excel-"))
        cls.apps = []

        def start():
            app = DispatchEx("Excel.Application")       # a separate Excel, never the user's
            app.Visible = False
            app.DisplayAlerts = False
            pid = ctypes.c_ulong()
            ctypes.windll.user32.GetWindowThreadProcessId(int(app.Hwnd), ctypes.byref(pid))
            cls.apps.append((app, pid.value))
            return app

        cls.start_excel = staticmethod(start)
        cls.app = start()

    @classmethod
    def tearDownClass(cls):
        pids = [pid for _, pid in cls.apps]

        def quit_all():
            for app, _ in cls.apps:
                try:
                    for index in range(int(app.Workbooks.Count), 0, -1):
                        app.Workbooks.Item(index).Close(False)
                    app.Quit()
                except Exception:  # noqa: BLE001
                    pass

        quit_all()      # in its own function so that nothing here keeps hold of Excel
        cls.apps.clear()
        cls.app = None
        import gc
        gc.collect()
        deadline = time.monotonic() + 20
        while time.monotonic() < deadline and any(process_running(pid) for pid in pids):
            time.sleep(0.5)
        cls.still_running = [pid for pid in pids if process_running(pid)]
        shutil.rmtree(cls.tmp, ignore_errors=True)
        # If the reader had kept hold of anything of Excel's, Excel could not have closed.
        assert not cls.still_running, f"Excel did not close: {cls.still_running}"

    def setUp(self):
        self.folder = self.tmp / self.id().split(".")[-1]
        self.folder.mkdir()
        self.path = wbk.create_workbook(self.folder / "Fitness_Tracker.xlsx")
        self.store = st.TrackerStore(self.path)

    def open(self, path=None, app=None):
        book = (app or self.app).Workbooks.Open(str(path or self.path))
        self.addCleanup(self.close, book)
        return book

    @staticmethod
    def close(book):
        try:
            book.Close(False)
        except Exception:  # noqa: BLE001
            pass

    @staticmethod
    def type_in(book, cells, sheet="October 2026"):
        """What a person typing does: the cell changes, the workbook is not saved."""
        ws = book.Worksheets(sheet)
        for ref, value in cells.items():
            if value is None:
                ws.Range(ref).ClearContents()
            else:
                ws.Range(ref).Value = value

    def test_open_workbook_is_detected_and_others_are_not(self):
        self.assertEqual(xl.inspect(self.path, read=False).state, xl.NOT_OPEN)
        self.open()
        self.assertEqual(xl.inspect(self.path, read=False).state, xl.SAVED)
        self.assertEqual(xl.inspect(self.folder / "Something Else.xlsx", read=False).state, xl.NOT_OPEN)

    def test_saved_state_is_read_and_matches_the_file_reader(self):
        book = self.open()
        self.type_in(book, {
            "B2": "Completed", "C2": "Had", "D2": "25:30", "G2": 2450, "B3": "Missed", "D3": "0:45",
            "G3": 1820.5, "G4": 0, "K2": 72.5, "L2": 84, "P5": 28,
        })
        book.Save()
        from_file = self.store.get_month(2026, 10)          # an ordinary folder can still be read
        with locked_file():
            from_excel = st.TrackerStore(self.path).get_month(2026, 10)
        self.assertEqual(from_excel, from_file)
        self.assertEqual(json.dumps(from_excel), json.dumps(from_file))
        self.assertEqual(from_excel["days"][0], {
            "date": "2026-10-01", "exercise": "Completed", "junk_food": "Had",
            "cardio_display": "25:30", "cardio_seconds": 1530,
            "calories_kcal": None, "protein_g": None, "weight_lifted_kg": 2450,
        })
        self.assertIsNone(from_excel["days"][3]["exercise"])            # blank stays blank
        self.assertEqual(from_excel["days"][2]["weight_lifted_kg"], 0)
        self.assertEqual(from_excel["measurements"][0]["weight_kg"], 72.5)
        self.assertEqual(from_excel["measurements"][3]["forearm_cm"], 28)
        self.assertEqual([d["date"] for d in from_excel["days"]][::30], ["2026-10-01", "2026-10-31"])
        self.assertEqual(from_excel["issues"], [])

    def test_unsaved_typing_is_not_shown_and_appears_after_saving(self):
        book = self.open()
        with locked_file():
            self.assertIsNone(self.store.get_month(2026, 10)["days"][0]["exercise"])
            self.type_in(book, {"B2": "Completed", "D2": "15:30", "G2": 1250})
            self.assertIs(bool(book.Saved), False)
            self.assertEqual(xl.inspect(self.path).state, xl.UNSAVED)
            for _ in range(2):
                day = self.store.get_month(2026, 10)["days"][0]
                self.assertEqual((day["exercise"], day["cardio_display"], day["weight_lifted_kg"]),
                                 (None, None, None))
            self.assertIs(self.store.status()["unsaved_changes"], True)
            book.Save()                                                  # Ctrl+S
            self.assertIs(self.store.status()["unsaved_changes"], False)
            day = self.store.get_month(2026, 10)["days"][0]
            self.assertEqual((day["exercise"], day["cardio_display"], day["weight_lifted_kg"]),
                             ("Completed", "15:30", 1250))
            self.type_in(book, {"B2": None})                             # changed again, not saved
            self.assertEqual(self.store.get_month(2026, 10)["days"][0]["exercise"], "Completed")
            book.Save()
            self.assertIsNone(self.store.get_month(2026, 10)["days"][0]["exercise"])

    def test_reading_leaves_excel_and_the_file_exactly_as_they_were(self):
        book = self.open()
        self.type_in(book, {"G2": 100})
        book.Save()
        before = (sha256(self.path), self.path.stat().st_mtime_ns)
        with locked_file():
            for _ in range(4):
                self.store.status()
                self.store.list_months()
                st.TrackerStore(self.path).get_month(2026, 10)
        self.assertEqual((sha256(self.path), self.path.stat().st_mtime_ns), before)
        self.assertIs(bool(book.Saved), True)                       # reading did not dirty it
        self.assertEqual(int(self.app.Workbooks.Count), 1)          # still open
        self.assertEqual(str(book.FullName).lower(), str(self.path).lower())
        self.assertIs(bool(self.app.Visible), False)                # settings untouched
        self.assertTrue(process_running(self.apps[0][1]))           # Excel was not closed

    def test_unsaved_workbook_is_not_saved_or_reset_by_reading(self):
        book = self.open()
        self.type_in(book, {"G2": 999})
        before = sha256(self.path)
        with locked_file():
            for _ in range(3):
                self.store.status()
                try:
                    st.TrackerStore(self.path).get_month(2026, 10)
                except st.WorkbookLockedError:
                    pass
        self.assertIs(bool(book.Saved), False)                              # still unsaved
        self.assertEqual(book.Worksheets("October 2026").Range("G2").Value, 999)   # typing kept
        self.assertEqual(sha256(self.path), before)                         # nothing written

    def test_right_workbook_among_several_excels(self):
        other_folder = self.folder / "another copy"
        other_folder.mkdir()
        other_path = wbk.create_workbook(other_folder / "Fitness_Tracker.xlsx")
        second = self.start_excel()                                  # a second Excel process
        self.assertNotEqual(self.apps[0][1], self.apps[-1][1])
        first_book, second_book = self.open(), self.open(other_path, second)
        self.type_in(first_book, {"G2": 111})
        self.type_in(second_book, {"G2": 222})
        first_book.Save()
        second_book.Save()
        with locked_file():
            self.assertEqual(st.TrackerStore(self.path).get_month(2026, 10)["days"][0]["weight_lifted_kg"], 111)
            self.assertEqual(st.TrackerStore(other_path).get_month(2026, 10)["days"][0]["weight_lifted_kg"], 222)
            self.type_in(second_book, {"G2": 333})                  # unsaved in the other Excel only
            self.assertEqual(xl.inspect(self.path, read=False).state, xl.SAVED)
            self.assertEqual(xl.inspect(other_path, read=False).state, xl.UNSAVED)

    def test_month_creation_stays_blocked_and_works_after_excel_closes_it(self):
        book = self.open()
        before = sha256(self.path)
        with self.assertRaises(st.WorkbookLockedError) as caught:
            self.store.create_month(2026, 11)
        self.assertIn("open in Excel or locked", str(caught.exception))
        self.assertEqual(sha256(self.path), before)
        self.assertEqual(int(book.Worksheets.Count), 1)
        book.Close(False)
        self.assertEqual(xl.inspect(self.path, read=False).state, xl.NOT_OPEN)
        self.assertEqual(self.store.create_month(2026, 11)["label"], "November 2026")
        self.assertEqual(len(self.store.get_month(2026, 11)["days"]), 30)       # file reader again

    def test_closing_and_reopening_excel_workbook_keeps_working(self):
        book = self.open()
        self.type_in(book, {"G2": 100})
        book.Save()
        with locked_file():
            self.assertEqual(self.store.get_month(2026, 10)["days"][0]["weight_lifted_kg"], 100)
        book.Close(False)
        self.assertEqual(self.store.get_month(2026, 10)["days"][0]["weight_lifted_kg"], 100)
        book = self.open()
        self.type_in(book, {"G2": 200})
        book.Save()
        with locked_file():
            self.assertEqual(self.store.get_month(2026, 10)["days"][0]["weight_lifted_kg"], 200)

    @unittest.skipUnless(EXE.is_file(), "the application has not been built")
    def test_packaged_program_reads_the_open_workbook(self):
        shutil.copy(EXE, self.folder)
        program = [str(self.folder / "Fitness Tracker.exe"), "--excel-status"]
        env = {k: v for k, v in os.environ.items() if not k.startswith("FITNESS_TRACKER")}

        def ask():
            done = subprocess.run(program, env=env, cwd=tempfile.gettempdir(), stdin=subprocess.DEVNULL,
                                  capture_output=True, text=True, timeout=120)
            self.assertEqual(done.returncode, 0, done.stdout + done.stderr)
            return json.loads(done.stdout.strip().splitlines()[-1])

        self.assertEqual(ask()["state"], xl.NOT_OPEN)
        book = self.open()
        report = ask()
        self.assertIs(report["reader_available"], True)                 # pywin32 is inside it
        self.assertEqual((report["state"], report["months"]), (xl.SAVED, ["October 2026"]))
        self.type_in(book, {"G2": 5})
        report = ask()
        self.assertEqual(report["state"], xl.UNSAVED)
        self.assertNotIn("sheets", report)
        self.assertIs(bool(book.Saved), False)                          # and it did not save it
        book.Save()
        self.assertEqual(ask()["state"], xl.SAVED)
        self.assertEqual(int(self.app.Workbooks.Count), 1)


if __name__ == "__main__":
    unittest.main()

"""Checks for writing daily rows and Sunday measurements to the workbook.

Every workbook here is a temporary one made for the test.

Run from the project root:  python -m unittest discover -s tests -v
"""

import json
import shutil
import tempfile
import threading
import time
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from backend import store as st
from backend import workbook as wbk
from backend.app import create_app
from tests.earlier_layout import add_month
from tests.test_migration import grid, layout, sha256

TODAY = date(2026, 10, 20)          # a Tuesday; the Sundays of October 2026 are 4, 11, 18 and 25

# October 2026: daily in A:G (row 2 is the 1st), measurements in J:P (row 2 is the 4th).
OCTOBER = {
    "B2": "Partial", "C2": "Had", "D2": "10:00", "E2": 1800, "F2": 90.5, "G2": 1000,
    "B6": "Completed", "G6": 2450.5, "B22": "Rest Day",
    "K2": 72.5, "L2": 84, "K3": 72, "P3": 28.5,
    "R1": "=SUM(G2:G32)",           # a formula of the user's own, beside the tables
}
# A month sheet that is not laid out the way the tracker lays one out. It must be left alone.
SEPTEMBER_OTHER = {"B2": True, "C2": False, "E2": 900, "I2": 73}
NOVEMBER = {"B2": "Missed", "K2": 71}

DAY_KEYS = {"date", "exercise", "junk_food", "cardio_display", "cardio_seconds",
            "calories_kcal", "protein_g", "weight_lifted_kg"}


class Case(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.path = self.tmp / "Fitness_Tracker.xlsx"
        wb = Workbook()
        notes = wb.active
        notes.title = "Notes"
        notes["A1"], notes["A2"] = "Gym notes", "=1+1"
        for (year, month, other), cells in (((2026, 9, True), SEPTEMBER_OTHER), ((2026, 10, False), OCTOBER),
                                            ((2026, 11, False), NOVEMBER)):
            ws = add_month(wb, year, month, earlier=other)
            for ref, value in cells.items():
                ws[ref] = value
        wb.save(self.path)
        self.before = sha256(self.path)
        self.store = st.TrackerStore(self.path)

    def day(self, number, fields, month=10, today=TODAY):
        return self.store.update_day(2026, month, number, fields, today)

    def sunday(self, number, fields, month=10, today=TODAY):
        return self.store.update_measurements(2026, month, number, fields, today)

    def sheet(self, name="October 2026"):
        return load_workbook(self.path)[name]

    def row(self, number, name="October 2026"):
        ws = self.sheet(name)
        return [ws.cell(number, col).value for col in range(2, 8)]

    def unchanged(self):
        self.assertEqual(sha256(self.path), self.before)
        self.assertEqual([p.name for p in self.tmp.iterdir()], ["Fitness_Tracker.xlsx"])

    def refused(self, call, fields=None):
        with self.assertRaises(st.InvalidUpdateError) as caught:
            call()
        self.unchanged()
        if fields is not None:
            self.assertEqual(sorted(caught.exception.fields), sorted(fields))
        return caught.exception


class DailyWriteTests(Case):
    def test_update_one_field(self):
        result = self.day(2, {"exercise": "Completed"})
        self.assertEqual(self.row(3), ["Completed", None, None, None, None, None])
        self.assertEqual(result["day"]["exercise"], "Completed")

    def test_update_every_field(self):
        result = self.day(2, {"exercise": "Completed", "junk_food": "None", "cardio": "25:30",
                              "calories_kcal": 2200, "protein_g": 140, "weight_lifted_kg": 2450})
        self.assertEqual(self.row(3), ["Completed", "None", "25:30", 2200, 140, 2450])
        self.assertEqual(result["day"], {
            "date": "2026-10-02", "exercise": "Completed", "junk_food": "None",
            "cardio_display": "25:30", "cardio_seconds": 1530,
            "calories_kcal": 2200, "protein_g": 140, "weight_lifted_kg": 2450,
        })

    def test_fields_left_out_are_kept(self):
        self.day(1, {"calories_kcal": 2000})
        self.assertEqual(self.row(2), ["Partial", "Had", "10:00", 2000, 90.5, 1000])
        self.day(1, {"exercise": "Completed", "protein_g": 120.5})
        self.assertEqual(self.row(2), ["Completed", "Had", "10:00", 2000, 120.5, 1000])

    def test_null_clears_a_field(self):
        result = self.day(1, {"cardio": None, "protein_g": None})
        self.assertEqual(self.row(2), ["Partial", "Had", None, 1800, None, 1000])
        self.assertEqual((result["day"]["cardio_display"], result["day"]["cardio_seconds"]), (None, None))
        self.day(1, {"exercise": None, "junk_food": "", "calories_kcal": None, "weight_lifted_kg": None})
        self.assertEqual(self.row(2), [None] * 6)

    def test_every_status_and_a_zero_can_be_stored(self):
        for status in ("Completed", "Partial", "Rest Day", "Missed"):
            self.day(3, {"exercise": status})
            self.assertEqual(self.row(4)[0], status)
        for status in ("None", "Controlled", "Had"):
            self.day(3, {"junk_food": status})
            self.assertEqual(self.row(4)[1], status)
        self.day(3, {"calories_kcal": 0, "protein_g": 0, "weight_lifted_kg": 0, "cardio": "0:00"})
        self.assertEqual(self.row(4)[2:], ["0:00", 0, 0, 0])

    def test_values_are_stored_in_their_tidy_form(self):
        self.day(3, {"exercise": " rest  day ", "junk_food": "controlled", "cardio": "05:07",
                     "calories_kcal": 2200.0, "protein_g": 140.0, "weight_lifted_kg": 2450.5})
        self.assertEqual(self.row(4), ["Rest Day", "Controlled", "5:07", 2200, 140, 2450.5])
        self.assertIs(type(self.row(4)[3]), int)

    def test_invalid_values_are_refused(self):
        for field, values in {
            "exercise": ["Done", True, 1, "Rest"],
            "junk_food": ["Yes", False, "none at all"],
            "cardio": ["25:75", 25.30, "1000:00", "25", "-1:00", 1530],
            "calories_kcal": [-1, 20001, 12.5, "2200", True],
            "protein_g": [-0.1, 1000.1, 12.55, "140", float("nan")],
            "weight_lifted_kg": [-5, 100000.5, 5.55, "heavy", float("inf")],
        }.items():
            for value in values:
                with self.subTest(field=field, value=value):
                    self.refused(lambda: self.day(2, {field: value}), [field])

    def test_every_invalid_field_is_reported_and_nothing_is_saved(self):
        error = self.refused(lambda: self.day(2, {
            "exercise": "Done", "junk_food": "None", "cardio": "25:75", "calories_kcal": 12.5,
            "protein_g": 140, "weight_lifted_kg": -1, "sleep": 8,
        }), ["exercise", "cardio", "calories_kcal", "weight_lifted_kg", "sleep"])
        self.assertIn("whole number", error.fields["calories_kcal"])
        self.assertIn("not a daily field", error.fields["sleep"])
        self.assertEqual(self.row(3), [None] * 6)             # the valid ones were not saved either

    def test_empty_or_wrongly_shaped_update_is_refused(self):
        for fields in ({}, [], None, "exercise", 5, [{"exercise": "Completed"}]):
            with self.subTest(fields=fields):
                self.refused(lambda: self.day(2, fields), ["body"])

    def test_measurement_fields_are_not_daily_fields(self):
        self.refused(lambda: self.day(4, {"weight_kg": 70}), ["weight_kg"])

    def test_past_day_and_today_can_be_edited(self):
        self.day(1, {"exercise": "Missed"})
        self.day(20, {"exercise": "Completed"})
        self.assertEqual((self.row(2)[0], self.row(21)[0]), ("Missed", "Completed"))
        self.day(1, {"exercise": "Completed"}, month=10, today=date(2027, 3, 1))      # long past
        self.assertEqual(self.row(2)[0], "Completed")

    def test_future_day_is_refused(self):
        self.before = sha256(self.path)
        error = self.refused(lambda: self.day(21, {"exercise": "Completed"}), ["day"])
        self.assertIn("future", error.fields["day"])
        self.refused(lambda: self.day(1, {"exercise": "Completed"}, month=11), ["day"])
        # Tomorrow becomes editable when it arrives.
        self.day(21, {"exercise": "Completed"}, today=date(2026, 10, 21))
        self.assertEqual(self.row(22)[0], "Completed")

    def test_future_day_and_bad_value_are_reported_together(self):
        self.refused(lambda: self.day(25, {"exercise": "Done"}), ["day", "exercise"])

    def test_day_that_does_not_exist_is_refused(self):
        for month, number in ((10, 0), (10, 32), (9, 31), (10, -1), (10, "1"), (10, 1.0), (10, True)):
            with self.subTest(month=month, day=number):
                with self.assertRaises(st.DayNotFoundError):
                    self.day(number, {"exercise": "Completed"}, month=month)
                self.unchanged()

    def test_bad_or_missing_month_is_refused(self):
        with self.assertRaises(st.InvalidMonthError):
            self.store.update_day(2026, 13, 1, {"exercise": "Completed"}, TODAY)
        with self.assertRaises(st.InvalidMonthError):
            self.store.update_day(1999, 1, 1, {"exercise": "Completed"}, TODAY)
        with self.assertRaises(st.MonthNotFoundError):
            self.store.update_day(2026, 8, 1, {"exercise": "Completed"}, TODAY)
        self.unchanged()

    def test_month_laid_out_another_way_is_refused_and_left_alone(self):
        with self.assertRaises(st.WorkbookFormatError):
            self.day(1, {"exercise": "Completed"}, month=9)
        self.unchanged()
        self.assertFalse(wbk.is_month_sheet(self.sheet("September 2026")))
        self.assertIs(self.sheet("September 2026")["B2"].value, True)

    def test_month_in_an_unknown_layout_is_refused(self):
        wb = load_workbook(self.path)
        wb["October 2026"]["C1"] = "Sweets"
        wb.save(self.path)
        self.before = sha256(self.path)
        with self.assertRaises(st.WorkbookFormatError):
            self.day(1, {"exercise": "Completed"})
        self.unchanged()

    def test_month_with_wrong_dates_is_refused(self):
        wb = load_workbook(self.path)
        wb["October 2026"]["A3"] = datetime(2026, 10, 1)                # the 1st twice, no 2nd
        wb.save(self.path)
        self.before = sha256(self.path)
        with self.assertRaises(st.WorkbookFormatError):
            self.day(1, {"exercise": "Completed"})
        self.unchanged()

    def test_default_today_is_the_real_date(self):
        with self.assertRaises(st.InvalidUpdateError):
            self.store.update_day(2099, 1, 1, {"exercise": "Completed"})
        self.assertEqual(self.store.update_day(2026, 10, 1, {"exercise": "Missed"})["day"]["exercise"], "Missed")

    def test_result_structure(self):
        result = self.day(2, {"exercise": "Completed"})
        self.assertEqual(set(result), {"year", "month", "label", "day", "workbook_modified"})
        self.assertEqual((result["year"], result["month"], result["label"]), (2026, 10, "October 2026"))
        self.assertEqual(set(result["day"]), DAY_KEYS)
        self.assertEqual(result["workbook_modified"], self.store.status()["workbook_modified"])
        self.assertEqual(result["day"], self.store.get_month(2026, 10)["days"][1])


class MeasurementWriteTests(Case):
    def test_update_one_measurement(self):
        result = self.sunday(18, {"weight_kg": 71.5})
        ws = self.sheet()
        self.assertEqual([ws.cell(4, col).value for col in range(11, 17)], [71.5, None, None, None, None, None])
        self.assertEqual(result["measurement"], {
            "sunday": "2026-10-18", "weight_kg": 71.5, "waist_cm": None, "chest_cm": None,
            "bicep_cm": None, "thigh_cm": None, "forearm_cm": None,
        })
        self.assertEqual(set(result), {"year", "month", "label", "measurement", "workbook_modified"})

    def test_update_several_measurements(self):
        self.sunday(18, {"weight_kg": 70, "waist_cm": 80, "chest_cm": 98, "bicep_cm": 35.5,
                         "thigh_cm": 55, "forearm_cm": 28.5})
        ws = self.sheet()
        self.assertEqual([ws.cell(4, col).value for col in range(11, 17)], [70, 80, 98, 35.5, 55, 28.5])

    def test_measurements_left_out_are_kept_and_null_clears(self):
        self.sunday(4, {"chest_cm": 98})
        ws = self.sheet()
        self.assertEqual([ws.cell(2, col).value for col in range(11, 17)], [72.5, 84, 98, None, None, None])
        result = self.sunday(4, {"waist_cm": None, "weight_kg": 72})
        ws = self.sheet()
        self.assertEqual([ws.cell(2, col).value for col in range(11, 17)], [72, None, 98, None, None, None])
        self.assertIsNone(result["measurement"]["waist_cm"])

    def test_sunday_that_is_today_can_be_edited(self):
        self.sunday(18, {"weight_kg": 71}, today=date(2026, 10, 18))
        self.assertEqual(self.sheet()["K4"].value, 71)

    def test_day_that_is_not_a_sunday_is_refused(self):
        for number in (1, 5, 17, 19):
            with self.subTest(day=number):
                error = self.refused(lambda: self.sunday(number, {"weight_kg": 70}), ["day"])
                self.assertIn("not a Sunday", error.fields["day"])

    def test_future_sunday_is_refused(self):
        error = self.refused(lambda: self.sunday(25, {"weight_kg": 70}), ["day"])
        self.assertIn("future", error.fields["day"])
        self.refused(lambda: self.sunday(1, {"weight_kg": 70}, month=11), ["day"])
        self.sunday(25, {"weight_kg": 70}, today=date(2026, 10, 25))
        self.assertEqual(self.sheet()["K5"].value, 70)

    def test_invalid_measurements_are_refused(self):
        for value in (0, -1, 500.1, 70.55, "70", True, float("nan")):
            for field in ("weight_kg", "forearm_cm"):
                with self.subTest(field=field, value=value):
                    self.refused(lambda: self.sunday(4, {field: value}), [field])
        self.refused(lambda: self.sunday(4, {"weight_kg": 0, "waist_cm": 80, "chest_cm": 600, "neck_cm": 40}),
                     ["weight_kg", "chest_cm", "neck_cm"])
        self.refused(lambda: self.sunday(4, {"exercise": "Completed"}), ["exercise"])
        self.refused(lambda: self.sunday(4, {}), ["body"])

    def test_missing_day_and_missing_month(self):
        with self.assertRaises(st.DayNotFoundError):
            self.sunday(32, {"weight_kg": 70})
        with self.assertRaises(st.MonthNotFoundError):
            self.sunday(2, {"weight_kg": 70}, month=8)                  # 2 August 2026 is a Sunday
        self.unchanged()

    def test_month_laid_out_another_way_is_refused(self):
        with self.assertRaises(st.WorkbookFormatError):
            self.sunday(6, {"weight_kg": 70}, month=9)
        self.unchanged()
        self.assertEqual(self.sheet("September 2026")["I2"].value, 73)


class PreservationTests(Case):
    def everything(self):
        wb = load_workbook(self.path)
        return {ws.title: grid(ws) for ws in wb}, {ws.title: layout(ws) for ws in wb}, wb.sheetnames

    def test_daily_write_changes_only_the_cells_asked_for(self):
        values, layouts, names = self.everything()
        self.day(5, {"exercise": "Missed", "calories_kcal": 2100})
        after_values, after_layouts, after_names = self.everything()
        self.assertEqual(after_names, names)
        self.assertEqual(after_layouts, layouts)
        changed = [(name, r + 1, c + 1) for name in names for r, line in enumerate(values[name])
                   for c, value in enumerate(line) if after_values[name][r][c] != value]
        self.assertEqual(changed, [("October 2026", 6, 2), ("October 2026", 6, 5)])
        self.assertEqual(self.row(6), ["Missed", None, None, 2100, None, 2450.5])       # same row, other cells

    def test_measurement_write_changes_only_the_cells_asked_for(self):
        values, layouts, names = self.everything()
        self.sunday(11, {"waist_cm": 83.5, "forearm_cm": None})
        after_values, after_layouts, _ = self.everything()
        self.assertEqual(after_layouts, layouts)
        changed = [(name, r + 1, c + 1) for name in names for r, line in enumerate(values[name])
                   for c, value in enumerate(line) if after_values[name][r][c] != value]
        self.assertEqual(changed, [("October 2026", 3, 12), ("October 2026", 3, 16)])
        self.assertEqual(self.sheet()["K3"].value, 72)                                  # that Sunday's weight
        self.assertEqual(self.sheet()["K2"].value, 72.5)                                # the Sunday before

    def test_formulas_headers_tables_and_validation_survive(self):
        self.day(1, {"weight_lifted_kg": 1200})
        self.sunday(4, {"weight_kg": 72})
        wb = load_workbook(self.path)
        ws = wb["October 2026"]
        self.assertEqual((ws["R1"].value, wb["Notes"]["A2"].value), ("=SUM(G2:G32)", "=1+1"))
        self.assertEqual([c.value for c in ws[1]][:7], wbk.DAILY_HEADERS)
        self.assertEqual([c.value for c in ws[1]][9:16], wbk.MEASUREMENT_HEADERS)
        self.assertEqual({t.name: t.ref for t in ws.tables.values()},
                         {"Daily_2026_10": "A1:G32", "Measurements_2026_10": "J1:P5"})
        rules = {str(rule.sqref): rule for rule in ws.data_validations.dataValidation}
        self.assertEqual(len(rules), 7)
        self.assertEqual(rules["B2:B32"].formula1, '"Completed,Partial,Rest Day,Missed"')
        self.assertEqual(rules["C2:C32"].formula1, '"None,Controlled,Had"')
        self.assertEqual([wbk.is_month_sheet(sheet) for sheet in wb], [False, False, True, True])
        self.assertEqual(wb.sheetnames, ["Notes", "September 2026", "October 2026", "November 2026"])

    def test_cell_formats_are_kept_when_values_are_written(self):
        self.day(2, {"exercise": "Completed", "cardio": "25:30", "calories_kcal": 2200,
                     "protein_g": 140.5, "weight_lifted_kg": 2450})
        self.sunday(18, {"weight_kg": 71.5})
        ws = self.sheet()
        self.assertEqual([ws[ref].number_format for ref in ("A3", "B3", "D3", "E3", "F3", "G3", "J4", "K4")],
                         [wbk.DATE_FORMAT, "@", "@", "0", "0.0", "0.0", wbk.DATE_FORMAT, "0.0"])
        self.assertEqual(ws["A3"].value, datetime(2026, 10, 2))
        self.assertIs(type(ws["D3"].value), str)

    def test_cells_the_user_typed_wrongly_elsewhere_are_left_alone(self):
        wb = load_workbook(self.path)
        wb["October 2026"]["B10"], wb["October 2026"]["E11"] = "Done", "lots"
        wb.save(self.path)
        self.day(1, {"exercise": "Completed"})
        ws = self.sheet()
        self.assertEqual((ws["B10"].value, ws["E11"].value, ws["B2"].value), ("Done", "lots", "Completed"))

    def test_many_writes_leave_a_readable_workbook(self):
        _, layouts, names = self.everything()
        for number in range(1, 21):
            self.day(number, {"exercise": ("Completed", "Partial", "Rest Day", "Missed")[number % 4],
                              "junk_food": ("None", "Controlled", "Had")[number % 3],
                              "cardio": f"{number}:{number:02d}", "calories_kcal": 2000 + number,
                              "protein_g": 100 + number / 2, "weight_lifted_kg": 1000 + number})
            self.day(number, {"cardio": None} if number % 5 == 0 else {"protein_g": 150})
        for number in (4, 11, 18):
            self.sunday(number, {"weight_kg": 70 + number / 10, "waist_cm": 80})
        _, after_layouts, after_names = self.everything()
        self.assertEqual((after_names, after_layouts), (names, layouts))
        data = self.store.get_month(2026, 10)
        self.assertEqual(data["issues"], [])
        self.assertEqual(data["days"][6], {
            "date": "2026-10-07", "exercise": "Missed", "junk_food": "Controlled",
            "cardio_display": "7:07", "cardio_seconds": 427,
            "calories_kcal": 2007, "protein_g": 150, "weight_lifted_kg": 1007,
        })
        self.assertIsNone(data["days"][9]["cardio_display"])
        self.assertEqual(data["days"][9]["protein_g"], 105)
        self.assertEqual(data["days"][20]["exercise"], "Rest Day")              # the 21st, never written
        self.assertEqual(data["measurements"][2]["weight_kg"], 71.8)
        self.assertIs(self.sheet("September 2026")["B2"].value, True)          # the other sheet is as it was
        streamed = load_workbook(self.path, read_only=True)
        self.assertEqual(streamed.sheetnames, names)
        streamed.close()                 # this mode keeps the file open until it is closed
        self.assertEqual([p.name for p in self.tmp.iterdir()], ["Fitness_Tracker.xlsx"])


class SafetyTests(Case):
    def test_workbook_open_in_excel_is_refused_and_untouched(self):
        with mock.patch("builtins.open", side_effect=PermissionError("locked")):
            with self.assertRaises(st.WorkbookLockedError):
                self.day(1, {"exercise": "Completed"})
            with self.assertRaises(st.WorkbookLockedError):
                self.sunday(4, {"weight_kg": 70})
        self.unchanged()

    def test_workbook_locked_at_the_swap_is_untouched(self):
        with mock.patch.object(st.os, "replace", side_effect=PermissionError("locked")):
            with self.assertRaises(st.WorkbookLockedError):
                self.day(1, {"exercise": "Completed"})
        self.unchanged()

    def test_save_that_fails_leaves_the_original_and_no_temporary_file(self):
        with mock.patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                self.day(1, {"exercise": "Completed"})
        self.unchanged()

    def test_saved_file_that_differs_from_what_was_asked_is_thrown_away(self):
        def drop_lists(ws):
            ws.data_validations.dataValidation = [
                rule for rule in ws.data_validations.dataValidation if rule.type != "list"]

        def drop_table(ws):
            del ws.tables["Measurements_2026_10"]

        changes = {
            "wrong value": lambda ws: ws.__setitem__("B2", "Missed"),
            "value not written": lambda ws: ws.__setitem__("B2", "Partial"),
            "another day changed": lambda ws: ws.__setitem__("B3", "Completed"),
            "another field changed": lambda ws: ws.__setitem__("G2", None),
            "measurement changed": lambda ws: ws.__setitem__("K2", 80),
            "date changed": lambda ws: ws.__setitem__("A2", datetime(2026, 10, 2)),
            "header changed": lambda ws: ws.__setitem__("G1", "Weight"),
            "formula lost": lambda ws: ws.__setitem__("R1", None),
            "format changed": lambda ws: setattr(ws["F2"], "number_format", "0.00"),
            "another month changed": lambda ws: ws.parent["November 2026"].__setitem__("B2", "Completed"),
            "the other month sheet changed": lambda ws: ws.parent["September 2026"].__setitem__("B2", False),
            "notes changed": lambda ws: ws.parent["Notes"].__setitem__("A1", "x"),
            "sheet removed": lambda ws: ws.parent.remove(ws.parent["Notes"]),
            "lists removed": drop_lists,
            "table removed": drop_table,
        }
        real = st._write_cells
        for what, change in changes.items():
            def wrong(ws, cells, change=change):
                real(ws, cells)
                change(ws)

            with self.subTest(what):
                with mock.patch.object(st, "_write_cells", wrong):
                    with self.assertRaises(st.TrackerError) as caught:
                        self.day(1, {"exercise": "Completed"})
                self.assertIn(caught.exception.code, ("workbook_unreadable", "workbook_format"))
                self.unchanged()

    def test_nothing_is_written_through_excel(self):
        import inspect
        source = inspect.getsource(st.TrackerStore._update) + inspect.getsource(st._write_cells) \
            + inspect.getsource(st._target_cells) + inspect.getsource(st.TrackerStore._save)
        for word in ("excel_live", "pythoncom", "win32com", "_read(", "_read_from_excel"):
            self.assertNotIn(word, source)

    def test_writes_never_ask_excel(self):
        with mock.patch.object(st.excel_live, "inspect", side_effect=AssertionError("asked Excel")):
            self.day(1, {"exercise": "Completed"})
            self.sunday(4, {"weight_kg": 70})

    def test_writes_happen_one_at_a_time(self):
        fields = {"exercise": "Completed", "junk_food": "None", "cardio": "25:30",
                  "calories_kcal": 2200, "protein_g": 140, "weight_lifted_kg": 2450}
        state = {"inside": 0, "most": 0}
        guard = threading.Lock()
        real_load, real_save = self.store._load, self.store._save

        def load():
            with guard:
                state["inside"] += 1
                state["most"] = max(state["most"], state["inside"])
            time.sleep(0.03)
            return real_load()

        def save(*args, **kwargs):
            try:
                return real_save(*args, **kwargs)
            finally:
                with guard:
                    state["inside"] -= 1

        errors = []

        def write(field, value):
            try:
                self.day(7, {field: value})
            except Exception as error:  # noqa: BLE001 - reported below
                errors.append(error)

        with mock.patch.object(self.store, "_load", load), mock.patch.object(self.store, "_save", save):
            threads = [threading.Thread(target=write, args=item) for item in fields.items()]
            threads.append(threading.Thread(target=lambda: errors.extend(
                [] if self.sunday(11, {"waist_cm": 83}) else ["no result"])))
            for thread in threads:
                thread.start()
            for thread in threads:
                thread.join()
        self.assertEqual(errors, [])
        self.assertEqual(state["most"], 1)
        # Each write was built on the one before it, so none was lost.
        self.assertEqual(self.row(8), list(fields.values()))
        self.assertEqual(self.sheet()["L3"].value, 83)

    def test_each_write_starts_from_the_file_as_it_is_now(self):
        self.day(1, {"exercise": "Completed"})
        wb = load_workbook(self.path)                       # saved from Excel in between
        wb["October 2026"]["G10"], wb["October 2026"]["B2"] = 555, "Missed"
        wb["Notes"]["A5"] = "added later"
        wb.save(self.path)
        result = self.day(1, {"calories_kcal": 2000})
        ws = self.sheet()
        self.assertEqual((ws["G10"].value, ws["B2"].value, ws["E2"].value), (555, "Missed", 2000))
        self.assertEqual(result["day"]["exercise"], "Missed")
        self.assertEqual(load_workbook(self.path)["Notes"]["A5"].value, "added later")

    def test_two_stores_on_one_file_do_not_undo_each_other(self):
        other = st.TrackerStore(self.path)
        self.day(1, {"exercise": "Completed"})
        other.update_day(2026, 10, 1, {"calories_kcal": 2000}, TODAY)
        self.day(1, {"protein_g": 150})
        self.assertEqual(self.row(2), ["Completed", "Had", "10:00", 2000, 150, 1000])

    def test_missing_workbook_is_reported(self):
        self.path.unlink()
        with self.assertRaises(st.WorkbookNotFoundError):
            self.day(1, {"exercise": "Completed"})
        self.assertEqual(list(self.tmp.iterdir()), [])


class ApiTests(Case):
    def setUp(self):
        super().setUp()
        self.client = create_app(self.path, today=lambda: TODAY).test_client()

    def put(self, url, body, **kwargs):
        return self.client.put(url, json=body, **kwargs)

    def error(self, response, status, code):
        self.assertEqual(response.status_code, status, response.get_data(as_text=True))
        error = response.get_json()["error"]
        self.assertEqual(error["code"], code)
        self.assertTrue(error["message"])
        return error

    def test_daily_put(self):
        response = self.put("/api/months/2026/10/days/2", {
            "exercise": "Completed", "junk_food": "None", "cardio": "25:30",
            "calories_kcal": 2200, "protein_g": 140, "weight_lifted_kg": 2450,
        })
        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertEqual(set(body), {"year", "month", "label", "day", "workbook_modified"})
        self.assertEqual(body["day"], {
            "date": "2026-10-02", "exercise": "Completed", "junk_food": "None",
            "cardio_display": "25:30", "cardio_seconds": 1530,
            "calories_kcal": 2200, "protein_g": 140, "weight_lifted_kg": 2450,
        })
        self.assertEqual((body["year"], body["month"], body["label"]), (2026, 10, "October 2026"))
        self.assertEqual(body["workbook_modified"], self.client.get("/api/status").get_json()["workbook_modified"])
        self.assertEqual(self.client.get("/api/months/2026/10").get_json()["days"][1], body["day"])
        self.assertEqual(self.row(3), ["Completed", "None", "25:30", 2200, 140, 2450])

    def test_measurement_put(self):
        response = self.put("/api/months/2026/10/measurements/18", {"weight_kg": 70, "waist_cm": 80, "chest_cm": 98})
        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertEqual(set(body), {"year", "month", "label", "measurement", "workbook_modified"})
        self.assertEqual(body["measurement"], {
            "sunday": "2026-10-18", "weight_kg": 70, "waist_cm": 80, "chest_cm": 98,
            "bicep_cm": None, "thigh_cm": None, "forearm_cm": None,
        })
        self.assertEqual(self.client.get("/api/months/2026/10").get_json()["measurements"][2], body["measurement"])

    def test_partial_put_and_null(self):
        self.assertEqual(self.put("/api/months/2026/10/days/1", {"cardio": None}).status_code, 200)
        self.assertEqual(self.row(2), ["Partial", "Had", None, 1800, 90.5, 1000])

    def test_malformed_json(self):
        for data in ("{not json", "", '{"exercise": "Completed"', "null"):
            with self.subTest(data=data):
                response = self.client.put("/api/months/2026/10/days/1", data=data,
                                           content_type="application/json")
                self.assertIn("body", self.error(response, 400, "invalid_request")["fields"])
        for body in ([], [1], "text", 5):
            with self.subTest(body=body):
                self.assertIn("body", self.error(self.put("/api/months/2026/10/days/1", body),
                                                 400, "invalid_request")["fields"])
        self.unchanged()

    def test_missing_or_wrong_content_type(self):
        data = json.dumps({"exercise": "Completed"})
        for content_type in (None, "text/plain", "application/x-www-form-urlencoded"):
            with self.subTest(content_type=content_type):
                response = self.client.put("/api/months/2026/10/days/1", data=data, content_type=content_type)
                self.assertIn("JSON", self.error(response, 400, "invalid_request")["message"])
        response = self.client.put("/api/months/2026/10/measurements/4", data='{"weight_kg": 70}')
        self.error(response, 400, "invalid_request")
        self.unchanged()

    def test_invalid_fields_are_all_listed(self):
        error = self.error(self.put("/api/months/2026/10/days/2", {
            "exercise": "Done", "calories_kcal": 12.5, "protein_g": 140}), 400, "invalid_request")
        self.assertEqual(sorted(error["fields"]), ["calories_kcal", "exercise"])
        error = self.error(self.put("/api/months/2026/10/measurements/4", {"weight_kg": 0, "waist_cm": 501}),
                           400, "invalid_request")
        self.assertEqual(sorted(error["fields"]), ["waist_cm", "weight_kg"])
        self.unchanged()

    def test_empty_update(self):
        for url in ("/api/months/2026/10/days/1", "/api/months/2026/10/measurements/4"):
            error = self.error(self.put(url, {}), 400, "invalid_request")
            self.assertEqual(error["fields"], {"body": "Nothing to change was sent."})
        self.unchanged()

    def test_future_date_and_not_a_sunday(self):
        self.assertIn("future", self.error(self.put("/api/months/2026/10/days/21", {"exercise": "Completed"}),
                                           400, "invalid_request")["fields"]["day"])
        self.assertIn("future", self.error(self.put("/api/months/2026/10/measurements/25", {"weight_kg": 70}),
                                           400, "invalid_request")["fields"]["day"])
        self.assertIn("not a Sunday", self.error(self.put("/api/months/2026/10/measurements/5", {"weight_kg": 70}),
                                                 400, "invalid_request")["fields"]["day"])
        self.unchanged()

    def test_month_laid_out_another_way(self):
        # The tracker has one layout. A month sheet in any other is reported, never guessed at or changed.
        self.error(self.put("/api/months/2026/9/days/1", {"exercise": "Completed"}), 500, "workbook_format")
        self.error(self.put("/api/months/2026/9/measurements/6", {"weight_kg": 70}), 500, "workbook_format")
        self.unchanged()
        for url in ("/api/months/2026/9", "/api/months/2026/9/analytics"):
            self.error(self.client.get(url), 500, "workbook_format")
        self.unchanged()

    def test_missing_month_and_missing_date(self):
        self.error(self.put("/api/months/2026/8/days/1", {"exercise": "Completed"}), 404, "month_not_found")
        self.error(self.put("/api/months/2026/8/measurements/2", {"weight_kg": 70}), 404, "month_not_found")
        self.error(self.put("/api/months/2026/9/days/31", {"exercise": "Completed"}), 404, "day_not_found")
        self.error(self.put("/api/months/2026/10/days/0", {"exercise": "Completed"}), 404, "day_not_found")
        self.error(self.put("/api/months/2026/10/measurements/32", {"weight_kg": 70}), 404, "day_not_found")
        self.error(self.put("/api/months/2026/13/days/1", {"exercise": "Completed"}), 400, "invalid_month")
        # A day that is not a number matches no edit route at all.
        self.assertIn(self.put("/api/months/2026/10/days/first", {"exercise": "Completed"}).status_code, (404, 405))
        self.unchanged()

    def test_locked_workbook(self):
        with mock.patch("builtins.open", side_effect=PermissionError("locked")):
            error = self.error(self.put("/api/months/2026/10/days/1", {"exercise": "Completed"}),
                               423, "workbook_locked")
            self.error(self.put("/api/months/2026/10/measurements/4", {"weight_kg": 70}), 423, "workbook_locked")
        self.assertIn("Nothing was changed", error["message"])
        self.unchanged()

    def test_client_cannot_choose_the_file_or_the_sheet(self):
        for extra in ({"path": "C:/other.xlsx"}, {"workbook": "other.xlsx"}, {"sheet": "Notes"},
                      {"date": "2026-10-09"}, {"year": 2025}):
            with self.subTest(extra=extra):
                error = self.error(self.put("/api/months/2026/10/days/1", {"exercise": "Completed", **extra}),
                                   400, "invalid_request")
                self.assertEqual(list(error["fields"]), list(extra))
        for url in ("/api/months/2026/10/days/1?path=C:/other.xlsx&sheet=Notes&day=9",):
            self.assertEqual(self.put(url, {"exercise": "Completed"}).status_code, 200)
        self.assertEqual(sorted(p.name for p in self.tmp.iterdir()), ["Fitness_Tracker.xlsx"])
        self.assertEqual((self.row(2)[0], self.row(10)[0]), ("Completed", None))    # the URL's day, nothing else

    def test_only_put_is_accepted_on_the_edit_routes(self):
        for method in ("post", "delete", "patch"):
            response = getattr(self.client, method)("/api/months/2026/10/days/1", json={"exercise": "Completed"})
            self.assertEqual(response.status_code, 405, method)
        self.assertEqual(self.client.get("/api/months/2026/10/days/1").status_code, 404)
        self.unchanged()

    def test_requests_from_another_website_are_refused(self):
        body = {"exercise": "Completed"}
        url = "/api/months/2026/10/days/1"
        for headers in ({"Origin": "http://evil.example"}, {"Origin": "null"},
                        {"Origin": "http://localhost:9999"}, {"Origin": "https://localhost"},
                        {"Sec-Fetch-Site": "cross-site"}, {"Sec-Fetch-Site": "same-site"},
                        {"Host": "evil.example"}, {"Host": "evil.example:5000", "Origin": "http://evil.example:5000"}):
            with self.subTest(headers=headers):
                self.error(self.put(url, body, headers=headers), 403, "forbidden_origin")
                self.error(self.put("/api/months/2026/10/measurements/4", {"weight_kg": 70}, headers=headers),
                           403, "forbidden_origin")
                self.error(self.client.post("/api/months", json={"year": 2026, "month": 12}, headers=headers),
                           403, "forbidden_origin")
        self.unchanged()

    def test_the_trackers_own_page_is_accepted(self):
        for headers in ({"Origin": "http://localhost"}, {"Origin": "http://localhost", "Sec-Fetch-Site": "same-origin"},
                        {"Host": "127.0.0.1:5000", "Origin": "http://127.0.0.1:5000"}, {}):
            with self.subTest(headers=headers):
                self.assertEqual(self.put("/api/months/2026/10/days/1", {"exercise": "Missed"},
                                          headers=headers).status_code, 200)
        created = self.client.post("/api/months", json={"year": 2026, "month": 12},
                                   headers={"Origin": "http://localhost"})
        self.assertEqual(created.status_code, 201)

    def test_reading_is_not_affected_by_the_check(self):
        response = self.client.get("/api/months", headers={"Origin": "http://evil.example"})
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("Access-Control-Allow-Origin", response.headers)
        options = self.client.open("/api/months/2026/10/days/1", method="OPTIONS",
                                   headers={"Origin": "http://evil.example",
                                            "Access-Control-Request-Method": "PUT"})
        self.assertNotIn("Access-Control-Allow-Origin", options.headers)      # so the browser blocks it

    def test_write_is_seen_by_the_next_read_without_a_restart(self):
        before = self.client.get("/api/months/2026/10").get_json()["days"][2]
        self.assertIsNone(before["exercise"])
        self.put("/api/months/2026/10/days/3", {"exercise": "Rest Day"})
        self.assertEqual(self.client.get("/api/months/2026/10").get_json()["days"][2]["exercise"], "Rest Day")


if __name__ == "__main__":
    unittest.main()

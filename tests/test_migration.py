"""Checks for the converter that turns a workbook in the tracker's earlier layout into its layout.

The converter is a tool run by hand; the tracker itself never calls it.

Every workbook here is a temporary one made for the test.

Run from the project root:  python -m unittest discover -s tests -v
"""

import hashlib
import shutil
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from backend import migrate
from backend import store as st
from backend import workbook as wbk
from backend.app import create_app
from tests.earlier_layout import add_month

EARLIER, CURRENT = True, False       # how a month in a test workbook is laid out

# Cells for October 2026 in the earlier layout: daily in A:E, measurements in H:N.
OCTOBER_EARLIER = {
    "B2": True, "C2": True, "D2": "25:30", "E2": 2450.5,
    "B3": False, "C3": False, "D3": "5:05",
    "E4": 0,
    "B5": True, "C5": False, "E5": 1800,
    "I2": 72.5, "J2": 84, "K2": 100, "L2": 35, "M2": 55, "N2": 28.5,
    "I3": 72.1, "N5": 30,
}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def grid(ws):
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def layout(ws):
    """Everything about a sheet except the values typed into it."""
    return {
        "tables": {table.name: (table.ref, table.tableStyleInfo.name) for table in ws.tables.values()},
        "validations": sorted((str(rule.sqref), rule.type, rule.operator, rule.formula1, rule.formula2,
                               rule.allow_blank, rule.showErrorMessage)
                              for rule in ws.data_validations.dataValidation),
        "formats": {cell.coordinate: cell.number_format for row in ws.iter_rows() for cell in row},
        "widths": {key: dim.width for key, dim in ws.column_dimensions.items()},
        "freeze": ws.freeze_panes,
        "bold_headers": [cell.font.b for cell in ws[1]],
    }


class Case(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.path = self.tmp / "Fitness_Tracker.xlsx"
        self.store = st.TrackerStore(self.path)

    def build(self, months, cells=None, extra=None):
        """Make the workbook: months is {(year, month): EARLIER or CURRENT}, cells is {sheet: {ref: value}}."""
        wb = Workbook()
        wb.remove(wb.active)
        for (year, month), earlier in months.items():
            add_month(wb, year, month, earlier=earlier)
        for sheet, values in (cells or {}).items():
            for ref, value in values.items():
                wb[sheet][ref] = value
        if extra:
            extra(wb)
        wb.save(self.path)
        return sha256(self.path)

    def october(self):
        return self.build({(2026, 10): EARLIER}, {"October 2026": OCTOBER_EARLIER})

    def convert(self):
        return migrate.convert(self.store)

    def files(self):
        return sorted(p.name for p in self.tmp.iterdir())

    def sheet(self, name="October 2026"):
        return load_workbook(self.path)[name]


class SingleSheetTests(Case):
    def test_one_month_is_converted(self):
        self.october()
        self.assertIs(migrate.is_earlier_sheet(self.sheet()), True)
        result = self.convert()
        self.assertEqual(result, {
            "status": "converted", "converted": ["October 2026"], "unchanged": [],
            "unsupported": [], "carried_over": [],
        })
        ws = self.sheet()
        self.assertIs(wbk.is_month_sheet(ws), True)
        self.assertIs(migrate.is_earlier_sheet(ws), False)
        self.assertEqual([c.value for c in ws[1]][:7], wbk.DAILY_HEADERS)
        self.assertEqual([c.value for c in ws[1]][9:16], wbk.MEASUREMENT_HEADERS)
        self.assertEqual(self.files(), ["Fitness_Tracker.xlsx"])           # no copy of any kind is made

    def test_exercise_true_false_blank(self):
        self.october()
        self.convert()
        ws = self.sheet()
        self.assertEqual([ws[f"B{row}"].value for row in (2, 3, 4, 5)], ["Completed", "Missed", None, "Completed"])
        self.assertEqual({ws[f"B{row}"].value for row in range(6, 33)}, {None})

    def test_no_junk_food_true_false_blank(self):
        self.october()
        self.convert()
        ws = self.sheet()
        self.assertEqual([ws[f"C{row}"].value for row in (2, 3, 4, 5)], ["None", "Had", None, "Had"])
        self.assertEqual({ws[f"C{row}"].value for row in range(6, 33)}, {None})

    def test_cardio_is_kept(self):
        self.october()
        self.convert()
        ws = self.sheet()
        self.assertEqual([ws[f"D{row}"].value for row in (2, 3, 4)], ["25:30", "5:05", None])

    def test_total_weight_lifted_becomes_weight_lifted(self):
        self.october()
        self.convert()
        ws = self.sheet()
        self.assertEqual(ws["G1"].value, "Weight Lifted (kg)")
        self.assertEqual([ws[f"G{row}"].value for row in (2, 3, 4, 5)], [2450.5, None, 0, 1800])

    def test_calories_and_protein_start_blank(self):
        self.october()
        self.convert()
        ws = self.sheet()
        self.assertEqual((ws["E1"].value, ws["F1"].value), ("Calories (kcal)", "Protein (g)"))
        self.assertEqual({ws.cell(row, col).value for row in range(2, 33) for col in (5, 6)}, {None})

    def test_measurements_are_kept_and_moved_to_j_to_p(self):
        self.october()
        self.convert()
        ws = self.sheet()
        self.assertEqual([ws.cell(2, col).value for col in range(10, 17)],
                         [datetime(2026, 10, 4), 72.5, 84, 100, 35, 55, 28.5])
        self.assertEqual((ws["K3"].value, ws["P5"].value, ws["K5"].value), (72.1, 30, None))
        self.assertEqual({ws.cell(row, col).value for row in range(1, 33) for col in (8, 9)}, {None})

    def test_dates_are_kept(self):
        self.october()
        self.convert()
        ws = self.sheet()
        self.assertEqual([ws[f"A{row}"].value.date() for row in range(2, 33)], wbk.month_dates(2026, 10))
        self.assertEqual([ws[f"J{row}"].value.date() for row in range(2, 6)], wbk.month_sundays(2026, 10))

    def test_converted_month_reads_through_the_store(self):
        self.october()
        self.convert()
        data = self.store.get_month(2026, 10)
        self.assertEqual(data["issues"], [])
        self.assertEqual(data["days"][0], {
            "date": "2026-10-01", "exercise": "Completed", "junk_food": "None",
            "cardio_display": "25:30", "cardio_seconds": 1530,
            "calories_kcal": None, "protein_g": None, "weight_lifted_kg": 2450.5,
        })
        self.assertEqual(data["days"][1]["exercise"], "Missed")
        self.assertEqual(data["days"][1]["junk_food"], "Had")
        self.assertEqual(data["measurements"][0], {
            "sunday": "2026-10-04", "weight_kg": 72.5, "waist_cm": 84, "chest_cm": 100,
            "bicep_cm": 35, "thigh_cm": 55, "forearm_cm": 28.5,
        })
        self.assertEqual(self.convert()["status"], "nothing_to_convert")

    def test_result_has_the_same_layout_as_a_new_workbook(self):
        self.october()
        self.convert()
        fresh = self.tmp / "fresh.xlsx"
        wbk.create_workbook(fresh, 2026, 10)
        migrated, expected = layout(self.sheet()), layout(load_workbook(fresh)["October 2026"])
        for part in expected:
            self.assertEqual(migrated[part], expected[part], part)

    def test_tables_validation_and_formats_survive(self):
        self.october()
        self.convert()
        ws = self.sheet()
        self.assertEqual({table.name: table.ref for table in ws.tables.values()},
                         {"Daily_2026_10": "A1:G32", "Measurements_2026_10": "J1:P5"})
        rules = {str(rule.sqref): rule for rule in ws.data_validations.dataValidation}
        self.assertEqual(len(rules), 7)
        self.assertEqual((rules["B2:B32"].type, rules["B2:B32"].formula1),
                         ("list", '"Completed,Partial,Rest Day,Missed"'))
        self.assertEqual((rules["C2:C32"].type, rules["C2:C32"].formula1), ("list", '"None,Controlled,Had"'))
        self.assertEqual((rules["E2:E32"].type, rules["F2:F32"].type, rules["G2:G32"].type),
                         ("whole", "decimal", "decimal"))
        self.assertEqual((rules["D2:D32"].type, rules["K2:P5"].type), ("custom", "custom"))
        self.assertEqual([ws[ref].number_format for ref in ("A2", "B2", "D2", "E2", "F2", "G2", "J2", "K2")],
                         [wbk.DATE_FORMAT, "@", "@", "0", "0.0", "0.0", wbk.DATE_FORMAT, "0.0"])

    def test_converted_workbook_opens_and_can_take_a_new_month(self):
        self.october()
        self.convert()
        self.assertEqual(load_workbook(self.path).sheetnames, ["October 2026"])
        streamed = load_workbook(self.path, read_only=True, data_only=True)
        self.addCleanup(streamed.close)
        self.assertEqual(streamed.sheetnames, ["October 2026"])
        streamed.close()                 # this mode keeps the file open until it is closed
        self.assertEqual(self.store.create_month(2026, 11)["label"], "November 2026")
        self.assertEqual(load_workbook(self.path).sheetnames, ["October 2026", "November 2026"])


class FailureTests(Case):
    def spoil(self, change):
        """Run the conversion with the rebuilt sheet altered just before it is saved."""
        real = migrate._write_sheet

        def wrong(ws, content):
            real(ws, content)
            change(ws)

        with mock.patch.object(migrate, "_write_sheet", wrong):
            with self.assertRaises(migrate.ConversionError) as caught:
                self.convert()
        return str(caught.exception)

    def test_any_difference_in_the_result_leaves_the_original_untouched(self):
        def drop_lists(ws):
            ws.data_validations.dataValidation = [
                rule for rule in ws.data_validations.dataValidation if rule.type != "list"]

        def drop_table(ws):
            del ws.tables["Measurements_2026_10"]

        changes = {
            "exercise mapped wrongly": lambda ws: ws.__setitem__("B2", "Partial"),
            "junk food mapped wrongly": lambda ws: ws.__setitem__("C3", "None"),
            "blank filled in": lambda ws: ws.__setitem__("B4", "Missed"),
            "cardio changed": lambda ws: ws.__setitem__("D2", "25:31"),
            "weight changed": lambda ws: ws.__setitem__("G2", 2450),
            "weight lost": lambda ws: ws.__setitem__("G5", None),
            "calories filled in": lambda ws: ws.__setitem__("E2", 2000),
            "protein filled in": lambda ws: ws.__setitem__("F9", 100),
            "measurement lost": lambda ws: ws.__setitem__("K2", None),
            "measurement changed": lambda ws: ws.__setitem__("P5", 31),
            "date changed": lambda ws: ws.__setitem__("A2", date(2026, 10, 2)),
            "sunday changed": lambda ws: ws.__setitem__("J2", date(2026, 10, 5)),
            "gap columns used": lambda ws: ws.__setitem__("H2", 1),
            "header changed": lambda ws: ws.__setitem__("G1", "Total Weight Lifted (kg)"),
            "something outside the tables": lambda ws: ws.__setitem__("R2", "x"),
            "row below the table": lambda ws: ws.__setitem__("A40", "x"),
            "lists missing": drop_lists,
            "table missing": drop_table,
        }
        for what, change in changes.items():
            with self.subTest(what):
                for leftover in self.tmp.iterdir():
                    leftover.unlink()
                before = self.october()
                message = self.spoil(change)
                self.assertIn("failed its check", message)
                self.assertIn("not changed", message)
                self.assertEqual(sha256(self.path), before)
                # The workbook is the only file there: no temporary file is left, and no copy was made.
                self.assertEqual(self.files(), ["Fitness_Tracker.xlsx"])
                self.assertIs(migrate.is_earlier_sheet(self.sheet()), True)

    def test_a_changed_untouched_sheet_fails_the_check(self):
        before = self.build({(2026, 10): EARLIER, (2026, 11): CURRENT}, {"October 2026": OCTOBER_EARLIER,
                                                              "November 2026": {"B2": "Partial"}})
        message = self.spoil(lambda ws: ws.parent["November 2026"].__setitem__("B2", "Completed"))
        self.assertIn("November 2026 should not have changed", message)
        self.assertEqual(sha256(self.path), before)

    def test_save_that_cannot_be_written_leaves_the_original_untouched(self):
        before = self.october()
        with mock.patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                self.convert()
        self.assertEqual(sha256(self.path), before)
        self.assertEqual(self.files(), ["Fitness_Tracker.xlsx"])

    def test_saved_file_that_does_not_open_leaves_the_original_untouched(self):
        before = self.october()
        real = load_workbook

        def broken(target, *args, **kwargs):
            if Path(target).name.startswith(".tracker-"):
                raise ValueError("not a workbook")
            return real(target, *args, **kwargs)

        with mock.patch.object(st, "load_workbook", broken):
            with self.assertRaises(ValueError):
                self.convert()
        self.assertEqual(sha256(self.path), before)
        self.assertEqual(self.files(), ["Fitness_Tracker.xlsx"])

    def test_workbook_open_in_excel_is_refused_before_anything_happens(self):
        before = self.october()
        with mock.patch("builtins.open", side_effect=PermissionError("locked")):
            with self.assertRaises(st.WorkbookLockedError):
                self.convert()
        self.assertEqual(sha256(self.path), before)
        self.assertEqual(self.files(), ["Fitness_Tracker.xlsx"])

    def test_workbook_locked_at_the_swap_leaves_the_original_untouched(self):
        before = self.october()
        with mock.patch.object(st.os, "replace", side_effect=PermissionError("locked")):
            with self.assertRaises(st.WorkbookLockedError):
                self.convert()
        self.assertEqual(sha256(self.path), before)
        self.assertEqual(self.files(), ["Fitness_Tracker.xlsx"])

    def test_missing_workbook_is_reported(self):
        with self.assertRaises(st.WorkbookNotFoundError):
            self.convert()
        self.assertEqual(self.files(), [])

    def test_conversion_never_goes_through_excel(self):
        import inspect
        source = inspect.getsource(migrate)
        for word in ("excel_live", "win32com", "pythoncom", "Dispatch", "_read("):
            self.assertNotIn(word, source)

    def test_a_failed_check_is_one_of_the_trackers_own_errors(self):
        self.assertEqual(migrate.ConversionError.code, "conversion_failed")
        self.assertTrue(issubclass(migrate.ConversionError, st.TrackerError))

    def test_no_copy_of_the_workbook_is_ever_made(self):
        import inspect
        source = inspect.getsource(migrate)
        for word in ("shutil", "copy2", "copyfile", "backup", ".bak"):
            self.assertNotIn(word, source, word)
        before = self.october()
        seen = []
        real = st.os.replace

        def watch(source_path, target):
            seen.append(sorted(p.name for p in self.tmp.iterdir()))
            return real(source_path, target)

        with mock.patch.object(st.os, "replace", watch):
            self.convert()
        # While it works there is the workbook and the one temporary file it is about to swap in; then only the workbook.
        self.assertEqual(len(seen), 1)
        self.assertEqual(len(seen[0]), 2)
        self.assertTrue(seen[0][0].startswith(".tracker-"))
        self.assertEqual(self.files(), ["Fitness_Tracker.xlsx"])
        self.assertNotEqual(sha256(self.path), before)


class IdempotencyTests(Case):
    def test_second_run_changes_nothing(self):
        self.october()
        self.convert()
        after = sha256(self.path)
        again = self.convert()
        self.assertEqual(again, {
            "status": "nothing_to_convert", "converted": [], "unchanged": ["October 2026"],
            "unsupported": [], "carried_over": [],
        })
        self.assertEqual(sha256(self.path), after)
        self.assertEqual(self.files(), ["Fitness_Tracker.xlsx"])

    def test_workbook_already_in_the_trackers_layout_is_left_alone(self):
        before = self.build({(2026, 10): CURRENT, (2026, 11): CURRENT},
                            {"October 2026": {"B2": "Rest Day", "E2": 2200, "K2": 72.5}})
        modified = self.path.stat().st_mtime_ns
        result = self.convert()
        self.assertEqual((result["status"], result["converted"]), ("nothing_to_convert", []))
        self.assertEqual(result["unchanged"], ["October 2026", "November 2026"])
        self.assertEqual(sha256(self.path), before)
        self.assertEqual(self.path.stat().st_mtime_ns, modified)
        self.assertEqual(self.files(), ["Fitness_Tracker.xlsx"])

    def test_statuses_are_never_mapped_a_second_time(self):
        self.october()
        self.convert()
        first = grid(self.sheet())
        for _ in range(3):
            self.convert()
        self.assertEqual(grid(self.sheet()), first)
        self.assertEqual(self.sheet()["B3"].value, "Missed")


class MultiMonthTests(Case):
    def test_several_months_are_all_converted(self):
        self.build({(2026, 10): EARLIER, (2026, 11): EARLIER, (2027, 2): EARLIER}, {
            "October 2026": OCTOBER_EARLIER,
            "November 2026": {"B2": False, "C2": True, "D31": "40:00", "E31": 900, "I2": 71},
            "February 2027": {"B29": True, "E29": 3000, "N5": 29.5},
        })
        result = self.convert()
        self.assertEqual(result["converted"], ["October 2026", "November 2026", "February 2027"])
        wb = load_workbook(self.path)
        self.assertEqual(wb.sheetnames, ["October 2026", "November 2026", "February 2027"])
        self.assertEqual({wbk.is_month_sheet(ws) for ws in wb}, {True})
        november, february = wb["November 2026"], wb["February 2027"]
        self.assertEqual((november["B2"].value, november["C2"].value, november["D31"].value,
                          november["G31"].value, november["K2"].value), ("Missed", "None", "40:00", 900, 71))
        self.assertEqual({t.name: t.ref for t in november.tables.values()},
                         {"Daily_2026_11": "A1:G31", "Measurements_2026_11": "J1:P6"})
        self.assertEqual((february["B29"].value, february["G29"].value, february["P5"].value),
                         ("Completed", 3000, 29.5))
        self.assertEqual({t.name: t.ref for t in february.tables.values()},
                         {"Daily_2027_02": "A1:G29", "Measurements_2027_02": "J1:P5"})
        self.assertEqual(len(self.store.list_months()), 3)
        for year, month in ((2026, 10), (2026, 11), (2027, 2)):
            self.assertEqual(self.store.get_month(year, month)["issues"], [])

    def test_only_the_months_in_the_earlier_layout_are_converted(self):
        self.build({(2026, 10): EARLIER, (2026, 11): CURRENT, (2026, 12): EARLIER}, {
            "October 2026": OCTOBER_EARLIER,
            "November 2026": {"B2": "Partial", "C2": "Controlled", "E2": 2200, "F2": 140.5, "G2": 900,
                              "K2": 71.5},
            "December 2026": {"B2": True},
        })
        november_before = load_workbook(self.path)["November 2026"]
        result = self.convert()
        self.assertEqual((result["status"], result["converted"], result["unchanged"]),
                         ("converted", ["October 2026", "December 2026"], ["November 2026"]))
        wb = load_workbook(self.path)
        self.assertEqual(wb.sheetnames, ["October 2026", "November 2026", "December 2026"])
        self.assertEqual(grid(wb["November 2026"]), grid(november_before))
        self.assertEqual(layout(wb["November 2026"]), layout(november_before))
        self.assertEqual((wb["November 2026"]["E2"].value, wb["November 2026"]["F2"].value), (2200, 140.5))
        self.assertEqual(wb["December 2026"]["B2"].value, "Completed")
        self.assertEqual({wbk.is_month_sheet(ws) for ws in wb}, {True})

    def test_other_sheets_are_untouched_and_stay_in_place(self):
        def notes(wb):
            ws = wb.create_sheet("Notes", 0)
            ws["A1"], ws["B7"], ws["C3"] = "Gym notes", 12.5, True
            ws["A2"] = "=1+1"
            ws.column_dimensions["A"].width = 40
            wb.create_sheet("Plan")["A1"] = "Push, pull, legs"

        self.build({(2026, 10): EARLIER}, {"October 2026": OCTOBER_EARLIER}, extra=notes)
        before = load_workbook(self.path)
        result = self.convert()
        self.assertEqual((result["converted"], result["unsupported"]), (["October 2026"], []))
        wb = load_workbook(self.path)
        self.assertEqual(wb.sheetnames, ["Notes", "October 2026", "Plan"])
        for name in ("Notes", "Plan"):
            self.assertEqual(grid(wb[name]), grid(before[name]))
        self.assertEqual((wb["Notes"]["A2"].value, wb["Notes"].column_dimensions["A"].width), ("=1+1", 40))

    def test_unsupported_month_is_left_exactly_as_it_was(self):
        self.build({(2026, 10): EARLIER, (2026, 11): EARLIER, (2026, 12): EARLIER}, {
            "October 2026": OCTOBER_EARLIER,
            "November 2026": {"B1": "Workout", "B2": True, "E2": 500},          # not a known layout
            "December 2026": {"A5": date(2026, 12, 25), "B2": True},            # dates are not the month's
        })
        before = load_workbook(self.path)
        result = self.convert()
        self.assertEqual((result["status"], result["converted"]), ("converted", ["October 2026"]))
        self.assertEqual(result["unsupported"], ["November 2026", "December 2026"])
        wb = load_workbook(self.path)
        for name in ("November 2026", "December 2026"):
            self.assertEqual(grid(wb[name]), grid(before[name]))
            self.assertEqual(layout(wb[name]), layout(before[name]))
        self.assertIs(wb["November 2026"]["B2"].value, True)
        self.assertEqual([(wbk.is_month_sheet(wb[name]), migrate.is_earlier_sheet(wb[name])) for name in wb.sheetnames],
                         [(True, False), (False, False), (False, True)])

    def test_workbook_with_nothing_it_can_convert_is_not_changed(self):
        def not_a_tracker(wb):
            wb.create_sheet("Sheet1")["A1"] = "hello"

        for months, cells, extra in (
            ({(2026, 10): EARLIER}, {"October 2026": {"C1": "Sweets"}}, None),
            ({(2026, 10): EARLIER}, {"October 2026": {"H3": date(2026, 10, 12)}}, None),
            ({(2026, 10): CURRENT, (2026, 11): CURRENT}, {"November 2026": {"K1": "Mass"}}, None),
            ({}, None, not_a_tracker),
        ):
            with self.subTest(cells=cells):
                before = self.build(months, cells, extra)
                result = self.convert()
                self.assertEqual((result["status"], result["converted"]), ("nothing_to_convert", []))
                self.assertEqual(sha256(self.path), before)
                self.assertEqual(self.files(), ["Fitness_Tracker.xlsx"])

    def test_file_that_is_not_a_workbook_is_reported_and_not_changed(self):
        self.path.write_bytes(b"this is not an Excel file")
        with self.assertRaises(st.WorkbookUnreadableError):
            self.convert()
        self.assertEqual(self.path.read_bytes(), b"this is not an Excel file")
        self.assertEqual(self.files(), ["Fitness_Tracker.xlsx"])


class BadCellTests(Case):
    def test_cells_that_are_not_true_or_false_are_carried_over_unchanged(self):
        self.build({(2026, 10): EARLIER}, {"October 2026": {
            "B2": "yes", "C2": 1, "D2": "25:75", "E2": "heavy", "B3": True, "C3": "no",
            "D3": 12.5, "E3": -40, "I2": "seventy", "J2": -3,
        }})
        result = self.convert()
        self.assertEqual(result["status"], "converted")
        self.assertEqual(result["carried_over"], [
            {"sheet": "October 2026", "date": "2026-10-01", "column": "Exercise", "value": "yes"},
            {"sheet": "October 2026", "date": "2026-10-01", "column": "Junk Food", "value": "1"},
            {"sheet": "October 2026", "date": "2026-10-02", "column": "Junk Food", "value": "no"},
        ])
        ws = self.sheet()
        self.assertEqual([ws[ref].value for ref in ("B2", "C2", "D2", "G2", "B3", "C3", "D3", "G3", "K2", "L2")],
                         ["yes", 1, "25:75", "heavy", "Completed", "no", 12.5, -40, "seventy", -3])
        # Nothing is guessed: the month reads, and each of those cells is reported.
        data = self.store.get_month(2026, 10)
        self.assertEqual(sorted((i["date"], i["field"]) for i in data["issues"]), [
            ("2026-10-01", "cardio"), ("2026-10-01", "exercise"), ("2026-10-01", "junk_food"),
            ("2026-10-01", "weight_lifted_kg"), ("2026-10-02", "cardio"), ("2026-10-02", "junk_food"),
            ("2026-10-02", "weight_lifted_kg"), ("2026-10-04", "waist_cm"), ("2026-10-04", "weight_kg"),
        ])
        self.assertEqual(data["days"][1]["exercise"], "Completed")

    def test_true_and_false_typed_as_text_are_not_treated_as_ticks(self):
        self.build({(2026, 10): EARLIER}, {"October 2026": {"B2": "TRUE", "C2": "FALSE"}})
        result = self.convert()
        self.assertEqual([c["value"] for c in result["carried_over"]], ["TRUE", "FALSE"])
        self.assertEqual((self.sheet()["B2"].value, self.sheet()["C2"].value), ("TRUE", "FALSE"))


class OnlyWhenAskedTests(Case):
    def test_the_tracker_itself_never_converts_and_never_reads_the_earlier_layout(self):
        before = self.october()
        client = create_app(self.path).test_client()
        for url in ("/", "/api/status", "/api/months"):
            response = client.get(url)
            self.assertEqual(response.status_code, 200, url)
            response.close()
        for url in ("/api/months/2026/10", "/api/months/2026/10/analytics"):
            response = client.get(url)
            self.assertEqual((response.status_code, response.get_json()["error"]["code"]), (500, "workbook_format"), url)
        self.assertEqual(client.put("/api/months/2026/10/days/1", json={"exercise": "Completed"}).status_code, 500)
        self.assertEqual(sha256(self.path), before)
        self.assertEqual(self.files(), ["Fitness_Tracker.xlsx"])
        # Nothing in the tracker's own modules knows the converter.
        import inspect
        from backend import analytics, app, launch
        for module in (st, wbk, app, analytics, launch):
            self.assertNotIn("migrate", inspect.getsource(module), module.__name__)

    def test_adding_a_month_converts_nothing(self):
        self.october()
        self.store.create_month(2026, 11)
        wb = load_workbook(self.path)
        self.assertEqual([wbk.is_month_sheet(ws) for ws in wb], [False, True])     # the new one is the tracker's
        self.assertIs(migrate.is_earlier_sheet(wb["October 2026"]), True)
        self.assertIs(wb["October 2026"]["B2"].value, True)
        self.assertEqual(self.files(), ["Fitness_Tracker.xlsx"])

if __name__ == "__main__":
    unittest.main()

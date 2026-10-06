"""Builds a monthly sheet in the tracker's earlier layout, for the converter's tests.

The tracker itself no longer makes or reads this layout; backend/migrate.py
converts it. Five daily columns in A to E (Date, Exercise, No Junk Food, Cardio,
Total Weight Lifted) with TRUE/FALSE habits, and the measurements in H to N.
"""

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from backend import workbook as wbk
from backend.migrate import EARLIER_DAILY_HEADERS, EARLIER_MEASUREMENT_HEADERS

FIRST_MEASUREMENT_COL = 8  # column H


def add_earlier_month_sheet(wb, year, month):
    """Add an empty month in the earlier layout and return the sheet."""
    name = wbk.sheet_name(year, month)
    ws = wb.create_sheet(name)
    dates, sundays = wbk.month_dates(year, month), wbk.month_sundays(year, month)
    wbk._write_headers(ws, 1, EARLIER_DAILY_HEADERS)
    wbk._write_headers(ws, FIRST_MEASUREMENT_COL, EARLIER_MEASUREMENT_HEADERS)

    first = wbk.HEADER_ROW + 1
    for row, day in enumerate(dates, start=first):
        ws.cell(row, 1, day).number_format = wbk.DATE_FORMAT
        ws.cell(row, 4).number_format = wbk.TEXT_FORMAT
        ws.cell(row, 5).number_format = wbk.NUMBER_FORMAT
    last_col = FIRST_MEASUREMENT_COL + len(EARLIER_MEASUREMENT_HEADERS) - 1
    for row, day in enumerate(sundays, start=first):
        ws.cell(row, FIRST_MEASUREMENT_COL, day).number_format = wbk.DATE_FORMAT
        for col in range(FIRST_MEASUREMENT_COL + 1, last_col + 1):
            ws.cell(row, col).number_format = wbk.NUMBER_FORMAT

    daily_last, meas_last = wbk.HEADER_ROW + len(dates), wbk.HEADER_ROW + len(sundays)
    suffix = f"{year}_{month:02d}"
    wbk._add_table(ws, f"Daily_{suffix}", f"A1:E{daily_last}")
    wbk._add_table(ws, f"Measurements_{suffix}", f"H1:{get_column_letter(last_col)}{meas_last}")

    done = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    done.add(f"B{first}:C{daily_last}")
    lifted = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    lifted.add(f"E{first}:E{daily_last}")
    measure = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
    measure.add(f"I{first}:{get_column_letter(last_col)}{meas_last}")
    for rule in (done, lifted, measure):
        rule.showErrorMessage = True
        ws.add_data_validation(rule)
    ws.freeze_panes = ws.cell(first, 1)
    return ws


def add_month(wb, year, month, earlier=False):
    """A month in the tracker's layout, or in the earlier one."""
    return add_earlier_month_sheet(wb, year, month) if earlier else wbk.add_month_sheet(wb, year, month)

"""Read and write access to Fitness_Tracker.xlsx for the local backend.

Every read loads the workbook fresh from disk, so edits saved in Excel show up
on the next request. Each operation touches exactly one monthly sheet; nothing
here combines data across months.

When Excel holds the file so that it cannot be read (it does this for files in
a OneDrive folder), the saved contents are read from Excel instead; see
excel_live.py. Only saved data is ever returned. Writing still needs the file
itself, so it stays blocked while Excel has the workbook open.
"""

import os
import tempfile
import threading
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

from backend import excel_live
from backend import validation
from backend import workbook as wbk

MIN_YEAR = 2000
MAX_YEAR = 2100

MEASUREMENT_KEYS = [
    "sunday", "weight_kg", "waist_cm", "chest_cm", "bicep_cm", "thigh_cm", "forearm_cm",
]


class TrackerError(Exception):
    code = "tracker_error"


class InvalidMonthError(TrackerError):
    code = "invalid_month"


class MonthNotFoundError(TrackerError):
    code = "month_not_found"


class MonthExistsError(TrackerError):
    code = "month_exists"


class WorkbookNotFoundError(TrackerError):
    code = "workbook_missing"


class WorkbookLockedError(TrackerError):
    code = "workbook_locked"


class WorkbookUnreadableError(TrackerError):
    code = "workbook_unreadable"


class WorkbookFormatError(TrackerError):
    code = "workbook_format"


class DayNotFoundError(TrackerError):
    code = "day_not_found"


class InvalidUpdateError(TrackerError):
    """Values that may not be stored. `fields` maps each one to what is wrong with it."""
    code = "invalid_request"

    def __init__(self, fields):
        super().__init__(" ".join(fields.values()))
        self.fields = dict(fields)


def validate_month(year, month):
    for name, value in (("year", year), ("month", month)):
        if isinstance(value, bool) or not isinstance(value, int):
            raise InvalidMonthError(f"{name} must be a whole number.")
    if not MIN_YEAR <= year <= MAX_YEAR:
        raise InvalidMonthError(f"year must be between {MIN_YEAR} and {MAX_YEAR}.")
    if not 1 <= month <= 12:
        raise InvalidMonthError("month must be between 1 and 12.")


class TrackerStore:
    def __init__(self, path=wbk.WORKBOOK_PATH):
        self.path = Path(path)
        self._write_lock = threading.Lock()
        self._cache_lock = threading.Lock()
        self._last_saved = None     # (file time, workbook) of the last saved state read

    def status(self):
        """Report when the workbook last changed, whether it can be written, and
        whether Excel is holding changes that have not been saved."""
        if not self.path.exists():
            raise WorkbookNotFoundError(f"Workbook not found: {self.path.name}")
        modified = datetime.fromtimestamp(self.path.stat().st_mtime)
        try:
            self._ensure_writable()
            locked = False
        except WorkbookLockedError:
            locked = True
        unsaved = locked and excel_live.inspect(self.path, read=False).state == excel_live.UNSAVED
        return {"workbook_modified": modified.isoformat(), "locked": locked,
                "unsaved_changes": unsaved}

    def list_months(self):
        wb = self._read()
        return [_month_summary(year, month) for year, month in wbk.list_months(wb)]

    def get_month(self, year, month):
        validate_month(year, month)
        wb = self._read()
        return _read_month(wb, year, month)

    def _read(self):
        """The saved workbook: from the file, or from Excel when Excel holds the file."""
        stamp = self._file_time()
        try:
            wb = self._load()
        except WorkbookLockedError as locked:
            return self._read_from_excel(locked)
        self._remember(stamp, wb)
        return wb

    def _read_from_excel(self, locked):
        stamp = self._file_time()
        with self._cache_lock:
            last = self._last_saved
        found = excel_live.inspect(self.path, read=False)
        if found.state == excel_live.SAVED:
            if last is not None and last[0] == stamp:
                return last[1]          # nothing has been saved since it was read
            found = excel_live.inspect(self.path, read=True)
            if found.state == excel_live.SAVED:
                if self._file_time() == stamp:
                    self._remember(stamp, found.workbook)
                return found.workbook
        if found.state in (excel_live.UNSAVED, excel_live.BUSY):
            # Never show what has not been saved. Keep the last saved state.
            if last is not None:
                return last[1]
            if found.state == excel_live.UNSAVED:
                raise WorkbookLockedError(
                    "The workbook is open in Excel with changes that have not been saved, "
                    "and its saved contents cannot be read while Excel has it open. "
                    "Save it in Excel (Ctrl+S) and the dashboard will load it."
                ) from locked
            raise WorkbookLockedError(
                "The workbook is open in Excel and Excel is busy (a cell is being edited "
                "or a dialog is open), so its saved contents cannot be read right now. "
                "The dashboard will try again."
            ) from locked
        raise locked    # not open in Excel, or Excel cannot be asked on this computer

    def _file_time(self):
        try:
            return self.path.stat().st_mtime_ns
        except OSError:
            return None

    def _remember(self, stamp, wb):
        with self._cache_lock:
            self._last_saved = (stamp, wb)

    def update_day(self, year, month, day, fields, today=None):
        """Change some of one day's fields.

        fields holds only what is to change: a value replaces, None clears,
        and a field left out is kept. Everything is checked first; if any of
        it is wrong nothing is written. Days after today cannot be edited.
        """
        return self._update(year, month, day, fields, today, sunday=False)

    def update_measurements(self, year, month, day, fields, today=None):
        """Change some of one Sunday's measurements. Same rules as update_day."""
        return self._update(year, month, day, fields, today, sunday=True)

    def _update(self, year, month, day, fields, today, sunday):
        validate_month(year, month)
        when, values = _checked_update(year, month, day, fields, today or date.today(), sunday)
        name = wbk.sheet_name(year, month)
        with self._write_lock:
            self._ensure_writable()         # open in Excel: refused, nothing done
            wb = self._load()               # always the file as it is on disk now
            if name not in wb.sheetnames:
                raise MonthNotFoundError(f"{name} does not exist.")
            ws = wb[name]
            _read_month(wb, year, month)    # raises if the tables or dates are not as expected

            expected = _snapshot(wb)
            cells = _target_cells(ws, year, month, when, values, sunday)
            expected[name]["cells"].update(cells)
            for cell in [cell for cell, value in cells.items() if value is None]:
                expected[name]["cells"].pop(cell, None)
            _write_cells(ws, cells)

            saved_record = {}

            def verify(saved):
                if _snapshot(saved) != expected:
                    raise WorkbookUnreadableError(
                        "The saved workbook failed verification. Nothing was changed."
                    )
                saved_record.update(_record(_read_month(saved, year, month), when, sunday))

            self._save(wb, name, verify)
            result = _month_summary(year, month)
            result["measurement" if sunday else "day"] = saved_record
            result["workbook_modified"] = datetime.fromtimestamp(self.path.stat().st_mtime).isoformat()
            return result

    def create_month(self, year, month):
        validate_month(year, month)
        with self._write_lock:
            self._ensure_writable()
            wb = self._load()
            name = wbk.sheet_name(year, month)
            if name in wb.sheetnames:
                raise MonthExistsError(f"{name} already exists.")
            ws = wbk.add_month_sheet(wb, year, month)
            _move_into_calendar_order(wb, ws, year, month)
            self._save(wb, name)
            return _read_month(wb, year, month)

    def _load(self):
        if not self.path.exists():
            raise WorkbookNotFoundError(f"Workbook not found: {self.path.name}")
        try:
            return load_workbook(self.path)
        except PermissionError as exc:
            raise WorkbookLockedError(
                "The workbook could not be read because another program has it locked."
            ) from exc
        except (BadZipFile, InvalidFileException, KeyError, ValueError, OSError) as exc:
            raise WorkbookUnreadableError(f"The workbook could not be read: {exc}") from exc

    def _ensure_writable(self):
        if not self.path.exists():
            raise WorkbookNotFoundError(f"Workbook not found: {self.path.name}")
        try:
            with open(self.path, "r+b"):
                pass
        except PermissionError as exc:
            raise WorkbookLockedError(
                "The workbook is open in Excel or locked by another program. "
                "Close it and try again. Nothing was changed."
            ) from exc

    def _save(self, wb, expected_sheet=None, verify=None):
        """Write to a temporary file, check it, then swap it in.

        The real workbook is only ever replaced by a complete, re-readable
        file. If anything fails it is left exactly as it was. verify, when
        given, is handed the re-read file and raises if it is not right.
        """
        fd, tmp = tempfile.mkstemp(dir=self.path.parent, prefix=".tracker-", suffix=".xlsx")
        os.close(fd)
        try:
            wb.save(tmp)
            saved = load_workbook(tmp)
            if expected_sheet is not None and expected_sheet not in saved.sheetnames:
                raise WorkbookUnreadableError("The saved workbook failed verification.")
            if verify is not None:
                verify(saved)
            os.replace(tmp, self.path)
        except PermissionError as exc:
            raise WorkbookLockedError(
                "The workbook is open in Excel or locked by another program. "
                "Close it and try again. Nothing was changed."
            ) from exc
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)


def _month_summary(year, month):
    return {"year": year, "month": month, "label": wbk.sheet_name(year, month)}


def _checked_update(year, month, day, fields, today, sunday):
    """The date and the values in stored form, or InvalidUpdateError listing everything wrong."""
    try:
        when = validation.validate_day(year, month, day)
    except validation.ValidationError as error:
        raise DayNotFoundError(f"{wbk.sheet_name(year, month)} has no day {day}.") from error
    errors = {}
    if sunday and when.weekday() != 6:
        errors["day"] = f"{when:%d %B %Y} is not a Sunday. Measurements go on Sundays."
    elif when > today:
        errors["day"] = ("A Sunday in the future cannot be edited yet." if sunday
                         else "A day in the future cannot be edited yet.")
    check = validation.validate_measurement_update if sunday else validation.validate_daily_update
    values = {}
    try:
        values = check(fields)
    except validation.ValidationError as error:
        errors.update(error.errors)
    if errors:
        raise InvalidUpdateError(errors)
    return when, values


def _target_cells(ws, year, month, when, values, sunday):
    """Where each value goes: {(row, column): value}, found by the date in the table."""
    table = f"{'Measurements' if sunday else 'Daily'}_{year}_{month:02d}"
    names = validation.MEASUREMENT_FIELDS if sunday else validation.DAILY_FIELDS
    min_col, min_row, _, max_row = range_boundaries(ws.tables[table].ref)
    rows = []
    for row in range(min_row + 1, max_row + 1):
        value = ws.cell(row, min_col).value
        if isinstance(value, datetime):
            value = value.date()
        if value == when:
            rows.append(row)
    if len(rows) != 1:
        raise WorkbookFormatError(f"{when.isoformat()} was not found exactly once in {ws.title}.")
    # The first column is the date; the fields follow in the order of the headers.
    return {(rows[0], min_col + 1 + names.index(field)): value for field, value in values.items()}


def _write_cells(ws, cells):
    for (row, column), value in cells.items():
        ws.cell(row, column).value = value


def _snapshot(wb):
    """Everything a write must leave alone, per sheet, in workbook order."""
    return {
        ws.title: {
            "position": position,
            "cells": {(cell.row, cell.column): cell.value
                      for row in ws.iter_rows() for cell in row if cell.value is not None},
            "formats": {(cell.row, cell.column): cell.number_format
                        for row in ws.iter_rows() for cell in row if cell.has_style},
            "tables": {table.name: table.ref for table in ws.tables.values()},
            "validations": sorted(
                (str(rule.sqref), str(rule.type), str(rule.operator), str(rule.formula1), str(rule.formula2))
                for rule in ws.data_validations.dataValidation
            ),
        }
        for position, ws in enumerate(wb.worksheets)
    }


def _record(month_data, when, sunday):
    rows, key = (month_data["measurements"], "sunday") if sunday else (month_data["days"], "date")
    return next(row for row in rows if row[key] == when.isoformat())


def _move_into_calendar_order(wb, ws, year, month):
    current = wb.index(ws)
    for index, name in enumerate(wb.sheetnames):
        other = wbk.parse_sheet_name(name)
        if other is not None and other > (year, month):
            wb.move_sheet(ws, offset=index - current)
            return


def _read_month(wb, year, month):
    """One month. Every cell is checked by the rules in validation.py."""
    name = wbk.sheet_name(year, month)
    if name not in wb.sheetnames:
        raise MonthNotFoundError(f"{name} does not exist.")
    ws = wb[name]
    suffix = f"{year}_{month:02d}"
    issues = []

    def checked(check, value, day, field):
        try:
            return check(value)
        except validation.ValidationError as error:
            _flag(issues, day, field, value, error.message)
            return None

    days = []
    rows = _table_rows(ws, f"Daily_{suffix}", wbk.DAILY_HEADERS)
    for day, row in _keyed_by_date(rows, wbk.month_dates(year, month), name, "Daily"):
        cardio = checked(validation.validate_cardio, row[3], day, "cardio")
        unreadable = cardio is None and row[3] is not None and str(row[3]).strip() != ""
        days.append({
            "date": day.isoformat(),
            "exercise": checked(validation.validate_exercise, row[1], day, "exercise"),
            "junk_food": checked(validation.validate_junk_food, row[2], day, "junk_food"),
            "cardio_display": str(row[3]) if unreadable else cardio,
            "cardio_seconds": validation.cardio_seconds(cardio),
            "calories_kcal": checked(validation.validate_calories, row[4], day, "calories_kcal"),
            "protein_g": checked(validation.validate_protein, row[5], day, "protein_g"),
            "weight_lifted_kg": checked(validation.validate_weight_lifted, row[6], day, "weight_lifted_kg"),
        })

    measurements = []
    rows = _table_rows(ws, f"Measurements_{suffix}", wbk.MEASUREMENT_HEADERS)
    for day, row in _keyed_by_date(rows, wbk.month_sundays(year, month), name, "Measurements"):
        record = {"sunday": day.isoformat()}
        for key, value in zip(MEASUREMENT_KEYS[1:], row[1:]):
            record[key] = checked(lambda v, key=key: validation.validate_measurement(v, key), value, day, key)
        measurements.append(record)

    result = _month_summary(year, month)
    result.update(days=days, measurements=measurements, issues=issues)
    return result


def _table_rows(ws, table_name, expected_headers):
    if table_name not in ws.tables:
        raise WorkbookFormatError(f"Table {table_name} is missing from sheet {ws.title}.")
    min_col, min_row, max_col, max_row = range_boundaries(ws.tables[table_name].ref)
    rows = list(ws.iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True
    ))
    if list(rows[0]) != expected_headers:
        raise WorkbookFormatError(
            f"Table {table_name} has unexpected columns: {list(rows[0])}. "
            f"Expected: {expected_headers}."
        )
    return rows[1:]


def _keyed_by_date(rows, expected_dates, sheet, table):
    """Pair each row with its date and insist the dates are exactly this month's."""
    keyed = []
    for row in rows:
        value = row[0]
        if isinstance(value, datetime):
            value = value.date()
        if not isinstance(value, date):
            raise WorkbookFormatError(
                f"{table} table in {sheet} has a row whose date is not a date: {row[0]!r}."
            )
        keyed.append((value, row))
    keyed.sort(key=lambda pair: pair[0])
    if [day for day, _ in keyed] != expected_dates:
        raise WorkbookFormatError(
            f"{table} table in {sheet} does not contain exactly the expected dates "
            f"for that month."
        )
    return keyed


def _flag(issues, day, field, value, message):
    issues.append({
        "date": day.isoformat(), "field": field, "value": str(value), "message": message,
    })

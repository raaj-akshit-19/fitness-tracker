"""Read the saved contents of the workbook out of Excel, while Excel has it open.

Normally the workbook file is read directly. But when the file is in a
OneDrive folder, Excel opens it through its cloud address and holds the local
file so that nothing else can read it. In that case the only place the data
can be read from is Excel itself, and that is what this module does.

It is strictly read-only. It only ever looks at a workbook that is already
open in an Excel that is already running: it never starts Excel, never opens,
saves or closes a workbook, and never changes a cell or a setting.

Only saved data is returned. Excel says whether a workbook has changes that
have not been saved; when it does, nothing is read from it.

Windows only. Without the pywin32 package (or on another system) every call
answers UNAVAILABLE and the caller behaves as it did before this existed.
"""

import datetime
import decimal
import gc
import os
import queue
import threading
from pathlib import Path
from types import SimpleNamespace
from typing import NamedTuple, Optional
from urllib.parse import unquote

from openpyxl.utils import range_boundaries

from backend import workbook as wbk

try:
    import pythoncom
    import winreg
    AVAILABLE = True
except ImportError:      # not Windows, or pywin32 is not installed
    pythoncom = winreg = None
    AVAILABLE = False

# Everything this module ever asks Excel for. Each one only reads something.
# Excel will carry out an action such as saving if it is asked for by name, so
# the names are fixed here and nothing outside this list can be requested.
READS = frozenset({
    "FullName", "Saved", "Worksheets", "ListObjects", "Name", "Range", "Address", "Value", "Formula",
})

NOT_OPEN = "not_open"        # the workbook is not open in any running Excel
SAVED = "saved"              # open in Excel, with nothing unsaved
UNSAVED = "unsaved"          # open in Excel, with changes that have not been saved
BUSY = "busy"                # open in Excel, but Excel would not answer just now
UNAVAILABLE = "unavailable"  # this computer cannot ask Excel at all

STATE_TIMEOUT = 4    # seconds to wait for Excel to say whether it is saved
READ_TIMEOUT = 10    # seconds to wait for Excel to hand over the data

# How Excel reports a cell holding an error, and what the file holds for it.
ERROR_VALUES = {
    -2146826288: "#NULL!", -2146826281: "#DIV/0!", -2146826273: "#VALUE!",
    -2146826265: "#REF!", -2146826259: "#NAME?", -2146826252: "#NUM!", -2146826246: "#N/A",
}


class Inspection(NamedTuple):
    state: str
    workbook: Optional["WorkbookSnapshot"] = None
    detail: str = ""


class SheetSnapshot:
    """One sheet's tables, as plain values. Reads like an openpyxl worksheet."""

    def __init__(self, title, tables, cells):
        self.title = title
        self.tables = {name: SimpleNamespace(ref=ref) for name, ref in tables.items()}
        self._cells = cells

    def iter_rows(self, min_row, max_row, min_col, max_col, values_only=True):
        for row in range(min_row, max_row + 1):
            yield tuple(self._cells.get((row, col)) for col in range(min_col, max_col + 1))


class WorkbookSnapshot:
    """The saved workbook as plain values. Reads like an openpyxl workbook."""

    def __init__(self, sheetnames, sheets):
        self.sheetnames = list(sheetnames)
        self._sheets = sheets

    def __getitem__(self, name):
        return self._sheets[name]


def inspect(path, read=True):
    """Ask Excel about the workbook at path.

    With read=True and the workbook saved, the result carries a snapshot of
    its monthly sheets. Never raises: any trouble comes back as a state.
    """
    if not AVAILABLE:
        return Inspection(UNAVAILABLE)
    return _worker.run(lambda: _inspect(Path(path), read), READ_TIMEOUT if read else STATE_TIMEOUT)


def _get(thing, name):
    """Read one property of one of Excel's objects.

    With _each below, this is the only way Excel is ever asked anything. Nothing
    is ever passed to Excel: there are no arguments, only a name from the list.
    """
    if name not in READS:
        raise ValueError(f"{name} is not something this module reads from Excel")
    return thing.Invoke(thing.GetIDsOfNames(name), 0, pythoncom.DISPATCH_PROPERTYGET, True)


def _each(collection):
    """Step through one of Excel's collections (its sheets, a sheet's tables), in order."""
    listing = collection.Invoke(pythoncom.DISPID_NEWENUM, 0, pythoncom.DISPATCH_PROPERTYGET, True)
    listing = listing.QueryInterface(pythoncom.IID_IEnumVARIANT)
    while True:
        batch = listing.Next(1)
        if not batch:
            break
        yield batch[0]


# ---------------------------------------------------------------- one thread talks to Excel

class _Worker:
    """Runs every conversation with Excel on a single background thread.

    One at a time, with a time limit for the caller: if Excel stops answering,
    callers get BUSY instead of waiting, and no further threads pile up.
    """

    def __init__(self):
        self._jobs = queue.Queue()
        self._turn = threading.Lock()
        self._start = threading.Lock()
        self._thread = None

    def run(self, job, timeout):
        if not self._turn.acquire(timeout=1):
            return Inspection(BUSY, detail="Excel has not answered an earlier request yet.")
        with self._start:
            if self._thread is None or not self._thread.is_alive():
                self._thread = threading.Thread(target=self._loop, name="excel-reader", daemon=True)
                self._thread.start()
        box = {"done": threading.Event(), "result": None}
        self._jobs.put((job, box))
        if box["done"].wait(timeout):
            return box["result"]
        return Inspection(BUSY, detail="Excel did not answer in time.")

    def _loop(self):
        pythoncom.CoInitialize()
        try:
            while True:
                job, box = self._jobs.get()
                try:
                    box["result"] = job()
                except Exception as error:  # noqa: BLE001 - never let the thread die
                    box["result"] = Inspection(BUSY, detail=f"{type(error).__name__}: {error}")
                finally:
                    gc.collect()        # let go of everything Excel handed over
                    box["done"].set()
                    self._turn.release()
        finally:
            pythoncom.CoUninitialize()


_worker = _Worker()


# ---------------------------------------------------------------- finding the workbook

def _normal(path):
    return os.path.normcase(os.path.realpath(str(path)))


def _onedrive_mounts():
    """(cloud address prefix, local folder) for every OneDrive folder on this computer."""
    mounts = []
    try:
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, r"Software\SyncEngines\Providers\OneDrive") as root:
            for index in range(winreg.QueryInfoKey(root)[0]):
                with winreg.OpenKey(root, winreg.EnumKey(root, index)) as provider:
                    try:
                        mounts.append((winreg.QueryValueEx(provider, "UrlNamespace")[0],
                                       winreg.QueryValueEx(provider, "MountPoint")[0]))
                    except OSError:
                        pass
    except OSError:
        pass
    for variable in ("OneDriveConsumer", "OneDrive"):
        if os.environ.get(variable):
            mounts.append(("https://d.docs.live.net", os.environ[variable]))
    return list(dict.fromkeys(mounts))


def _local_candidates(address, mounts=None):
    """Local paths a cloud address may stand for."""
    candidates = []
    for prefix, folder in (_onedrive_mounts() if mounts is None else mounts):
        prefix = prefix.rstrip("/")
        if not prefix or not address.lower().startswith(prefix.lower() + "/"):
            continue
        parts = [unquote(part) for part in address[len(prefix):].split("/") if part]
        if parts:
            candidates.append(os.path.join(folder, *parts))
        if len(parts) > 1:      # a personal OneDrive puts the account id first
            candidates.append(os.path.join(folder, *parts[1:]))
    return candidates


def _is_workbook(name, target, mounts=None):
    """Whether the name Excel gave a document is the file at target (already normalised)."""
    if name.lower().startswith(("http://", "https://")):
        return any(_normal(candidate) == target for candidate in _local_candidates(name, mounts))
    return _normal(name) == target


def _running_documents():
    """(name, object) for every document registered as running on this computer.

    Each Excel that is running lists the workbooks it has open here, so this
    covers every Excel, not just the first one started.
    """
    context = pythoncom.CreateBindCtx(0)
    table = pythoncom.GetRunningObjectTable()
    listing = table.EnumRunning()
    while True:
        batch = listing.Next(1)     # stepped through by hand: a plain "for" loop over it
        if not batch:               # would make pywin32 load its scripting helper library
            break
        moniker = batch[0]
        try:
            name = moniker.GetDisplayName(context, None)
        except pythoncom.com_error:
            continue
        yield name, lambda moniker=moniker: table.GetObject(moniker).QueryInterface(
            pythoncom.IID_IDispatch)


class ExcelBusy(Exception):
    """The workbook is open in Excel, but Excel would not answer about it."""


def _find(path):
    """The open workbook for exactly this file, or None. Never opens anything.

    Raises ExcelBusy when Excel has this file open but will not answer, which
    is what happens while a cell is being typed in or a dialog is open.
    """
    target = _normal(path)
    filename = os.path.basename(target)
    busy = None
    for name, fetch in _running_documents():
        if not name.lower().endswith(filename) or not _is_workbook(name, target):
            continue
        try:
            book = fetch()
            if _is_workbook(str(_get(book, "FullName")), target):
                return book
        except (pythoncom.com_error, AttributeError) as error:
            # It is registered under this file's name but will not say more.
            busy = error.args[0] if error.args else error
    if busy is not None:
        raise ExcelBusy(f"Excel is busy ({busy}).")
    return None


# ---------------------------------------------------------------- reading it

def _inspect(path, read):
    try:
        book = _find(path)
        if book is None:
            return Inspection(NOT_OPEN)
        if not bool(_get(book, "Saved")):
            return Inspection(UNSAVED)
        if not read:
            return Inspection(SAVED)
        snapshot = _snapshot(book)
        # If something was typed while the cells were being read, what was read
        # may include it. Throw it away; the next request tries again.
        if not bool(_get(book, "Saved")):
            return Inspection(UNSAVED)
        return Inspection(SAVED, snapshot)
    except ExcelBusy as error:
        return Inspection(BUSY, detail=str(error))
    except (pythoncom.com_error, AttributeError) as error:
        # Typically: a cell is being edited, or a dialog is open in Excel.
        return Inspection(BUSY, detail=f"Excel is busy ({error.args[0] if error.args else error}).")
    finally:
        book = snapshot = None


def _snapshot(book):
    sheets = list(_each(_get(book, "Worksheets")))
    names = [str(_get(sheet, "Name")) for sheet in sheets]
    monthly = {}
    for name, sheet in zip(names, sheets):
        if wbk.parse_sheet_name(name) is None:
            continue
        tables, cells = {}, {}
        for table in _each(_get(sheet, "ListObjects")):
            area = _get(table, "Range")
            ref = str(_get(area, "Address")).replace("$", "")
            tables[str(_get(table, "Name"))] = ref
            min_col, min_row, max_col, max_row = range_boundaries(ref)
            values, formulas = _grid(_get(area, "Value")), _grid(_get(area, "Formula"))
            for r, row in enumerate(values):
                for c, value in enumerate(row):
                    plain = _plain(value, formulas[r][c])
                    if plain is not None:
                        cells[(min_row + r, min_col + c)] = plain
        monthly[name] = SheetSnapshot(name, tables, cells)
    return WorkbookSnapshot(names, monthly)


def _grid(value):
    """Excel gives one cell as a bare value and several as rows of values."""
    return value if isinstance(value, tuple) else ((value,),)


def _plain(value, formula):
    """Turn a value from Excel into what reading the file would have given."""
    if isinstance(formula, str) and formula.startswith("="):
        return formula                      # the file holds the formula, not its result
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, datetime.datetime):
        return datetime.datetime(value.year, value.month, value.day,
                                 value.hour, value.minute, value.second)
    if isinstance(value, decimal.Decimal):
        value = float(value)
    if isinstance(value, int):
        return ERROR_VALUES.get(value, value)
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    if isinstance(value, str):
        return value if value != "" else None
    return str(value)

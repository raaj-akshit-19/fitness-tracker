"""Excel workbook layout for the local fitness tracker.

One sheet per month, named "<Month> <Year>" (for example "October 2026").
Every monthly sheet holds two independent Excel tables:

  Daily table         A1  one row per calendar day of the month
  Measurements table  J1  one row per Sunday of the month

This module is the single definition of that layout. Anything that reads or
writes the workbook should go through the constants and helpers here. The
rules for what each cell may hold live in validation.py.
"""

import calendar
import os
import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from backend import validation
from backend.paths import APP_DIR

# The workbook sits in the app's own folder: the project folder, or beside the
# .exe when packaged. FITNESS_TRACKER_WORKBOOK points the app and the tests at
# a different file.
WORKBOOK_PATH = Path(
    os.environ.get("FITNESS_TRACKER_WORKBOOK") or APP_DIR / "Fitness_Tracker.xlsx"
)

# English names are fixed here so sheet names never depend on the OS locale.
MONTH_NAMES = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]

# Daily table A to G, two blank columns, measurements J to P.
DAILY_HEADERS = [
    "Date",
    "Exercise",
    "Junk Food",
    "Cardio (min:sec)",
    "Calories (kcal)",
    "Protein (g)",
    "Weight Lifted (kg)",
]
MEASUREMENT_HEADERS = [
    "Sunday",
    "Weight (kg)",
    "Waist (cm)",
    "Chest (cm)",
    "Bicep (cm)",
    "Thigh (cm)",
    "Forearm (cm)",
]

HEADER_ROW = 1
DAILY_FIRST_COL = 1  # column A
MEASUREMENT_FIRST_COL = 10  # column J

DATE_FORMAT = "ddd, dd mmm yyyy"
TEXT_FORMAT = "@"
NUMBER_FORMAT = "0.0"
WHOLE_NUMBER_FORMAT = "0"

_CARDIO_RE = re.compile(r"^(\d{1,3}):([0-5]\d)$")
_SHEET_NAME_RE = re.compile(r"^(%s) (\d{4})$" % "|".join(MONTH_NAMES))


def sheet_name(year, month):
    return f"{MONTH_NAMES[month - 1]} {year}"


def parse_sheet_name(name):
    """Return (year, month) for a monthly sheet name, or None if it is not one."""
    match = _SHEET_NAME_RE.match(name)
    if not match:
        return None
    return int(match.group(2)), MONTH_NAMES.index(match.group(1)) + 1


def month_dates(year, month):
    days = calendar.monthrange(year, month)[1]
    return [date(year, month, day) for day in range(1, days + 1)]


def month_sundays(year, month):
    return [d for d in month_dates(year, month) if d.weekday() == calendar.SUNDAY]


def parse_cardio(value):
    """Convert a "min:sec" cardio entry to total seconds. Empty returns None."""
    if value is None or str(value).strip() == "":
        return None
    match = _CARDIO_RE.match(str(value).strip())
    if not match:
        raise ValueError(f"Cardio must be min:sec, for example 25:30. Got: {value!r}")
    return int(match.group(1)) * 60 + int(match.group(2))


def format_cardio(total_seconds):
    minutes, seconds = divmod(int(total_seconds), 60)
    return f"{minutes}:{seconds:02d}"


def list_months(wb):
    """Return [(year, month)] for every monthly sheet, in calendar order."""
    months = [parse_sheet_name(name) for name in wb.sheetnames]
    return sorted(m for m in months if m is not None)


def table_headers(ws, table_name):
    """The header names of one of a sheet's tables, or None if it has no such table."""
    if table_name not in ws.tables:
        return None
    min_col, min_row, max_col, _ = range_boundaries(ws.tables[table_name].ref)
    rows = ws.iter_rows(min_row=min_row, max_row=min_row, min_col=min_col, max_col=max_col,
                        values_only=True)
    return list(next(iter(rows)))


def is_month_sheet(ws):
    """True when a sheet is a month laid out as the tracker lays one out.

    Decided by the header names of the sheet's two tables, found by table
    name, not by where on the sheet they sit. Only looks; never changes anything.
    """
    parsed = parse_sheet_name(ws.title)
    if parsed is None:
        return False
    suffix = "%d_%02d" % parsed
    return (table_headers(ws, f"Daily_{suffix}") == DAILY_HEADERS
            and table_headers(ws, f"Measurements_{suffix}") == MEASUREMENT_HEADERS)


def add_month_sheet(wb, year, month):
    """Add an empty monthly sheet to the workbook and return it."""
    name = sheet_name(year, month)
    if name in wb.sheetnames:
        raise ValueError(f"Sheet already exists: {name}")
    ws = wb.create_sheet(name)
    dates = month_dates(year, month)
    sundays = month_sundays(year, month)
    first_col = MEASUREMENT_FIRST_COL
    last_col = first_col + len(MEASUREMENT_HEADERS) - 1

    _write_headers(ws, DAILY_FIRST_COL, DAILY_HEADERS)
    _write_headers(ws, first_col, MEASUREMENT_HEADERS)

    first = HEADER_ROW + 1
    # Every cell starts blank, which means "not entered".
    daily_formats = {2: TEXT_FORMAT, 3: TEXT_FORMAT, 4: TEXT_FORMAT,
                     5: WHOLE_NUMBER_FORMAT, 6: NUMBER_FORMAT, 7: NUMBER_FORMAT}
    for row, day in enumerate(dates, start=first):
        ws.cell(row, 1, day).number_format = DATE_FORMAT
        for col, number_format in daily_formats.items():
            ws.cell(row, col).number_format = number_format

    for row, day in enumerate(sundays, start=first):
        ws.cell(row, first_col, day).number_format = DATE_FORMAT
        for col in range(first_col + 1, last_col + 1):
            ws.cell(row, col).number_format = NUMBER_FORMAT

    daily_last = HEADER_ROW + len(dates)
    meas_last = HEADER_ROW + len(sundays)
    suffix = f"{year}_{month:02d}"
    _add_table(ws, f"Daily_{suffix}", f"A{HEADER_ROW}:G{daily_last}")
    _add_table(ws, f"Measurements_{suffix}",
               f"{get_column_letter(first_col)}{HEADER_ROW}:{get_column_letter(last_col)}{meas_last}")
    _add_validations(ws, first, daily_last, meas_last, first_col, last_col)

    widths = {1: 18, 2: 13, 3: 13, 4: 17, 5: 16, 6: 13, 7: 19, 8: 3, 9: 3, first_col: 18}
    for col in range(first_col + 1, last_col + 1):
        widths[col] = 14
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(first, 1)
    return ws


def create_workbook(path=WORKBOOK_PATH, year=2026, month=10):
    """Create a new workbook holding one empty month. Never overwrites."""
    path = Path(path)
    if path.exists():
        raise FileExistsError(f"Refusing to overwrite existing workbook: {path}")
    wb = Workbook()
    wb.remove(wb.active)
    add_month_sheet(wb, year, month)
    wb.save(path)
    return path


def open_workbook(path=WORKBOOK_PATH):
    return load_workbook(path)


def _write_headers(ws, first_col, headers):
    for offset, header in enumerate(headers):
        cell = ws.cell(HEADER_ROW, first_col + offset, header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="left")


def _add_table(ws, name, ref):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight1", showRowStripes=False, showColumnStripes=False
    )
    ws.add_table(table)


def _add_validations(ws, first, daily_last, meas_last, meas_first_col, meas_last_col):
    """What Excel itself accepts when a cell is typed into. Limits come from validation.py."""
    def listed(values, label, cells):
        rule = DataValidation(type="list", formula1='"%s"' % ",".join(values), allow_blank=True)
        rule.error = f"{label} must be one of: {', '.join(values)}; or left blank."
        rule.add(cells)
        return rule

    def between(kind, low, high, message, cells):
        rule = DataValidation(type=kind, operator="between", formula1=str(low), formula2=str(high),
                              allow_blank=True)
        rule.error = message
        rule.add(cells)
        return rule

    exercise = listed(validation.EXERCISE_STATUSES, "Exercise", f"B{first}:B{daily_last}")
    junk_food = listed(validation.JUNK_FOOD_STATUSES, "Junk Food", f"C{first}:C{daily_last}")

    # Text cells, so that Excel keeps 25:30 as typed instead of making it a time of day.
    cardio = DataValidation(
        type="custom",
        formula1=(
            f'AND(LEN(D{first})-FIND(":",D{first})=2,'
            f'ISNUMBER(--LEFT(D{first},FIND(":",D{first})-1)),'
            f"ISNUMBER(--RIGHT(D{first},2)),--RIGHT(D{first},2)<60)"
        ),
        allow_blank=True,
    )
    cardio.error = "Enter cardio as min:sec, for example 25:30."
    cardio.add(f"D{first}:D{daily_last}")

    calories = between("whole", 0, validation.CALORIES_MAX,
                       f"Enter calories as a whole number of kcal, 0 to {validation.CALORIES_MAX}.",
                       f"E{first}:E{daily_last}")
    protein = between("decimal", 0, validation.PROTEIN_MAX,
                      f"Enter protein in grams, 0 to {validation.PROTEIN_MAX}.",
                      f"F{first}:F{daily_last}")
    lifted = between("decimal", 0, validation.WEIGHT_LIFTED_MAX,
                     f"Enter the day's total weight lifted in kg, 0 to {validation.WEIGHT_LIFTED_MAX}.",
                     f"G{first}:G{daily_last}")

    # "More than zero" cannot be said with "between", so a formula checks each cell.
    top_left = f"{get_column_letter(meas_first_col + 1)}{first}"
    measure = DataValidation(
        type="custom",
        formula1=f"AND(ISNUMBER({top_left}),{top_left}>0,{top_left}<={validation.MEASUREMENT_MAX})",
        allow_blank=True,
    )
    measure.error = f"Enter a number more than 0, up to {validation.MEASUREMENT_MAX}."
    measure.add(f"{top_left}:{get_column_letter(meas_last_col)}{meas_last}")

    for rule in (exercise, junk_food, cardio, calories, protein, lifted, measure):
        rule.showErrorMessage = True
        ws.add_data_validation(rule)

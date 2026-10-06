"""Convert a workbook made with the tracker's earlier layout, once.

The tracker itself never runs this: it reads and writes the one layout that
workbook.py defines, and nothing here is called at startup, while reading, or
from the page. It is a tool run by hand, on a workbook that still holds months
laid out the earlier way:

    python -m backend.migrate [path to the workbook]

What it does, in order, and stops at the first thing that goes wrong:

  1. Looks at every monthly sheet. If none is in the earlier layout, it
     changes nothing.
  2. Rebuilds each such month in the tracker's layout, in memory.
  3. Saves that to a temporary file, opens it again and checks every value
     against what the original held.
  4. Only then swaps it in for the original.

Until step 4 the original file is not touched. No copy of the workbook is
made or kept. Months already in the tracker's layout, sheets that are in
neither, and sheets that are not months are left exactly as they are.

The earlier layout had five daily columns (Date, Exercise, No Junk Food,
Cardio, Total Weight Lifted) with the measurements starting in column H.

Mapping:   Exercise      TRUE -> Completed   FALSE -> Missed   blank -> blank
           No Junk Food  TRUE -> None        FALSE -> Had      blank -> blank
           Cardio, Total Weight Lifted and the Sunday measurements are kept.
           Calories and Protein start blank.
A cell holding something else (a typing slip) is carried over unchanged, so
nothing is lost or guessed; reading the month afterwards reports it.
"""

import sys
from datetime import date, datetime

from openpyxl.utils import range_boundaries

from backend import validation
from backend import workbook as wbk
from backend.store import TrackerError, TrackerStore, WorkbookFormatError

CONVERTED = "converted"
NOTHING_TO_CONVERT = "nothing_to_convert"

# The earlier layout. Only this module knows it.
EARLIER_DAILY_HEADERS = ["Date", "Exercise", "No Junk Food", "Cardio (min:sec)", "Total Weight Lifted (kg)"]
EARLIER_MEASUREMENT_HEADERS = wbk.MEASUREMENT_HEADERS

EXERCISE_MAP = {True: validation.EXERCISE_COMPLETED, False: validation.EXERCISE_MISSED}
JUNK_FOOD_MAP = {True: validation.JUNK_FOOD_NONE, False: validation.JUNK_FOOD_HAD}


class ConversionError(TrackerError):
    code = "conversion_failed"


def is_earlier_sheet(ws):
    """True when a monthly sheet is laid out the earlier way. Only looks."""
    parsed = wbk.parse_sheet_name(ws.title)
    if parsed is None:
        return False
    suffix = "%d_%02d" % parsed
    return (wbk.table_headers(ws, f"Daily_{suffix}") == EARLIER_DAILY_HEADERS
            and wbk.table_headers(ws, f"Measurements_{suffix}") == EARLIER_MEASUREMENT_HEADERS)


def convert(store):
    """Convert the store's workbook. Returns a summary of what was done."""
    with store._write_lock:
        store._ensure_writable()            # open in Excel: refused, nothing done
        wb = store._load()                  # always the file itself, never through Excel
        result = {"status": CONVERTED, "converted": [], "unchanged": [], "unsupported": [], "carried_over": []}
        earlier = []
        for year, month in wbk.list_months(wb):
            label = wbk.sheet_name(year, month)
            if wbk.is_month_sheet(wb[label]):
                result["unchanged"].append(label)
            elif is_earlier_sheet(wb[label]):
                earlier.append((label, year, month))
            else:
                result["unsupported"].append(label)

        # Take everything out of the earlier sheets first. A sheet whose dates
        # are not that month's is not converted.
        contents = {}
        for label, year, month in earlier:
            try:
                contents[label] = _read_earlier(wb[label], year, month)
            except WorkbookFormatError:
                result["unsupported"].append(label)
        if not contents:
            result["status"] = NOTHING_TO_CONVERT
            return result
        untouched = {name: _grid(wb[name]) for name in wb.sheetnames if name not in contents}
        order = list(wb.sheetnames)

        for label, year, month in earlier:
            if label in contents:
                _rebuild(wb, label, year, month, contents[label])
                result["converted"].append(label)
                result["carried_over"] += contents[label]["carried_over"]

        def verify(saved):
            _verify(saved, order, contents, untouched)

        store._save(wb, verify=verify)      # temporary file, checked, then swapped in
        return result


# ---------------------------------------------------------------- reading the earlier layout

def _table_cells(ws, name):
    min_col, min_row, max_col, max_row = range_boundaries(ws.tables[name].ref)
    return list(ws.iter_rows(min_row=min_row + 1, max_row=max_row, min_col=min_col,
                             max_col=max_col, values_only=True))


def _as_date(value):
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


def _read_earlier(ws, year, month):
    """The values of a month in the earlier layout, with what each becomes."""
    suffix = f"{year}_{month:02d}"
    daily = {_as_date(row[0]): row for row in _table_cells(ws, f"Daily_{suffix}")}
    sundays = {_as_date(row[0]): row for row in _table_cells(ws, f"Measurements_{suffix}")}
    if list(daily) != wbk.month_dates(year, month) or list(sundays) != wbk.month_sundays(year, month):
        raise WorkbookFormatError(f"{ws.title} does not hold exactly that month's dates.")

    carried = []

    def mapped(value, mapping, day, column):
        if value is None:
            return None
        if isinstance(value, bool):
            return mapping[value]
        carried.append({"sheet": ws.title, "date": day.isoformat(), "column": column, "value": str(value)})
        return value

    days = {}
    for day, row in daily.items():
        days[day] = {
            2: mapped(row[1], EXERCISE_MAP, day, "Exercise"),
            3: mapped(row[2], JUNK_FOOD_MAP, day, "Junk Food"),
            4: row[3],          # cardio, as it is
            5: None,            # calories
            6: None,            # protein
            7: row[4],          # weight lifted, as it is
        }
    return {
        "days": days,
        "measurements": {day: list(row[1:]) for day, row in sundays.items()},
        "carried_over": carried,
    }


# ---------------------------------------------------------------- writing the tracker's layout

def _rebuild(wb, label, year, month, content):
    """Replace a sheet in the earlier layout with one in the tracker's, in the same place."""
    position = wb.index(wb[label])
    wb.remove(wb[label])
    ws = wbk.add_month_sheet(wb, year, month)
    wb.move_sheet(ws, offset=position - wb.index(ws))
    _write_sheet(ws, content)


def _write_sheet(ws, content):
    first = wbk.HEADER_ROW + 1
    for row, values in enumerate(content["days"].values(), start=first):
        for col, value in values.items():
            ws.cell(row, col).value = value
    for row, values in enumerate(content["measurements"].values(), start=first):
        for offset, value in enumerate(values, start=1):
            ws.cell(row, wbk.MEASUREMENT_FIRST_COL + offset).value = value


# ---------------------------------------------------------------- checking the result

def _grid(ws):
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def _verify(saved, order, contents, untouched):
    """Check the saved file against the original's contents. Raises on any difference."""
    def fail(what):
        raise ConversionError(f"The converted workbook failed its check ({what}). "
                              "The workbook was not changed.")

    if saved.sheetnames != order:
        fail("the sheets are not the same")
    for name, before in untouched.items():
        if _grid(saved[name]) != before:
            fail(f"{name} should not have changed")

    for label, content in contents.items():
        year, month = wbk.parse_sheet_name(label)
        ws = saved[label]
        days, sundays = wbk.month_dates(year, month), wbk.month_sundays(year, month)
        suffix = f"{year}_{month:02d}"
        if not wbk.is_month_sheet(ws):
            fail(f"{label} does not have the tracker's headers")
        if {n: ws.tables[n].ref for n in ws.tables} != {
            f"Daily_{suffix}": f"A1:G{len(days) + 1}",
            f"Measurements_{suffix}": f"J1:P{len(sundays) + 1}",
        }:
            fail(f"{label} does not have the two tables")
        lists = {str(rule.sqref): rule.formula1 for rule in ws.data_validations.dataValidation
                 if rule.type == "list"}
        if lists != {
            f"B2:B{len(days) + 1}": '"%s"' % ",".join(validation.EXERCISE_STATUSES),
            f"C2:C{len(days) + 1}": '"%s"' % ",".join(validation.JUNK_FOOD_STATUSES),
        }:
            fail(f"{label} is missing the Exercise or Junk Food list")

        for row, (day, values) in enumerate(content["days"].items(), start=2):
            if _as_date(ws.cell(row, 1).value) != day:
                fail(f"{label}: a date changed")
            for col, value in values.items():
                if ws.cell(row, col).value != value:
                    fail(f"{label}: {day:%d %b} column {col} differs")
            if ws.cell(row, 5).value is not None or ws.cell(row, 6).value is not None:
                fail(f"{label}: Calories and Protein must start blank")
            if ws.cell(row, 8).value is not None or ws.cell(row, 9).value is not None:
                fail(f"{label}: columns H and I must be blank")
        for row, (day, values) in enumerate(content["measurements"].items(), start=2):
            if _as_date(ws.cell(row, wbk.MEASUREMENT_FIRST_COL).value) != day:
                fail(f"{label}: a Sunday changed")
            found = [ws.cell(row, wbk.MEASUREMENT_FIRST_COL + offset).value for offset in range(1, 7)]
            if found != values:
                fail(f"{label}: measurements for {day:%d %b} differ")
        if ws.max_row > len(days) + 1 or ws.max_column > 16:
            fail(f"{label}: something lies outside the two tables")


if __name__ == "__main__":
    outcome = convert(TrackerStore(sys.argv[1]) if len(sys.argv) > 1 else TrackerStore())
    print(outcome["status"])
    for key in ("converted", "unchanged", "unsupported"):
        if outcome[key]:
            print(f"  {key}: {', '.join(outcome[key])}")
    for cell in outcome["carried_over"]:
        print(f"  carried over as it was: {cell['sheet']} {cell['date']} {cell['column']} = {cell['value']}")
